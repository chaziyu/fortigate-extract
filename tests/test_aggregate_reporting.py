from __future__ import annotations

import io
import unittest

from openpyxl import load_workbook

from fortigate_extract.config import ExtractionConfig
from fortigate_extract.derived import build_derived_views
from fortigate_extract.export import export_excel
from fortigate_extract.extraction.extractor import extract_fortigate_config
from fortigate_extract.extraction.result import ExtractionResult
from fortigate_extract.model.interface import FGInterface
from fortigate_extract.model.source import FGConfig
from fortigate_extract.parser import parse_fortigate_config
from fortigate_extract.validation.validator import validate_config


_CONFIG = r'''
config system interface
    edit "port1"
    next
    edit "port2"
    next
    edit "agg1"
        set type aggregate
        set member "port1" "port2"
    next
    edit "vlan100"
        set type vlan
        set interface "agg1"
    next
end
'''


class AggregateReportingTest(unittest.TestCase):
    def test_standalone_logical_interfaces_do_not_report_missing_parents(self):
        config = FGConfig(
            interfaces=[
                FGInterface(name="l2t.root", type="tunnel"),
                FGInterface(name="naf.root", type="tunnel"),
                FGInterface(name="ssl.root", type="tunnel"),
                FGInterface(name="loop1", type="loopback"),
                FGInterface(name="vlan100", type="vlan"),
                FGInterface(name="agg1", type="aggregate", members=[]),
            ]
        )

        topology = build_derived_views(config).topology
        by_name = {item.name: item for item in topology.interfaces}

        for name in ("l2t.root", "naf.root", "ssl.root", "loop1"):
            self.assertEqual(by_name[name].path, (name,))
            self.assertIsNone(by_name[name].aggregate)
            self.assertEqual(by_name[name].physical_interfaces, ())
            self.assertEqual(by_name[name].issues, ())

        self.assertEqual(
            by_name["vlan100"].issues,
            ("Logical interface has no resolvable physical parent.",),
        )
        self.assertEqual(
            by_name["agg1"].issues,
            ("Aggregate/redundant interface has no configured members.",),
        )

    def test_real_path_populates_interface_aggregate_columns(self):
        extracted = extract_fortigate_config(
            parse_fortigate_config(_CONFIG),
            config=ExtractionConfig(),
        )
        derived = build_derived_views(extracted.config)
        validation = validate_config(extracted.config, derived=derived)

        by_name = {item.name: item for item in derived.topology.interfaces}
        self.assertEqual(by_name["port1"].aggregate, "agg1")
        self.assertEqual(by_name["port2"].aggregate, "agg1")
        self.assertEqual(by_name["agg1"].aggregate, "agg1")
        self.assertEqual(by_name["vlan100"].aggregate, "agg1")
        self.assertEqual(by_name["port1"].path, ("port1",))
        self.assertEqual(by_name["port2"].path, ("port2",))
        self.assertEqual(by_name["vlan100"].path, ("vlan100", "agg1"))
        self.assertEqual(by_name["agg1"].physical_interfaces, ("port1", "port2"))

        output = io.BytesIO()
        export_excel(
            extracted=extracted,
            derived=derived,
            validation=validation,
            output=output,
        )
        sheet = load_workbook(io.BytesIO(output.getvalue()), data_only=False)["Interfaces"]
        headers = [sheet.cell(3, column).value for column in range(1, sheet.max_column + 1)]
        rows = {
            row[headers.index("Name")]: row
            for row in sheet.iter_rows(min_row=4, values_only=True)
        }
        self.assertEqual(rows["├─ ● port1"][headers.index("Aggregate")], "agg1")
        self.assertEqual(rows["├─ ● port2"][headers.index("Aggregate")], "agg1")
        self.assertEqual(rows["◆ agg1"][headers.index("Members")], "port1\nport2")
        self.assertEqual(rows["◆ agg1"][headers.index("Physical Interfaces")], "port1\nport2")
        self.assertEqual(rows["└─ ▣ vlan100"][headers.index("Aggregate")], "agg1")

    def test_invalid_membership_is_scoped_and_reviewable(self):
        config = FGConfig(
            interfaces=[
                FGInterface(name="port1", vdom="a"),
                FGInterface(name="agg1", vdom="a", type="aggregate", members=[]),
                FGInterface(name="port1", vdom="b"),
                FGInterface(name="agg1", vdom="b", type="aggregate", members=["port1"]),
                FGInterface(name="agg2", vdom="b", type="aggregate", members=["port1", "vlan100"]),
                FGInterface(name="vlan100", vdom="b", type="vlan", interface="agg1"),
            ]
        )
        derived = build_derived_views(config)
        validation = validate_config(config, derived=derived)
        by_key = {(item.vdom, item.name): item for item in derived.topology.interfaces}

        self.assertIsNone(by_key[("a", "port1")].aggregate)
        self.assertIsNone(by_key[("b", "port1")].aggregate)
        self.assertEqual(by_key[("b", "vlan100")].aggregate, "agg1")
        self.assertTrue(by_key[("b", "vlan100")].issues)
        self.assertTrue(by_key[("a", "agg1")].issues)
        self.assertTrue(by_key[("b", "port1")].issues)
        self.assertTrue(
            any(
                issue.domain == "interface_topology"
                and issue.object_name == "agg1"
                for issue in validation.issues
            )
        )
        self.assertTrue(
            any(
                issue.domain == "interface_topology"
                and issue.object_name == "port1"
                for issue in validation.issues
            )
        )

        output = io.BytesIO()
        extracted = ExtractionResult(config=config)
        export_excel(
            extracted=extracted,
            derived=derived,
            validation=validation,
            output=output,
        )
        sheet = load_workbook(io.BytesIO(output.getvalue()), data_only=False)["Interfaces"]
        headers = [sheet.cell(3, column).value for column in range(1, sheet.max_column + 1)]
        rows = {
            (row[headers.index("VDOM")], row[headers.index("Name")]): row
            for row in sheet.iter_rows(min_row=4, values_only=True)
        }
        empty = rows[("a", "◆ agg1")]
        self.assertIsNone(empty[headers.index("Physical Interfaces")])
        self.assertEqual(empty[headers.index("Analysis Status")], "REVIEW_REQUIRED")
        self.assertTrue(empty[headers.index("Topology Issues")])


if __name__ == "__main__":
    unittest.main()
