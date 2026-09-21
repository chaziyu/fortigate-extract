from __future__ import annotations

import io
import unittest

from openpyxl import load_workbook

from fortigate_extract.config import ExtractionConfig
from fortigate_extract.derived import build_derived_views
from fortigate_extract.export import export_excel
from fortigate_extract.extraction.extractor import extract_fortigate_config
from fortigate_extract.parser import parse_fortigate_config
from fortigate_extract.validation.validator import validate_config


_SAMPLE_CONFIG = r"""
config system interface
    edit "port1"
        set ip 192.0.2.1 255.255.255.0
        set allowaccess ping https
    next
    edit "port2"
    next
    edit "agg1"
        set type aggregate
        set member "port1" "port2"
    next
    edit "vlan100"
        set type vlan
        set interface "agg1"
        set vlanid 100
        set ip 10.100.0.1 255.255.255.0
    next
end

config vpn ipsec phase1-interface
    edit "VPN-HQ"
        set interface "vlan100"
        set remote-gw 203.0.113.2
        set psksecret "do-not-export-this-secret"
    next
end

config user local
    edit "alice"
        set type password
        set passwd "another-do-not-export-secret"
    next
end
"""


class ExcelReportTest(unittest.TestCase):
    def _workbook(self):
        tree = parse_fortigate_config(_SAMPLE_CONFIG)
        extracted = extract_fortigate_config(
            tree,
            config=ExtractionConfig(),
        )
        derived = build_derived_views(
            extracted.config
        )
        validation = validate_config(
            extracted.config,
            derived=derived,
        )

        output = io.BytesIO()
        export_excel(
            extracted=extracted,
            derived=derived,
            validation=validation,
            output=output,
            source_name="sample.conf",
        )
        output.seek(0)
        return load_workbook(
            output,
            data_only=False,
        )

    def test_workbook_contract_and_topology(self):
        workbook = self._workbook()

        self.assertIn(
            "Summary",
            workbook.sheetnames,
        )
        self.assertIn(
            "Review Required",
            workbook.sheetnames,
        )
        self.assertIn(
            "Interfaces",
            workbook.sheetnames,
        )
        self.assertIn(
            "DHCP Exclude Ranges",
            workbook.sheetnames,
        )
        self.assertNotIn(
            "Extraction Evidence",
            workbook.sheetnames,
        )

        interfaces = workbook["Interfaces"]
        headers = [
            interfaces.cell(3, column).value
            for column in range(
                1,
                interfaces.max_column + 1,
            )
        ]

        self.assertNotIn(
            "Migration Status",
            headers,
        )
        self.assertIn(
            "Analysis Status",
            headers,
        )
        self.assertIn(
            "Relationship",
            headers,
        )
        self.assertIn(
            "Topology Kind",
            headers,
        )

        name_column = headers.index("Name") + 1
        names = [
            str(
                interfaces.cell(
                    row,
                    name_column,
                ).value
                or ""
            )
            for row in range(
                4,
                interfaces.max_row + 1,
            )
        ]

        aggregate_index = next(
            index
            for index, value in enumerate(names)
            if "◆ agg1" in value
        )
        port1_index = next(
            index
            for index, value in enumerate(names)
            if "port1" in value
        )
        port2_index = next(
            index
            for index, value in enumerate(names)
            if "port2" in value
        )
        vlan_index = next(
            index
            for index, value in enumerate(names)
            if "▣ vlan100" in value
        )
        vpn_index = next(
            index
            for index, value in enumerate(names)
            if "◈ VPN-HQ" in value
        )

        self.assertLess(
            aggregate_index,
            port1_index,
        )
        self.assertLess(
            aggregate_index,
            port2_index,
        )
        self.assertLess(
            port1_index,
            vlan_index,
        )
        self.assertLess(
            port2_index,
            vlan_index,
        )
        self.assertLess(
            vlan_index,
            vpn_index,
        )

    def test_secrets_are_not_exported(self):
        workbook = self._workbook()

        forbidden = {
            "do-not-export-this-secret",
            "another-do-not-export-secret",
        }

        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue

                    value = str(cell.value)

                    for secret in forbidden:
                        self.assertNotIn(
                            secret,
                            value,
                            msg=(
                                f"Secret leaked in "
                                f"{sheet.title}!{cell.coordinate}"
                            ),
                        )

    def test_legacy_ir_columns_are_removed(self):
        workbook = self._workbook()

        system_settings = workbook["System Settings"]
        system_headers = {
            system_settings.cell(3, column).value
            for column in range(
                1,
                system_settings.max_column + 1,
            )
        }

        legacy_management_headers = {
            "Management IPv4 Address",
            "Management Netmask",
            "Management Default Gateway",
            "Management Address Type",
            "Management IPv6 Address",
            "Management IPv6 Default Gateway",
            "Management IPv6 Enabled",
            "Management IPv6 Address Type",
            "Management IPv6 Gateway Type",
            "Explicit Management Services",
            "System Permitted IPs",
        }

        self.assertTrue(
            system_headers.isdisjoint(
                legacy_management_headers
            )
        )

        policies = workbook["Policies"]
        policy_headers = {
            policies.cell(3, column).value
            for column in range(
                1,
                policies.max_column + 1,
            )
        }

        self.assertTrue(
            {
                "Source Addresses",
                "Destination Addresses",
                "Services",
                "Action",
                "Schedule",
                "UTM Status",
                "Source Explicit Fields",
            }.issubset(policy_headers)
        )

        self.assertTrue(
            policy_headers.isdisjoint(
                {
                    "Source Address (Original)",
                    "Source Address (Normalized)",
                    "Destination Address (Original)",
                    "Destination Address (Normalized)",
                    "Service (Original)",
                    "Service (Normalized)",
                    "Action (Original)",
                    "Action (Normalized)",
                    "Schedule (Original)",
                    "Schedule (Normalized)",
                    "Effective UTM Status",
                }
            )
        )

        routes = workbook["Routes"]
        route_headers = {
            routes.cell(3, column).value
            for column in range(
                1,
                routes.max_column + 1,
            )
        }

        self.assertTrue(
            route_headers.isdisjoint(
                {
                    "Source Route ID",
                    "Destination Prefix (Normalized)",
                    "Source Destination",
                    "Destination Object / Group",
                    "Device",
                    "Next Hop",
                    "Administrative Distance",
                }
            )
        )

    def test_target_vendor_columns_are_removed(self):
        workbook = self._workbook()

        ip_pools = workbook["IP Pools"]
        headers = {
            ip_pools.cell(3, column).value
            for column in range(
                1,
                ip_pools.max_column + 1,
            )
        }

        self.assertNotIn(
            "Check Point Pool Object Type",
            headers,
        )

        security = workbook[
            "Security Profiles"
        ]
        security_headers = {
            security.cell(3, column).value
            for column in range(
                1,
                security.max_column + 1,
            )
        }

        self.assertNotIn(
            "Anti-Spyware",
            security_headers,
        )
        self.assertNotIn(
            "WildFire",
            security_headers,
        )


if __name__ == "__main__":
    unittest.main()
