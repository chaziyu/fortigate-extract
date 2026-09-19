"""Human-oriented Excel presentation layer for firewall inventory workbooks.

This module changes workbook presentation only. Canonical IR, extracted values,
and target-generation semantics remain unchanged. Main inventory sheets are
reordered into Core -> Review -> Advanced columns; Advanced columns are retained
and collapsed using Excel outline groups so users can expand them with +/-. 
"""

from __future__ import annotations

from copy import copy
from dataclasses import dataclass
import math
import time
from typing import Any, Sequence

from fwmigrate.report.fortigate_semantics_excel import FortiGateSemanticsExcelExporter
from fwmigrate.report.excel_vendor_visibility import VendorAwareIRExcelExporter
from fwmigrate.report.excel_audit import build_audit_classifier
from fwmigrate.report.excel_options import ExcelExportProfile


EXTRACTION_EVIDENCE_SHEET = "Extraction Evidence"
_PARENT_SHEET_ORDER = tuple(FortiGateSemanticsExcelExporter.SHEET_ORDER)
_READABLE_SHEET_ORDER = (
    "Summary",
    FortiGateSemanticsExcelExporter.REVIEW_SHEET,
    EXTRACTION_EVIDENCE_SHEET,
    *(
        sheet_name
        for sheet_name in _PARENT_SHEET_ORDER
        if sheet_name
        not in {"Summary", FortiGateSemanticsExcelExporter.REVIEW_SHEET}
    ),
)


@dataclass(frozen=True)
class _ReviewWorkbookAnalysis:
    review_rows: list[tuple[str, str, str, str, str, int]]
    evidence_rows: list[tuple[str, str, str, str, str, int]]


class ReadableFortiGateExcelExporter(FortiGateSemanticsExcelExporter):
    """Apply compact, expandable review views to the generated workbook."""

    EXTRACTION_EVIDENCE_SHEET = EXTRACTION_EVIDENCE_SHEET
    SHEET_ORDER = _READABLE_SHEET_ORDER
    ALWAYS_VISIBLE_SHEETS = (
        FortiGateSemanticsExcelExporter.ALWAYS_VISIBLE_SHEETS
        | frozenset({EXTRACTION_EVIDENCE_SHEET})
    )
    LARGE_SHEET_ROW_THRESHOLD = 1000

    _OBJECT_HEADERS = (
        "Name",
        "Rule Name",
        "Object Name",
        "Item",
        "ID",
        "Source ID",
        "Section",
        "Interface",
        "Profile Name",
    )
    _REASON_HEADERS = (
        "Review Reasons",
        "Review Reason",
        "Reason",
        "Message",
        "Notes",
        "Audit Note",
    )
    _REVIEW_STATUS_HEADERS = (
        "Extraction Status",
        "Migration Status",
        "Status",
        "Confidence",
        "Result",
    )
    _EVIDENCE_STATUS_HEADERS = (
        "Extraction Status",
        "Migration Status",
        "Status",
    )

    # Order is intentional. Existing columns not listed below are preserved after
    # these fields and collapsed as one Advanced outline group.
    COLUMN_GROUPS = {
        "Interfaces": {
            "core": (
                "Name",
                "Alias",
                "Zone",
                "IP / Prefix",
                "Interface Type",
                "Role",
                "Addressing Mode",
                "Management Access",
                "VLAN ID",
                "Enabled",
                "Source VDOM",
                "Description",
            ),
            "review": (
                "Migration Status",
                "Extraction Status",
                "Manual Review",
                "Review Reasons",
            ),
        },
        "Addresses": {
            "core": (
                "Name",
                "Type",
                "Value",
                "Address Family",
                "Associated Interface",
                "Allow Routing",
                "Tags",
                "Description",
            ),
            "review": (
                "Migration Status",
                "Manual Review",
                "Audit Note",
                "Parse Error",
            ),
        },
        "Address Groups": {
            "core": (
                "Name",
                "Members",
                "Dynamic",
                "Dynamic Filter",
                "Address Family",
                "Exclusion Enabled",
                "Exclude Members",
                "Description",
            ),
            "review": (
                "Migration Status",
                "Manual Review",
                "Audit Note",
            ),
        },
        "Services": {
            "core": (
                "Name",
                "Category",
                "Configured Protocol",
                "Effective Protocol",
                "Protocol / Destination Port",
                "Proxy",
                "Description",
            ),
            "review": (
                "Migration Status",
                "Manual Review",
                "Audit Note",
            ),
        },
        "Policies": {
            "core": (
                "Rule #",
                "Name",
                "Source Interface",
                "Source Address (Normalized)",
                "Destination Interface",
                "Destination Address (Normalized)",
                "Service (Normalized)",
                "Action (Normalized)",
                "Schedule (Normalized)",
                "NAT Enabled",
                "Effective UTM Status",
                "Disabled",
            ),
            "review": (
                "Extraction Status",
                "Manual Review",
                "Review Reasons",
            ),
        },
        "NAT Rules": {
            "core": (
                "Rule #",
                "Name",
                "Type",
                "Enabled",
                "Source Interface",
                "Destination Interface",
                "Original Source",
                "Original Destination",
                "Services",
                "Source Translation Mode",
                "Translated Source",
                "Destination Translation Mode",
                "Translated Destination",
                "Description",
            ),
            "review": (
                "Migration Status",
                "Manual Review",
                "Review Reasons",
            ),
        },
        "Routes": {
            "core": (
                "Name",
                "Destination Prefix (Normalized)",
                "Interface",
                "Next Hop",
                "Administrative Distance",
                "Priority",
                "Enabled",
                "SD-WAN Zone",
                "VRF",
                "Description",
            ),
            "review": (
                "Migration Status",
                "Manual Review",
                "Review Reasons",
                "Parse Error",
            ),
        },
        "VPN Tunnels": {
            "core": (
                "Name",
                "Type",
                "Peer Address",
                "Local Interface",
                "Remote Gateway IPv4",
                "Remote Gateway IPv6",
                "IKE Version",
                "Authentication Method",
                "IKE Proposal",
                "NAT Traversal",
                "DPD Mode",
                "Description",
            ),
            "review": (
                "Extraction Status",
                "Manual Review",
                "Review Reasons",
            ),
        },
        "SD-WAN Rules": {
            "core": (
                "ID",
                "Name",
                "Status",
                "Mode",
                "Source",
                "Destination",
                "Internet Service",
                "Priority Members",
                "Health Checks",
                "SLA Compare Method",
                "VDOM",
            ),
            "review": (),
        },
        "Administrators": {
            "core": (
                "Name",
                "Access Profile",
                "VDOMs",
                "IPv4 Trusted Hosts",
                "Two Factor",
                "Remote Auth",
                "Remote Group",
                "Credential Configured",
            ),
            "review": (
                "Migration Status",
                "Manual Review",
                "Unresolved References",
            ),
        },
        "User Groups": {
            "core": (
                "Name",
                "ID",
                "Type",
                "Members",
                "Resolved Members",
                "Unresolved Members",
                "Match Count",
            ),
            "review": (
                "Extraction Status",
                "Manual Review",
            ),
        },
    }

    STATUS_HEADERS = frozenset(
        {
            "Migration Status",
            "Extraction Status",
            "Status",
            "Confidence",
            "Result",
        }
    )
    REVIEW_REASON_HEADERS = frozenset(
        {
            "Review Reasons",
            "Review Reason",
            "Audit Note",
            "Parse Error",
            "Unresolved References",
        }
    )

    def _workbook_sheet_order(
        self,
        active_order: tuple[str, ...],
    ) -> tuple[str, ...]:
        return (
            "Summary",
            self.REVIEW_SHEET,
            self.EXTRACTION_EVIDENCE_SHEET,
            *(
                sheet_name
                for sheet_name in active_order
                if sheet_name != "Summary"
            ),
        )

    def _sheet_category(self, sheet_name: str) -> str:
        if sheet_name == self.EXTRACTION_EVIDENCE_SHEET:
            return "Overview"
        return super()._sheet_category(sheet_name)

    def _is_large_sheet(self, sheet: Any) -> bool:
        return self._record_count(sheet) > self.LARGE_SHEET_ROW_THRESHOLD

    @staticmethod
    def _normalized_status(value: Any) -> str:
        return str(value or "").strip().upper().replace(" ", "_")

    @classmethod
    def _friendly_status(cls, value: Any) -> str:
        normalized = cls._normalized_status(value)
        labels = {
            "PARTIALLY_NORMALIZED": "Partial",
            "PARTIAL": "Partial",
            "EXTRACT_ONLY": "Extract only",
            "UNSUPPORTED": "Unsupported",
            "UNRESOLVED": "Unresolved",
            "PARSE_ERROR": "Parse error",
            "MANUAL": "Manual review",
            "REVIEW": "Review",
            "NORMALIZED": "Normalized",
        }
        return labels.get(normalized, str(value or "").replace("_", " ").title())

    def _review_rows(self, workbook: Any) -> list[tuple[str, str, str, str, str, int]]:
        return self._analyze_review_workbook(workbook).review_rows

    def _extraction_evidence_rows(
        self,
        workbook: Any,
    ) -> list[tuple[str, str, str, str, str, int]]:
        return self._analyze_review_workbook(workbook).evidence_rows

    def _analyze_review_workbook(self, workbook: Any) -> _ReviewWorkbookAnalysis:
        accumulator = getattr(self, "_audit_accumulator", None)
        if accumulator is not None:
            review_rows, evidence_rows = accumulator.for_sheets(
                sheet.title for sheet in workbook.worksheets
            )
            return _ReviewWorkbookAnalysis(review_rows, evidence_rows)

        review_rows: list[tuple[str, str, str, str, str, int]] = []
        evidence_rows: list[tuple[str, str, str, str, str, int]] = []
        excluded_sheets = {
            "Summary",
            self.REVIEW_SHEET,
            self.EXTRACTION_EVIDENCE_SHEET,
        }

        for sheet in workbook.worksheets:
            if sheet.title in excluded_sheets or sheet.max_row < 4:
                continue

            headers = tuple(
                sheet.cell(3, column).value
                for column in range(1, sheet.max_column + 1)
            )
            classifier = build_audit_classifier(
                sheet.title,
                headers,
                self._sheet_category,
            )
            for row_number, values in enumerate(
                sheet.iter_rows(min_row=4, values_only=True),
                4,
            ):
                classified = classifier.classify(values, row_number)
                if classified.review is not None:
                    review_rows.append(classified.review)
                if classified.evidence is not None:
                    evidence_rows.append(classified.evidence)

        return _ReviewWorkbookAnalysis(review_rows, evidence_rows)

    def _build_review_required(self, workbook: Any) -> None:
        from openpyxl.styles import PatternFill

        if self.REVIEW_SHEET in workbook.sheetnames:
            workbook.remove(workbook[self.REVIEW_SHEET])
        if self.EXTRACTION_EVIDENCE_SHEET in workbook.sheetnames:
            workbook.remove(workbook[self.EXTRACTION_EVIDENCE_SHEET])

        analysis_started = time.perf_counter()
        analysis = self._analyze_review_workbook(workbook)
        metrics = getattr(self, "_last_export_metrics", None)
        if metrics is not None:
            metrics.timings["worksheet analysis"] = time.perf_counter() - analysis_started
        review_rows = []
        for category, obj, issue, status, source_sheet, source_row in analysis.review_rows:
            normalized = self._normalized_status(status)
            if normalized in {"UNSUPPORTED", "PARSE_ERROR"}:
                severity = "Critical"
            elif normalized == "UNRESOLVED":
                severity = "High"
            else:
                severity = "Review"
            review_rows.append(
                (
                    severity,
                    category,
                    obj,
                    issue,
                    self._friendly_status(status),
                    source_sheet,
                    source_row,
                )
            )

        review_started = time.perf_counter()
        review_sheet = self._table_sheet(
            workbook,
            self.REVIEW_SHEET,
            (
                "Severity",
                "Category",
                "Object",
                "Issue / Review Reason",
                "Status",
                "Source Sheet",
                "Source Row",
            ),
            review_rows,
            empty_note="No actionable items currently require manual review.",
            subtitle=(
                "Actionable unsupported, unresolved, partial, manual-review, and "
                "parse-error findings. Extract-only evidence is listed separately."
            ),
        )
        if metrics is not None:
            metrics.timings["Review Required generation"] = (
                time.perf_counter() - review_started
            )

        evidence_rows = analysis.evidence_rows
        evidence_started = time.perf_counter()
        evidence_sheet = self._table_sheet(
            workbook,
            self.EXTRACTION_EVIDENCE_SHEET,
            (
                "Category",
                "Object",
                "Evidence / Note",
                "Status",
                "Source Sheet",
                "Source Row",
            ),
            evidence_rows,
            empty_note="No extract-only source evidence was recorded.",
            subtitle=(
                "Source-only configuration retained for audit and migration reference; "
                "these rows are not treated as actionable review findings."
            ),
        )
        if metrics is not None:
            metrics.timings["Extraction Evidence generation"] = (
                time.perf_counter() - evidence_started
            )

        fills = {
            "red": PatternFill("solid", fgColor=self._LIGHT_RED),
            "amber": PatternFill("solid", fgColor=self._LIGHT_AMBER),
            "blue": PatternFill("solid", fgColor=self._LIGHT_BLUE),
        }
        for sheet, source_col, row_col, status_col in (
            (review_sheet, 6, 7, 5),
            (evidence_sheet, 5, 6, 4),
        ):
            for row_number in range(4, sheet.max_row + 1):
                source_sheet = str(sheet.cell(row_number, source_col).value or "")
                source_row = sheet.cell(row_number, row_col).value
                if source_sheet and source_sheet in workbook.sheetnames and source_row:
                    escaped = source_sheet.replace("'", "''")
                    sheet.cell(row_number, source_col).hyperlink = (
                        f"#'{escaped}'!A{source_row}"
                    )
                    sheet.cell(row_number, source_col).style = "Hyperlink"

                status_cell = sheet.cell(row_number, status_col)
                normalized = self._normalized_status(status_cell.value)
                if normalized in {"UNSUPPORTED", "PARSE_ERROR", "UNRESOLVED"}:
                    status_cell.fill = fills["red"]
                elif normalized in {"PARTIAL", "PARTIALLY_NORMALIZED", "MANUAL_REVIEW"}:
                    status_cell.fill = fills["amber"]
                elif normalized == "EXTRACT_ONLY":
                    status_cell.fill = fills["blue"]

        for row_number in range(4, review_sheet.max_row + 1):
            severity_cell = review_sheet.cell(row_number, 1)
            if severity_cell.value in {"Critical", "High"}:
                severity_cell.fill = fills["red"]
            elif severity_cell.value == "Review":
                severity_cell.fill = fills["amber"]

    def _apply_sheet_view(self, sheet: Any) -> None:
        # Bypass the previous FortiGate hide-only layer. This class replaces it
        # with ordered Core/Review columns and a collapsible Advanced group.
        if sheet.title == "Summary":
            VendorAwareIRExcelExporter._apply_sheet_view(self, sheet)
            self._collapse_hidden_navigation_rows(sheet)
            return

        groups = self.COLUMN_GROUPS.get(sheet.title)
        options = getattr(self, "_export_options", None)
        large = self._is_large_sheet(sheet) or getattr(
            options,
            "profile",
            None,
        ) in {ExcelExportProfile.FAST, ExcelExportProfile.DATA_ONLY}
        if groups and sheet.max_row >= 3:
            if getattr(self, "_preserve_column_order", False):
                self._group_columns_in_place(sheet, groups)
            elif large:
                self._group_columns_in_place(sheet, groups)
            else:
                self._reorder_and_group_columns(sheet, groups)
        elif (
            sheet.title not in self.ALWAYS_VISIBLE_SHEETS
            and sheet.max_row >= 4
            and self._record_count(sheet) > 0
        ):
            self._hide_empty_columns(sheet)

        VendorAwareIRExcelExporter._apply_sheet_view(self, sheet)
        self._apply_status_formatting(sheet)
        if not large:
            self._fit_visible_row_heights(sheet)

    @staticmethod
    def _column_order(
        headers: Sequence[str],
        groups: dict[str, tuple[str, ...]],
    ) -> tuple[list[int], list[int], list[int], list[int]]:
        positions: dict[str, list[int]] = {}
        for index, header in enumerate(headers, 1):
            name = str(header or "").strip()
            if name:
                positions.setdefault(name, []).append(index)

        used: set[int] = set()

        def select(names: Sequence[str]) -> list[int]:
            selected = []
            for name in names:
                for index in positions.get(name, ()):
                    if index not in used:
                        selected.append(index)
                        used.add(index)
                        break
            return selected

        core_columns = select(groups.get("core", ()))
        review_columns = select(groups.get("review", ()))
        advanced_columns = [
            index
            for index, header in enumerate(headers, 1)
            if index not in used and str(header or "").strip()
        ]
        empty_columns = [
            index
            for index, header in enumerate(headers, 1)
            if index not in used and not str(header or "").strip()
        ]
        return (
            core_columns,
            review_columns,
            advanced_columns,
            core_columns + review_columns + advanced_columns + empty_columns,
        )

    def _classified_columns(
        self,
        sheet: Any,
        groups: dict[str, tuple[str, ...]],
    ) -> tuple[list[str], list[int], list[int], list[int], list[int]]:
        headers = [
            str(sheet.cell(3, column).value or "").strip()
            for column in range(1, sheet.max_column + 1)
        ]
        core, review, advanced, ordered = self._column_order(headers, groups)
        return headers, core, review, advanced, ordered

    def _group_columns_in_place(
        self,
        sheet: Any,
        groups: dict[str, tuple[str, ...]],
    ) -> None:
        from openpyxl.utils import get_column_letter

        _, core_columns, review_columns, advanced_columns, _ = self._classified_columns(
            sheet, groups
        )
        for column in core_columns + review_columns:
            dimension = sheet.column_dimensions[get_column_letter(column)]
            dimension.hidden = False
            dimension.outlineLevel = 0
        for column in advanced_columns:
            dimension = sheet.column_dimensions[get_column_letter(column)]
            dimension.hidden = True
            dimension.outlineLevel = 1
        if advanced_columns:
            sheet.sheet_properties.outlinePr.summaryRight = True

    def _reorder_and_group_columns(
        self,
        sheet: Any,
        groups: dict[str, tuple[str, ...]],
    ) -> None:
        from openpyxl.utils import get_column_letter

        headers, core_columns, review_columns, _, ordered_columns = self._classified_columns(
            sheet, groups
        )
        if ordered_columns == list(range(1, sheet.max_column + 1)):
            self._group_columns_in_place(sheet, groups)
            return

        # Snapshot row 3 onward because rows 1-2 are title/subtitle merged ranges.
        snapshots = []
        dimensions = []
        for source_column in ordered_columns:
            source_letter = get_column_letter(source_column)
            source_dimension = sheet.column_dimensions[source_letter]
            dimensions.append(
                (
                    source_dimension.width,
                    source_dimension.bestFit,
                )
            )
            column_cells = []
            for row in range(3, sheet.max_row + 1):
                source = sheet.cell(row, source_column)
                column_cells.append(
                    (
                        source.value,
                        copy(source._style),
                        copy(source.hyperlink),
                        copy(source.comment),
                    )
                )
            snapshots.append(column_cells)

        for destination_column, (column_cells, dimension) in enumerate(
            zip(snapshots, dimensions),
            1,
        ):
            destination_letter = get_column_letter(destination_column)
            target_dimension = sheet.column_dimensions[destination_letter]
            target_dimension.width = dimension[0]
            target_dimension.bestFit = dimension[1]
            target_dimension.hidden = False
            target_dimension.outlineLevel = 0

            for offset, snapshot in enumerate(column_cells, 3):
                target = sheet.cell(offset, destination_column)
                value, style, hyperlink, comment = snapshot
                target.value = value
                target._style = style
                target.hyperlink = hyperlink
                target.comment = comment

        self._group_columns_in_place(sheet, groups)

    def _hide_empty_columns(self, sheet: Any) -> None:
        from openpyxl.utils import get_column_letter

        headers = next(
            sheet.iter_rows(
                min_row=3,
                max_row=3,
                max_col=sheet.max_column,
                values_only=True,
            ),
            (),
        )
        candidates = {
            column
            for column, header in enumerate(headers, 1)
            if header
        }
        populated = set()
        remaining = set(candidates)
        for values in sheet.iter_rows(
            min_row=4,
            max_row=sheet.max_row,
            max_col=sheet.max_column,
            values_only=True,
        ):
            for column in tuple(remaining):
                if values[column - 1] not in (None, ""):
                    populated.add(column)
                    remaining.remove(column)
            if not remaining:
                break

        for column in candidates:
            sheet.column_dimensions[get_column_letter(column)].hidden = (
                column not in populated
            )

    def _apply_status_formatting(self, sheet: Any) -> None:
        from openpyxl.styles import PatternFill

        if sheet.max_row < 4:
            return
        headers = self._header_map(sheet)
        status_columns = [
            headers[header] for header in self.STATUS_HEADERS if header in headers
        ]
        manual_column = headers.get("Manual Review")
        reason_columns = [
            headers[header]
            for header in self.REVIEW_REASON_HEADERS
            if header in headers
        ]
        fills = {
            "red": PatternFill("solid", fgColor=self._LIGHT_RED),
            "amber": PatternFill("solid", fgColor=self._LIGHT_AMBER),
            "blue": PatternFill("solid", fgColor=self._LIGHT_BLUE),
            "teal": PatternFill("solid", fgColor=self._LIGHT_TEAL),
        }
        status_fills = {
            **dict.fromkeys(
                ("UNSUPPORTED", "PARSE_ERROR", "UNRESOLVED"), fills["red"]
            ),
            **dict.fromkeys(
                ("PARTIALLY_NORMALIZED", "PARTIAL", "MANUAL"), fills["amber"]
            ),
            "EXTRACT_ONLY": fills["blue"],
            **dict.fromkeys(
                ("NORMALIZED", "COMPLETE", "SUCCESS"), fills["teal"]
            ),
        }

        for row in range(4, sheet.max_row + 1):
            for column in status_columns:
                cell = sheet.cell(row, column)
                fill = status_fills.get(self._normalized_status(cell.value))
                if fill is not None:
                    cell.fill = fill
            if manual_column is not None:
                cell = sheet.cell(row, manual_column)
                if self._truthy_review(cell.value):
                    cell.fill = fills["amber"]
            for column in reason_columns:
                cell = sheet.cell(row, column)
                if cell.value not in (None, ""):
                    cell.fill = fills["amber"]

    def _fit_visible_row_heights(self, sheet: Any) -> None:
        from openpyxl.utils import get_column_letter

        if sheet.max_row < 4:
            return
        visible_columns = [
            column
            for column in range(1, sheet.max_column + 1)
            if not sheet.column_dimensions[get_column_letter(column)].hidden
        ]
        if not visible_columns:
            return

        for row in range(4, sheet.max_row + 1):
            max_lines = 1
            for column in visible_columns:
                value = sheet.cell(row, column).value
                if value in (None, ""):
                    continue
                text = str(value)
                width = sheet.column_dimensions[get_column_letter(column)].width or 14
                chars_per_line = max(int(width), 8)
                lines = 0
                for segment in text.splitlines() or [text]:
                    lines += max(1, math.ceil(len(segment) / chars_per_line))
                max_lines = max(max_lines, lines)
            sheet.row_dimensions[row].height = min(60, 20 + 15 * (max_lines - 1))

    def _collapse_hidden_navigation_rows(self, summary: Any) -> None:
        header_row = next(
            (
                row
                for row in range(1, summary.max_row + 1)
                if summary.cell(row, 1).value == "Category"
                and summary.cell(row, 2).value == "Sheet"
            ),
            None,
        )
        if header_row is None:
            return

        visibility_column = next(
            (
                column
                for column in range(1, summary.max_column + 1)
                if summary.cell(header_row, column).value == "Visibility"
            ),
            None,
        )
        if visibility_column is None:
            return

        hidden_rows = []
        for row in range(header_row + 1, summary.max_row + 1):
            if not summary.cell(row, 2).value:
                break
            if summary.cell(row, visibility_column).value == "Hidden":
                hidden_rows.append(row)

        if not hidden_rows:
            return

        start = previous = hidden_rows[0]
        for row in hidden_rows[1:] + [None]:
            if row is not None and row == previous + 1:
                previous = row
                continue
            summary.row_dimensions.group(
                start,
                previous,
                outline_level=1,
                hidden=True,
            )
            if row is not None:
                start = previous = row
        summary.sheet_properties.outlinePr.summaryBelow = True

    def _add_summary_visibility(self, summary: Any, workbook: Any) -> None:
        from openpyxl.styles import Alignment, Font, PatternFill

        super()._add_summary_visibility(summary, workbook)

        # Use the otherwise-empty right side of the metadata area for a compact
        # dashboard without shifting navigation rows or breaking hyperlinks.
        summary.merge_cells("D4:E4")
        title = summary["D4"]
        title.value = "Key Counts"
        title.font = Font(name="Aptos", bold=True, color=self._WHITE)
        title.fill = PatternFill("solid", fgColor=self._NAVY)
        title.alignment = Alignment(vertical="center")

        metrics = (
            ("Interfaces", self._sheet_records(workbook, "Interfaces")),
            ("Addresses", self._sheet_records(workbook, "Addresses")),
            ("Policies", self._sheet_records(workbook, "Policies")),
            ("NAT Rules", self._sheet_records(workbook, "NAT Rules")),
            ("Routes", self._sheet_records(workbook, "Routes")),
            ("Needs Review", self._sheet_records(workbook, self.REVIEW_SHEET)),
            (
                "Extract-only Evidence",
                self._sheet_records(workbook, self.EXTRACTION_EVIDENCE_SHEET),
            ),
        )
        for row, (label, count) in enumerate(metrics, 5):
            summary.cell(row, 4, label)
            summary.cell(row, 5, count)
            summary.cell(row, 4).font = Font(name="Aptos", size=10, color=self._TEXT)
            summary.cell(row, 5).font = Font(
                name="Aptos",
                size=10,
                bold=True,
                color=self._TEXT,
            )
            summary.cell(row, 5).alignment = Alignment(horizontal="right")

        summary.column_dimensions["D"].width = max(summary.column_dimensions["D"].width or 0, 22)
        summary.column_dimensions["E"].width = max(summary.column_dimensions["E"].width or 0, 14)

    def _sheet_records(self, workbook: Any, sheet_name: str) -> int:
        if sheet_name not in workbook.sheetnames:
            return 0
        return self._record_count(workbook[sheet_name])
