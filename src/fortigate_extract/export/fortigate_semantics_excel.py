"""Excel projection for FortiGate route devindex and NGFW pre-match rules."""

from __future__ import annotations

from copy import copy
import json
from typing import Any

from fwmigrate.report.excel_effective_order import EffectiveOrderIRExcelExporter


PREMATCH_SHEET = "NGFW Pre-Match Policies"
_PREMATCH_FAMILY = "ngfw-pre-match-policy"


def _join(value: Any) -> Any:
    if isinstance(value, (list, tuple, set)):
        return "\n".join(str(item) for item in value)
    return value


class FortiGateSemanticsExcelExporter(EffectiveOrderIRExcelExporter):
    """Expose corrected FortiGate source semantics in dedicated inventory views."""

    # Keep the full source inventory in the workbook, but make the most common
    # review sheets open with a compact set of migration-relevant columns.
    # Non-core columns remain available to users by unhiding them in Excel.
    CORE_COLUMNS = {
        "Interfaces": frozenset(
            {
                "Name",
                "Source VDOM",
                "Zone",
                "IP / Prefix",
                "Enabled",
                "Interface Type",
                "Role",
                "Addressing Mode",
                "Management Access",
                "Alias",
                "VLAN ID",
                "Description",
            }
        ),
        "Addresses": frozenset(
            {
                "Name",
                "Type",
                "Value",
                "Address Family",
                "Associated Interface",
                "Allow Routing",
                "Tags",
                "Description",
            }
        ),
        "Address Groups": frozenset(
            {
                "Name",
                "Members",
                "Dynamic",
                "Dynamic Filter",
                "Address Family",
                "Exclusion Enabled",
                "Exclude Members",
                "Description",
            }
        ),
        "Services": frozenset(
            {
                "Name",
                "Category",
                "Configured Protocol",
                "Effective Protocol",
                "Protocol / Destination Port",
                "Proxy",
                "Description",
            }
        ),
        "Policies": frozenset(
            {
                "Rule #",
                "Source Policy ID",
                "Name",
                "Source Interface",
                "From Zone",
                "Destination Interface",
                "To Zone",
                "Source Address (Normalized)",
                "Destination Address (Normalized)",
                "User Groups",
                "Service (Normalized)",
                "Action (Normalized)",
                "Schedule (Normalized)",
                "Disabled",
                "NAT Enabled",
                "Effective UTM Status",
                "Security Profile Group",
                "Description",
            }
        ),
        "NAT Rules": frozenset(
            {
                "Rule #",
                "Name",
                "Type",
                "Source Policy ID",
                "Enabled",
                "Source Interface",
                "From Zone",
                "Destination Interface",
                "To Zone",
                "Original Source",
                "Original Destination",
                "Services",
                "Source Translation Mode",
                "Translated Source",
                "Destination Translation Mode",
                "Translated Destination",
                "Description",
            }
        ),
        "Routes": frozenset(
            {
                "Name",
                "Source Route ID",
                "Address Family",
                "Destination Prefix (Normalized)",
                "Interface",
                "Next Hop",
                "Administrative Distance",
                "Priority",
                "Enabled",
                "Description",
                "Device Index",
            }
        ),
        "SD-WAN Rules": frozenset(
            {
                "ID",
                "Name",
                "Mode",
                "Status",
                "Source",
                "Destination",
                "Health Checks",
                "Priority Members",
                "Priority Zones",
                "Internet Service",
                "Internet Service Names",
                "SLA Compare Method",
                "Tie Break",
                "VDOM",
            }
        ),
        "Administrators": frozenset(
            {
                "Name",
                "Access Profile",
                "VDOMs",
                "IPv4 Trusted Hosts",
                "IPv6 Trusted Hosts",
                "Two Factor",
                "FortiToken",
                "Remote Auth",
                "Remote Group",
                "Credential Configured",
            }
        ),
        "User Groups": frozenset(
            {
                "Name",
                "ID",
                "Type",
                "Members",
                "Resolved Members",
                "Unresolved Members",
                "Match Count",
            }
        ),
    }

    # These fields explain migration risk and must remain visible even when the
    # sheet's normal core view would otherwise hide them.
    REVIEW_COLUMNS = frozenset(
        {
            "Migration Status",
            "Extraction Status",
            "Manual Review",
            "Review Reasons",
            "Review Reason",
            "Audit Note",
            "Parse Error",
        }
    )

    def _active_sheet_order(self) -> tuple[str, ...]:
        order = super()._active_sheet_order()
        if self._source_vendor() == "fortigate":
            return order
        return tuple(sheet_name for sheet_name in order if sheet_name != PREMATCH_SHEET)

    def _apply_sheet_view(self, sheet: Any) -> None:
        """Apply the normal view plus compact, data-aware column visibility."""
        super()._apply_sheet_view(sheet)

        if sheet.title in self.ALWAYS_VISIBLE_SHEETS:
            return
        if sheet.max_row < 4:
            return

        from openpyxl.utils import get_column_letter

        core_columns = self.CORE_COLUMNS.get(sheet.title)
        record_count = self._record_count(sheet)

        for column in range(1, sheet.max_column + 1):
            header = str(sheet.cell(3, column).value or "").strip()
            if not header:
                continue

            if core_columns is not None:
                # Core views are intentionally stable across different source
                # configurations. This avoids the visible layout changing just
                # because one firewall happens to populate an uncommon field.
                hidden = (
                    header not in core_columns
                    and header not in self.REVIEW_COLUMNS
                )
            elif record_count > 0:
                # For all other inventory sheets, retain every populated field
                # and only suppress columns that contain no data at all.
                hidden = not any(
                    sheet.cell(row, column).value not in (None, "")
                    for row in range(4, sheet.max_row + 1)
                )
            else:
                hidden = False

            sheet.column_dimensions[get_column_letter(column)].hidden = hidden

    def _build_routes(self, workbook: Any) -> None:
        super()._build_routes(workbook)
        if "Routes" not in workbook.sheetnames:
            return

        sheet = workbook["Routes"]
        headers = self._header_map(sheet)
        if "Device Index" in headers:
            return

        column = sheet.max_column + 1
        header = sheet.cell(3, column, "Device Index")
        if column > 1:
            template = sheet.cell(3, column - 1)
            header.font = copy(template.font)
            header.fill = copy(template.fill)
            header.border = copy(template.border)
            header.alignment = copy(template.alignment)
            header.protection = copy(template.protection)
            header.number_format = template.number_format

        context_column = next(
            (
                headers[name]
                for name in ("Source VDOM", "Source Context", "VDOM")
                if name in headers
            ),
            None,
        )
        id_column = next(
            (
                headers[name]
                for name in ("Source Route ID", "Route ID", "Source ID", "ID")
                if name in headers
            ),
            None,
        )

        route_lookup: dict[tuple[str, str], Any] = {}
        routes_by_id: dict[str, list[Any]] = {}
        for route in self.ir.routes:
            source_id = getattr(route, "source_route_id", None)
            if source_id is None:
                continue
            source_id_text = str(source_id)
            source_context = str(getattr(route, "source_context", "root") or "root")
            route_lookup[(source_context, source_id_text)] = route
            routes_by_id.setdefault(source_id_text, []).append(route)

        for row_number in range(4, sheet.max_row + 1):
            route = None
            source_id_text = (
                str(sheet.cell(row_number, id_column).value)
                if id_column and sheet.cell(row_number, id_column).value is not None
                else ""
            )
            if source_id_text:
                if context_column:
                    source_context = str(
                        sheet.cell(row_number, context_column).value or "root"
                    )
                    route = route_lookup.get((source_context, source_id_text))
                if route is None:
                    candidates = routes_by_id.get(source_id_text, [])
                    if len(candidates) == 1:
                        route = candidates[0]

            if route is None:
                route_index = row_number - 4
                if 0 <= route_index < len(self.ir.routes):
                    route = self.ir.routes[route_index]

            devindex = None
            if route is not None:
                devindex = dict(getattr(route, "source_attributes", {}) or {}).get(
                    "devindex"
                )
            cell = sheet.cell(row_number, column, devindex)
            if column > 1:
                template = sheet.cell(row_number, column - 1)
                cell.font = copy(template.font)
                cell.fill = copy(template.fill)
                cell.border = copy(template.border)
                cell.alignment = copy(template.alignment)
                cell.protection = copy(template.protection)
                cell.number_format = template.number_format

        sheet.column_dimensions[header.column_letter].width = 14

    def _build_security_policies(self, workbook: Any) -> None:
        super()._build_security_policies(workbook)
        rules = [
            rule
            for rule in self.ir.vendor_extensions.fortios.security_policies
            if getattr(rule, "family", None) == _PREMATCH_FAMILY
        ]

        displayed_keys = {
            "srcintf",
            "dstintf",
            "srcaddr",
            "dstaddr",
            "srcaddr6",
            "dstaddr6",
            "service",
            "users",
            "groups",
            "fsso_groups",
            "ssl_ssh_profile",
            "comments",
            "source_tree",
        }
        rows = []
        for rule in sorted(
            rules,
            key=lambda item: (
                str(getattr(item, "source_context", "") or ""),
                int(getattr(item, "source_order", 0) or 0),
            ),
        ):
            attributes = dict(getattr(rule, "source_attributes", {}) or {})
            additional = {
                key: value
                for key, value in attributes.items()
                if key not in displayed_keys
            }
            rows.append(
                (
                    getattr(rule, "source_context", None),
                    getattr(rule, "source_id", None),
                    getattr(rule, "source_order", None),
                    getattr(rule, "name", None),
                    "Yes" if getattr(rule, "enabled", None) is True else (
                        "No" if getattr(rule, "enabled", None) is False else None
                    ),
                    _join(attributes.get("srcintf")),
                    _join(attributes.get("dstintf")),
                    _join(attributes.get("srcaddr")),
                    _join(attributes.get("dstaddr")),
                    _join(attributes.get("srcaddr6")),
                    _join(attributes.get("dstaddr6")),
                    _join(attributes.get("users")),
                    _join(attributes.get("groups")),
                    _join(attributes.get("fsso_groups")),
                    _join(attributes.get("service")),
                    attributes.get("ssl_ssh_profile"),
                    attributes.get("comments"),
                    getattr(rule, "migration_status", None),
                    "Yes" if getattr(rule, "requires_manual_review", False) else "No",
                    "\n".join(getattr(rule, "review_reasons", []) or []),
                    json.dumps(additional, sort_keys=True, default=str)
                    if additional
                    else None,
                )
            )

        self._table_sheet(
            workbook,
            PREMATCH_SHEET,
            (
                "Source VDOM",
                "Rule ID",
                "Source Order",
                "Name",
                "Enabled",
                "Source Interfaces",
                "Destination Interfaces",
                "Source Addresses",
                "Destination Addresses",
                "IPv6 Source Addresses",
                "IPv6 Destination Addresses",
                "Users",
                "Groups",
                "FSSO Groups",
                "Services",
                "SSL Inspection Reference",
                "Comments",
                "Migration Status",
                "Manual Review",
                "Review Reasons",
                "Additional Settings",
            ),
            rows,
            empty_note="No FortiGate policy-based NGFW pre-match rules were extracted.",
            subtitle=(
                "FortiGate policy-based NGFW SSL inspection/authentication pre-match "
                "inventory. These rules have no portable allow/deny action and remain "
                "separate from NGFW Security Policies."
            ),
        )
