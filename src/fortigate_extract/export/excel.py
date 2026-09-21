from __future__ import annotations

from collections import defaultdict
from dataclasses import asdict, is_dataclass
from datetime import datetime, timezone
from enum import Enum
from pathlib import Path
from typing import Any, BinaryIO, Iterable, Mapping, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..config import ExtractionConfig
from ..derived import DerivedViews, build_derived_views
from ..extraction.result import ExtractionResult
from ..extraction.source_inventory import SourceObjectRecord
from ..security.extraction import sanitize_source_attributes
from ..section_registry import registered_sections
from ..validation.models import ValidationIssue, ValidationResult
from .excel_schema import (
    DERIVED_COLUMNS_BY_SHEET,
    HIDDEN_COLUMNS_BY_DEFAULT,
    SHEET_HEADERS,
    SHEET_ORDER,
)


_TITLE_FILL = PatternFill("solid", fgColor="17324D")
_HEADER_FILL = PatternFill("solid", fgColor="0F766E")
_DERIVED_FILL = PatternFill("solid", fgColor="D7F0EC")
_REVIEW_FILL = PatternFill("solid", fgColor="FEF3C7")
_ERROR_FILL = PatternFill("solid", fgColor="FEE2E2")
_ALT_FILL = PatternFill("solid", fgColor="F8FAFC")
_WHITE_FONT = Font(color="FFFFFF", bold=True)
_TITLE_FONT = Font(color="FFFFFF", bold=True, size=14)
_HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
_MUTED_FONT = Font(color="667085", italic=True, size=9)
_LINK_FONT = Font(color="0563C1", underline="single")
_THIN = Side(style="thin", color="D9E2DF")
_BORDER = Border(bottom=_THIN)


_SOURCE_PATHS_BY_SHEET: dict[str, tuple[str, ...]] = {
    "System Settings": ("system global", "system settings"),
    "DNS Settings": ("system dns",),
    "NTP Settings": ("system ntp", "system ntp ntpserver"),
}

_SOURCE_HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "Alias": ("alias",),
    "IP / Prefix": ("ip",),
    "Interface Type": ("type",),
    "Role": ("role",),
    "Addressing Mode": ("mode",),
    "Management Access": ("allowaccess",),
    "VLAN ID": ("vlanid",),
    "Members": ("member", "members"),
    "Description": ("description", "comment", "comments"),
    "Protocol": ("protocol",),
    "Source Port": ("src-port", "source-port"),
    "Destination Port": ("dst-port", "destination-port"),
    "Auto Negotiate": ("auto-negotiate", "auto_negotiate"),
}


def export_excel(
    *,
    extracted: ExtractionResult,
    validation: ValidationResult,
    output: BinaryIO | Path | str,
    config: ExtractionConfig | None = None,
    derived: DerivedViews | None = None,
    source_name: str | None = None,
) -> None:
    """Write the FortiGate configuration report."""

    del config

    derived = derived or build_derived_views(extracted.config)

    context = _ExcelContext(
        extracted=extracted,
        derived=derived,
        validation=validation,
        source_name=source_name,
    )

    workbook = _build_workbook(context)

    if isinstance(output, (str, Path)):
        path = Path(output)
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(path)
        return

    workbook.save(output)


class _ExcelContext:
    def __init__(
        self,
        *,
        extracted: ExtractionResult,
        derived: DerivedViews,
        validation: ValidationResult,
        source_name: str | None,
    ) -> None:
        self.extracted = extracted
        self.config = extracted.config
        self.derived = derived
        self.validation = validation
        self.source_name = source_name or ""
        self.source_by_path: dict[str, list[SourceObjectRecord]] = defaultdict(list)

        for record in extracted.source_objects:
            self.source_by_path[record.source_path].append(record)

        self.issues_by_object: dict[tuple[str, str, str], list[ValidationIssue]] = defaultdict(list)

        for issue in validation.issues:
            if issue.object_name:
                self.issues_by_object[(issue.domain, issue.vdom, str(issue.object_name))].append(issue)

    def issues_for(
        self,
        *,
        vdom: str,
        names: Iterable[Any],
        domains: Iterable[str] | None = None,
    ) -> list[ValidationIssue]:
        found: list[ValidationIssue] = []
        seen: set[tuple[str, str, str]] = set()
        allowed_domains = set(domains) if domains is not None else None

        for name in names:
            if name in (None, ""):
                continue

            issue_domains = (
                allowed_domains
                if allowed_domains is not None
                else {
                    domain
                    for domain, issue_vdom, issue_name in self.issues_by_object
                    if issue_vdom == vdom and issue_name == str(name)
                }
            )
            for domain in issue_domains:
                for issue in self.issues_by_object.get((domain, vdom, str(name)), ()):
                    key = (issue.domain, issue.field or "", issue.message)
                    if key in seen:
                        continue
                    seen.add(key)
                    found.append(issue)

        return found


def _build_workbook(context: _ExcelContext) -> Workbook:
    workbook = Workbook()
    workbook.remove(workbook.active)

    order = list(SHEET_ORDER)

    for sheet_name in order:
        if sheet_name == "Summary":
            _build_summary(
                workbook,
                context,
                order,
            )
            continue

        headers = list(
            SHEET_HEADERS[sheet_name]
        )

        rows = _rows_for_sheet(
            sheet_name,
            context,
            headers,
        )

        _write_table_sheet(
            workbook,
            sheet_name,
            headers,
            rows,
            context=context,
        )

    return workbook


def _rows_for_sheet(
    sheet_name: str,
    context: _ExcelContext,
    headers: Sequence[str],
) -> list[dict[str, Any]]:
    builders = {
        "Review Required": _review_rows,
        "Interfaces": _interface_rows,
        "Interface Secondary IPs": _interface_secondary_rows,
        "Zones": _zone_rows,
        "Addresses": _address_rows,
        "Wildcard FQDN": _wildcard_fqdn_rows,
        "Address Groups": _address_group_rows,
        "Service Categories": _service_category_rows,
        "Services": _service_rows,
        "Service Groups": _service_group_rows,
        "Policies": _policy_rows,
        "IP Pools": _ip_pool_rows,
        "Virtual IPs": _vip_rows,
        "VIP Real Servers": _vip_real_server_rows,
        "VIP Groups": _vip_group_rows,
        "NAT Rules": _nat_rows,
        "Routes": _route_rows,
        "VPN Tunnels": _vpn_phase1_rows,
        "VPN Phase 2": _vpn_phase2_rows,
        "DHCP Servers": _dhcp_server_rows,
        "DHCP IP Ranges": _dhcp_ip_range_rows,
        "DHCP Reservations": _dhcp_reservation_rows,
        "SD-WAN": _sdwan_rows,
        "SD-WAN Zones": _sdwan_zone_rows,
        "SD-WAN Members": _sdwan_member_rows,
        "SD-WAN Health Checks": _sdwan_health_rows,
        "SD-WAN Rules": _sdwan_rule_rows,
        "SSL VPN Settings": _ssl_settings_rows,
        "SSL VPN Portals": _ssl_portal_rows,
        "SSL VPN Authentication Rules": _ssl_auth_rule_rows,
        "SSL VPN Host Checks": _ssl_host_check_rows,
        "SSL VPN Host Check Items": _ssl_host_check_item_rows,
        "Local Users": _local_user_rows,
        "User Groups": _user_group_rows,
        "User Group Matches": _user_group_match_rows,
        "User Group Guests": _user_group_guest_rows,
        "Administrators": _administrator_rows,
        "Admin Profiles": _admin_profile_rows,
        "Admin Profile Permissions": _admin_permission_rows,
        "IPS Sensors": _ips_sensor_rows,
        "IPS Sensor Entries": _ips_entry_rows,
        "IPS Exempt IPs": _ips_exempt_rows,
        "Security Profiles": _security_profile_rows,
        "External Resources": _external_resource_rows,
        "Firewall Policy Source Settings": _policy_source_rows,
        "Interface Source Settings": _interface_source_rows,
        "Interface Nested Configuration": _interface_nested_rows,
        "Unresolved References": _unresolved_reference_rows,
        "Unsupported": _unsupported_rows,
        "FortiGate Source Inventory": _fortigate_source_inventory_rows,
        "Extraction Coverage": _coverage_rows,
    }

    builder = builders.get(sheet_name)

    if builder is not None:
        return builder(context, headers)

    return _generic_source_rows(sheet_name, context, headers)


def _build_summary(
    workbook: Workbook,
    context: _ExcelContext,
    order: Sequence[str],
) -> None:
    sheet = workbook.create_sheet("Summary")
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A5"

    sheet.merge_cells("A1:F1")
    sheet["A1"] = "FortiGate Configuration Report"
    sheet["A1"].fill = _TITLE_FILL
    sheet["A1"].font = Font(color="FFFFFF", bold=True, size=18)
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 30

    metadata = [
        ("Source File", context.source_name),
        ("Hostname", _hostname(context)),
        ("FortiOS Version", None),
        ("VDOMs", "\n".join(_vdoms(context))),
        ("Generated UTC", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")),
        ("Validation Errors", len(context.validation.errors)),
        ("Validation Warnings", len(context.validation.warnings)),
    ]

    for row_number, (label, value) in enumerate(metadata, start=3):
        sheet.cell(row_number, 1, label).font = Font(bold=True, color="41504C")
        sheet.cell(row_number, 2, _excel_safe(value))
        sheet.cell(row_number, 2).alignment = Alignment(wrap_text=True, vertical="top")

    section_row = 11
    sheet.cell(section_row, 1, "Inventory")
    sheet.cell(section_row, 1).fill = _HEADER_FILL
    sheet.cell(section_row, 1).font = _WHITE_FONT
    sheet.merge_cells(start_row=section_row, start_column=1, end_row=section_row, end_column=3)

    counts = _inventory_counts(context)

    row_number = section_row + 1
    for label, count, target in counts:
        sheet.cell(row_number, 1, label)
        sheet.cell(row_number, 2, count)
        if target in order:
            cell = sheet.cell(row_number, 3, "Open sheet")
            cell.hyperlink = f"#'{target}'!A1"
            cell.font = _LINK_FONT
        row_number += 1

    indicator_row = row_number + 1
    sheet.cell(indicator_row, 1, "Migration Indicators")
    sheet.cell(indicator_row, 1).fill = _HEADER_FILL
    sheet.cell(indicator_row, 1).font = _WHITE_FONT
    sheet.merge_cells(
        start_row=indicator_row,
        start_column=1,
        end_row=indicator_row,
        end_column=3,
    )

    for row_number, (label, value) in enumerate(
        _migration_indicators(context),
        start=indicator_row + 1,
    ):
        sheet.cell(row_number, 1, label)
        sheet.cell(row_number, 2, value)

    nav_col = 5
    sheet.cell(3, nav_col, "Workbook navigation")
    sheet.cell(3, nav_col).fill = _HEADER_FILL
    sheet.cell(3, nav_col).font = _WHITE_FONT

    nav_row = 4
    for name in order:
        if name == "Summary":
            continue
        cell = sheet.cell(nav_row, nav_col, name)
        cell.hyperlink = f"#'{name}'!A1"
        cell.font = _LINK_FONT
        nav_row += 1

    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 24
    sheet.column_dimensions["C"].width = 16
    sheet.column_dimensions["E"].width = 34


def _write_table_sheet(
    workbook: Workbook,
    sheet_name: str,
    headers: Sequence[str],
    rows: Sequence[Mapping[str, Any]],
    *,
    context: _ExcelContext,
) -> None:
    sheet = workbook.create_sheet(sheet_name)
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 90

    max_col = max(len(headers), 1)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
    sheet.cell(1, 1, sheet_name)
    sheet.cell(1, 1).fill = _TITLE_FILL
    sheet.cell(1, 1).font = _TITLE_FONT
    sheet.cell(1, 1).alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 24

    if max_col > 2:
        sheet.merge_cells(
            start_row=2,
            start_column=1,
            end_row=2,
            end_column=max_col - 1,
        )
    note = sheet.cell(2, 1)
    note.value = _sheet_note(sheet_name, len(rows))
    note.font = _MUTED_FONT
    note.alignment = Alignment(wrap_text=True, vertical="top")

    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(3, column, header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = _BORDER


    for index, row in enumerate(rows, start=4):
        for column, header in enumerate(headers, start=1):
            value = _excel_safe(row.get(header))
            cell = sheet.cell(index, column, value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")

        if index % 2 == 1:
            for column in range(1, max_col + 1):
                sheet.cell(index, column).fill = _ALT_FILL

        for column, header in enumerate(headers, start=1):
            if header in DERIVED_COLUMNS_BY_SHEET.get(sheet_name, ()):
                sheet.cell(index, column).fill = _DERIVED_FILL

        outline = row.get("__outline_level__")
        if isinstance(outline, int) and outline > 0:
            sheet.row_dimensions[index].outlineLevel = min(outline, 7)

        if row.get("__review__"):
            fill = _ERROR_FILL if row.get("__error__") else _REVIEW_FILL
            for column in range(1, max_col + 1):
                sheet.cell(index, column).fill = fill

    if rows:
        sheet.auto_filter.ref = f"A3:{get_column_letter(max_col)}{len(rows) + 3}"

    sheet.freeze_panes = _freeze_pane(sheet_name, headers)
    _apply_widths(sheet, headers)
    for column, header in enumerate(headers, start=1):
        if header in HIDDEN_COLUMNS_BY_DEFAULT.get(sheet_name, ()):
            sheet.column_dimensions[get_column_letter(column)].hidden = True
    _add_back_link(sheet)

    if sheet_name == "Review Required":
        _apply_review_colors(sheet, headers, len(rows))


def _add_back_link(sheet) -> None:
    if sheet.max_column <= 1:
        return

    cell = sheet.cell(2, sheet.max_column)
    cell.value = "Back to Summary"
    cell.hyperlink = "#'Summary'!A1"
    cell.font = _LINK_FONT
    cell.alignment = Alignment(horizontal="right")


def _apply_review_colors(sheet, headers: Sequence[str], row_count: int) -> None:
    if "Severity" not in headers:
        return

    severity_col = headers.index("Severity") + 1

    for row in range(4, row_count + 4):
        severity = str(sheet.cell(row, severity_col).value or "").lower()
        fill = _ERROR_FILL if severity == "error" else _REVIEW_FILL

        for column in range(1, len(headers) + 1):
            sheet.cell(row, column).fill = fill


def _apply_widths(sheet, headers: Sequence[str]) -> None:
    for index, header in enumerate(headers, start=1):
        normalized = header.lower()

        if any(token in normalized for token in ("additional settings", "review reason", "raw capture", "notes")):
            width = 36
        elif any(token in normalized for token in ("description", "comment", "source path", "setting", "value")):
            width = 28
        elif any(token in normalized for token in ("members", "addresses", "interfaces", "services", "profiles", "references")):
            width = 26
        elif any(token in normalized for token in ("name", "object", "category", "status", "type")):
            width = 20
        elif "id" in normalized or "port" in normalized or "vlan" in normalized:
            width = 14
        else:
            width = 18

        sheet.column_dimensions[get_column_letter(index)].width = width


def _freeze_pane(sheet_name: str, headers: Sequence[str]) -> str:
    if sheet_name == "Interfaces":
        return "D4"
    if sheet_name == "Policies":
        return "D4"
    if len(headers) > 12:
        return "C4"
    return "A4"


def _sheet_note(
    sheet_name: str,
    row_count: int,
) -> str:
    if sheet_name == "FortiGate Source Inventory":
        return (
            "Explicit FortiGate source commands retained for traceability. "
            "Extraction Status identifies whether the source section has "
            "dedicated typed extraction support."
        )

    if row_count:
        return (
            f"{row_count} record(s). "
            "Values are explicit FortiGate source data unless the column is "
            "identified as derived or analysis output."
        )

    return "No matching explicit FortiGate source records were extracted."


def _inventory_counts(context: _ExcelContext) -> list[tuple[str, int, str]]:
    config = context.config
    unsupported = {
        record.source_path
        for record in context.extracted.source_objects
        if record.source_path not in set(registered_sections())
    }
    return [
        ("Interfaces", len(config.interfaces), "Interfaces"),
        ("Zones", len(config.zones), "Zones"),
        ("Addresses", len(config.addresses), "Addresses"),
        ("Address Groups", len(config.address_groups), "Address Groups"),
        ("Services", len(context.derived.services.services), "Services"),
        ("Service Groups", len(context.derived.services.groups), "Service Groups"),
        ("Policies", len(config.policies), "Policies"),
        ("NAT Rules", len(context.derived.nat), "NAT Rules"),
        ("IP Pools", len(config.ip_pools), "IP Pools"),
        ("Virtual IPs", len(config.vips), "Virtual IPs"),
        ("Routes", len(config.static_routes), "Routes"),
        ("VPN Tunnels", len(config.ipsec_phase1), "VPN Tunnels"),
        ("VPN Phase 2", len(config.ipsec_phase2), "VPN Phase 2"),
        ("Local Users", len(config.local_users), "Local Users"),
        ("DHCP Servers", len(config.dhcp_servers), "DHCP Servers"),
        ("SD-WAN Members", sum(len(item.members) for item in config.sdwans), "SD-WAN Members"),
        ("SSL VPN Portals", len(config.ssl_vpn_portals), "SSL VPN Portals"),
        ("Administrators", len(config.administrators), "Administrators"),
        ("IPS Sensors", len(config.ips_sensors), "IPS Sensors"),
        ("External Resources", len(config.external_resources), "External Resources"),
        ("Review Required", len(context.validation.issues), "Review Required"),
        ("Unsupported Source Sections", len(unsupported), "Unsupported"),
    ]


def _hostname(context: _ExcelContext) -> str | None:
    for record in context.source_by_path.get("system global", ()):
        for key, value in record.values.items():
            if _normalize_key(key) == "hostname":
                return value
    return None


def _migration_indicators(context: _ExcelContext) -> list[tuple[str, Any]]:
    config = context.config
    topology_issues = sum(
        len(item.issues)
        for item in (
            *context.derived.topology.interfaces,
            *context.derived.topology.vpns,
        )
    )
    nat_ambiguities = sum(
        "Multiple possible outgoing interfaces" in issue
        for item in context.derived.nat
        for issue in item.issues
    )
    ipv6_explicit = any(
        "ipv6" in _normalize_key(key) or "ip6" in _normalize_key(key)
        for record in context.extracted.source_objects
        for key in record.values
    )
    dynamic_wan = any(
        (interface.mode or "").lower() in {"dhcp", "pppoe"}
        for interface in config.interfaces
    )
    return [
        ("IPv6 Explicit Configuration Present", "Yes" if ipv6_explicit else "No"),
        ("Dynamic WAN Addressing Present", "Yes" if dynamic_wan else "No"),
        ("SD-WAN Present", "Yes" if config.sdwans else "No"),
        ("Interface-NAT Ambiguities", nat_ambiguities),
        ("Broken References", len(context.derived.broken_references)),
        ("Policy Names Truncated", sum(item.truncated for item in context.derived.policy_names)),
        ("Policy Name Collisions", sum(item.collision for item in context.derived.policy_names)),
        ("Topology Issues", topology_issues),
    ]


def _vdoms(context: _ExcelContext) -> list[str]:
    values = {
        record.vdom
        for record in context.extracted.source_objects
    }
    values.update(
        getattr(item, "vdom", "root")
        for collection_name in (
            "interfaces",
            "zones",
            "addresses",
            "policies",
            "static_routes",
            "ipsec_phase1",
        )
        for item in getattr(context.config, collection_name, ())
    )
    return sorted(values)


def _review_rows(context: _ExcelContext, headers: Sequence[str]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []

    for issue in context.validation.issues:
        row = {
            "Severity": issue.severity.value,
            "Category": issue.domain,
            "Object": issue.object_name,
            "VDOM": issue.vdom,
            "Field": issue.field,
            "Issue / Review Reason": issue.message,
            "Source Sheet": _sheet_for_domain(issue.domain),
        }
        rows.append(row)

    return rows


def _interface_rows(context: _ExcelContext, headers: Sequence[str]) -> list[dict[str, Any]]:
    topology = {(item.vdom, item.name): item for item in context.derived.topology.interfaces}
    interfaces = {(item.vdom, item.name): item for item in context.config.interfaces}
    aggregate_members: dict[tuple[str, str], list[tuple[str, str]]] = defaultdict(list)
    children: dict[tuple[str, str], list[tuple[str, str]]] = defaultdict(list)
    owned: set[tuple[str, str]] = set()
    for interface in context.config.interfaces:
        key = (interface.vdom, interface.name)
        if interface.interface and (interface.vdom, interface.interface) in interfaces:
            children[(interface.vdom, interface.interface)].append(key)
        for member in interface.members:
            member_key = (interface.vdom, member)
            aggregate_members[key].append(member_key)
            owned.add(member_key)

    vpns_by_interface: dict[tuple[str, str], list[Any]] = defaultdict(list)
    vpns = {(item.vdom, item.name): item for item in context.config.ipsec_phase1}
    for vpn in context.derived.topology.vpns:
        if vpn.attached_interface:
            vpns_by_interface[(vpn.vdom, vpn.attached_interface)].append(vpn)

    def rank(key: tuple[str, str]) -> tuple[int, str]:
        return _interface_rank(topology.get(key), interfaces[key])

    roots = sorted(
        (key for key in interfaces if key not in owned and not interfaces[key].interface),
        key=rank,
    )
    zones_by_interface: dict[tuple[str, str], list[str]] = defaultdict(list)
    for zone in context.config.zones:
        for member in zone.members:
            zones_by_interface[(zone.vdom, member)].append(zone.name)
    rows: list[dict[str, Any]] = []
    visited: set[tuple[str, str]] = set()
    emitted_vpns: set[tuple[str, str]] = set()

    def emit_vpn(vpn_top: Any, prefix: str = "") -> None:
        vpn_key = (vpn_top.vdom, vpn_top.name)
        if vpn_key in emitted_vpns:
            return
        emitted_vpns.add(vpn_key)
        vpn = vpns.get(vpn_key)
        if vpn is None:
            return
        row = {
            "Name": f"{prefix}{_topology_symbol('vpn')} {vpn.name}",
            "Type": vpn.type,
            "Parent Interface": vpn_top.attached_interface,
            "Aggregate": vpn_top.aggregate,
            "Physical Interfaces": list(vpn_top.physical_interfaces),
            "Topology Path": list(vpn_top.path),
            "VDOM": vpn.vdom,
            "Topology Issues": list(vpn_top.issues),
            "Source Explicit Fields": sorted(vpn.explicit_fields),
            "Additional Settings": sanitize_source_attributes(vpn.raw_extra),
        }
        _add_analysis_status(
            row,
            context,
            vdom=vpn.vdom,
            names=(vpn.name,),
            domains=("ipsec_phase1", "vpn_topology"),
            extra_reasons=vpn_top.issues,
        )
        _overlay_raw(row, vpn.raw_extra, headers)
        rows.append(row)

    def emit_interface(key: tuple[str, str], prefix: str = "", child_indent: str = "") -> None:
        if key in visited:
            return
        visited.add(key)
        interface = interfaces[key]
        key = (interface.vdom, interface.name)
        top = topology.get(key)
        row = {
            "Name": f"{prefix}{_topology_symbol(top.kind if top else None)} {interface.name}",
            "Alias": interface.alias,
            "Zone": zones_by_interface.get(key, []),
            "IP / Prefix": interface.ip,
            "Type": interface.type,
            "Role": interface.role,
            "Addressing Mode": interface.mode,
            "Management Access": interface.allowaccess,
            "VLAN ID": interface.vlanid,
            "Parent Interface": interface.interface,
            "Aggregate": top.aggregate if top else None,
            "Physical Interfaces": list(top.physical_interfaces) if top else [],
            "Topology Path": list(top.path) if top else [interface.name],
            "Members": interface.members,
            "Status": _enabled_text(interface.status),
            "Description": interface.description,
            "VDOM": interface.vdom,
            "Topology Issues": list(top.issues) if top else [],
            "Source Explicit Fields": sorted(interface.explicit_fields),
            "Additional Settings": sanitize_source_attributes(interface.raw_extra),
        }

        _add_analysis_status(
            row,
            context,
            vdom=interface.vdom,
            names=(interface.name,),
            domains=("interface", "interface_topology"),
        )
        source_values = _interface_source_values(
            context,
            vdom=interface.vdom,
            name=interface.name,
        )

        row["Secondary IPv4 Addresses"] = [secondary.ip for secondary in interface.secondary_ips if secondary.ip]
        _overlay_raw(row, source_values, headers)
        _overlay_raw(row, interface.raw_extra, headers)
        row["Additional Settings"] = _additional_source_settings(
            source_values,
            row,
            headers,
        )
        rows.append(row)

        child_keys = tuple(aggregate_members.get(key, ())) + tuple(children.get(key, ()))
        child_keys = tuple(sorted((item for item in dict.fromkeys(child_keys) if item in interfaces), key=rank))
        vpn_children = tuple(vpns_by_interface.get(key, ()))
        children_count = len(child_keys) + len(vpn_children)
        child_number = 0
        for child_key in child_keys:
            child_number += 1
            last = child_number == children_count
            emit_interface(child_key, child_indent + ("└─ " if last else "├─ "), child_indent + ("   " if last else "│  "))
        for vpn_top in vpn_children:
            child_number += 1
            last = child_number == children_count
            emit_vpn(vpn_top, child_indent + ("└─ " if last else "├─ "))

    for key in roots:
        emit_interface(key)
    for key in sorted(interfaces, key=rank):
        emit_interface(key)
    for vpn_top in context.derived.topology.vpns:
        if (vpn_top.vdom, vpn_top.name) not in emitted_vpns:
            emit_vpn(vpn_top)

    return rows


def _interface_rank(topology: Any, interface: Any) -> tuple[int, str]:
    return (
        {"aggregate": 0, "physical": 1, "vlan": 2, "logical": 3, "tunnel": 4}.get(
            topology.kind if topology else None,
            5,
        ),
        interface.name,
    )


def _topology_symbol(kind: str | None) -> str:
    return {"aggregate": "◆", "physical": "●", "vlan": "▣", "logical": "◇", "tunnel": "◇", "vpn": "◈"}.get(kind, "◇")


def _interface_secondary_rows(context: _ExcelContext, headers: Sequence[str]) -> list[dict[str, Any]]:
    rows = []
    for interface in context.config.interfaces:
        for item in interface.secondary_ips:
            row = {
                "Interface": interface.name,
                "ID": item.id,
                "IP / Prefix": item.ip,
                "Management Access": item.allowaccess,
                "HA Priority": item.ha_priority,
                "Additional Settings": _additional_settings(item),
            }
            _add_analysis_status(
                row,
                context,
                vdom=interface.vdom,
                names=(interface.name, item.id),
            )
            _overlay_raw(row, item.raw_extra, headers)
            rows.append(row)
    return rows


def _zone_rows(context: _ExcelContext, headers: Sequence[str]) -> list[dict[str, Any]]:
    rows = []
    for item in context.config.zones:
        row = {
            "Name": item.name,
            "Members": item.members,
            "Description": item.description,
            "Intrazone": item.intrazone,
            "VDOM": item.vdom,
            "Additional Settings": _additional_settings(item),
        }
        _add_analysis_status(row, context, vdom=item.vdom, names=(item.name,))
        _overlay_raw(row, item.raw_extra, headers)
        rows.append(row)
    return rows


def _address_rows(context: _ExcelContext, headers: Sequence[str]) -> list[dict[str, Any]]:
    rows = []
    for item in context.config.addresses:
        value = item.subnet or (
            f"{item.start_ip}-{item.end_ip}" if item.start_ip and item.end_ip else None
        ) or item.fqdn or item.wildcard or item.wildcard_fqdn
        row = {
            "Name": item.name,
            "Type": item.type,
            "Value": value,
            "Address Family": item.address_family,
            "Associated Interface": item.associated_interface,
            "Allow Routing": item.allow_routing,
            "Tags": [
                tag
                for entry in item.tagging
                for tag in entry.tags
            ],
            "Description": item.comment,
            "VDOM": item.vdom,
            "Additional Settings": sanitize_source_attributes(item.raw_extra),
        }
        _add_analysis_status(
            row,
            context,
            vdom=item.vdom,
            names=(item.name,),
            domains=("address6",) if item.address_family == "ipv6" else ("address",),
        )
        _overlay_raw(row, item.raw_extra, headers)
        rows.append(row)
    return rows


def _wildcard_fqdn_rows(context: _ExcelContext, headers: Sequence[str]) -> list[dict[str, Any]]:
    rows = []
    for item in context.config.wildcard_fqdns:
        row = {
            "Name": item.name,
            "Wildcard FQDN": item.wildcard_fqdn,
            "Description": item.comment,
            "VDOM": item.vdom,
            "Additional Settings": sanitize_source_attributes(item.raw_extra),
        }
        _add_analysis_status(row, context, vdom=item.vdom, names=(item.name,))
        rows.append(row)
    return rows


def _address_group_rows(context: _ExcelContext, headers: Sequence[str]) -> list[dict[str, Any]]:
    rows = []
    for item in context.config.address_groups:
        row = {
            "Name": item.name,
            "Members": item.members,
            "Address Family": item.address_family,
            "Exclusion Enabled": item.exclude,
            "Exclude Members": item.exclude_members,
            "Description": item.comment,
            "Allow Routing": item.allow_routing,
            "Group Type": item.type,
            "Tags": [tag for entry in item.tagging for tag in entry.tags],
            "VDOM": item.vdom,
            "Additional Settings": sanitize_source_attributes(item.raw_extra),
        }
        _add_analysis_status(
            row,
            context,
            vdom=item.vdom,
            names=(item.name,),
            domains=("address_group6",)
            if item.address_family == "ipv6"
            else ("address_group",),
        )
        _overlay_raw(row, item.raw_extra, headers)
        rows.append(row)
    return rows


def _address_group_tag_rows(context: _ExcelContext, headers: Sequence[str]) -> list[dict[str, Any]]:
    rows = []
    for group in context.config.address_groups:
        for tag in group.tagging:
            rows.append(
                {
                    "Group Name": group.name,
                    "Address Family": group.address_family,
                    "Tag Entry": tag.name,
                    "Category": tag.category,
                    "Tags": tag.tags,
                    "Extraction Status": "EXTRACTED",
                    "Manual Review": "No",
                    "Additional Settings": sanitize_source_attributes(tag.raw_extra),
                }
            )
    return rows


def _service_category_rows(context: _ExcelContext, headers: Sequence[str]) -> list[dict[str, Any]]:
    rows = []
    for item in context.config.service_categories:
        row = {
            "Name": item.name,
            "Description": item.comment,
            "Extraction Status": "EXTRACTED",
            "Additional Settings": sanitize_source_attributes(item.raw_extra),
        }
        _overlay_raw(row, item.raw_extra, headers)
        rows.append(row)
    return rows


def _service_rows(context: _ExcelContext, headers: Sequence[str]) -> list[dict[str, Any]]:
    source = {(item.vdom, item.name): item for item in context.config.services}
    rows = []
    for item in context.derived.services.services:
        source_item = source.get((item.vdom, item.source_name or item.name))
        row = {
            "Name": item.name,
            "Source Service": item.source_name,
            "Protocol": item.protocol,
            "Destination Port": item.port,
            "Source Port": item.source_port,
            "Protocol Number": item.protocol_number,
            "ICMP Type": item.icmp_type,
            "ICMP Code": item.icmp_code,
            "Generated": item.generated,
            "Description": item.comment,
            "VDOM": item.vdom,
        }
        _add_analysis_status(
            row,
            context,
            vdom=item.vdom,
            names=(item.source_name, item.name),
        )
        if source_item is not None:
            _overlay_raw(row, source_item.raw_extra, headers)
        rows.append(row)
    return rows


def _service_group_rows(context: _ExcelContext, headers: Sequence[str]) -> list[dict[str, Any]]:
    source = {(item.vdom, item.name): item for item in context.config.service_groups}
    rows = []
    for item in context.derived.services.groups:
        source_item = source.get((item.vdom, item.name))
        row = {
            "Name": item.name,
            "Members": list(item.members),
            "Generated": item.generated,
            "Description": item.comment,
            "VDOM": item.vdom,
        }
        _add_analysis_status(row, context, vdom=item.vdom, names=(item.name,))
        if source_item is not None:
            _overlay_raw(row, source_item.raw_extra, headers)
        rows.append(row)
    return rows


def _policy_rows(context: _ExcelContext, headers: Sequence[str]) -> list[dict[str, Any]]:
    normalized = {
        (item.vdom, item.policy_id): item
        for item in context.derived.policy_names
    }
    nat = {
        (item.vdom, item.policy_id): item
        for item in context.derived.nat
    }
    rows = []
    for item in context.config.policies:
        name_info = normalized.get((item.vdom, item.policy_id))
        nat_item = nat.get((item.vdom, item.policy_id))
        row = {
            "Rule #": item.policy_id,
            "Policy Name": name_info.normalized_name if name_info else item.name,
            "Source Name": (
                name_info.source_name
                if name_info is not None and name_info.truncated
                else None
            ),
            "Source Interface": item.srcintf,
            "Source Addresses": item.srcaddr,
            "Destination Interface": item.dstintf,
            "Destination Addresses": item.dstaddr,
            "Services": item.service,
            "Action": item.action,
            "Schedule": item.schedule,
            "NAT Enabled": _enabled_text(item.nat),
            "UTM Status": item.utm_status,
            "Source Address Negate": item.srcaddr_negate,
            "Destination Address Negate": item.dstaddr_negate,
            "User Groups": item.groups,
            "Users": item.users,
            "Service Negate": item.service_negate,
            "VPN Tunnel": item.vpntunnel,
            "Security Profile Group": item.profile_group,
            "Antivirus": item.av_profile,
            "IPS Sensor": item.ips_sensor,
            "Web Filter": item.webfilter_profile,
            "Application List": item.application_list,
            "DNS Filter": item.dnsfilter_profile,
            "SSL/SSH Profile": item.ssl_ssh_profile,
            "Log Setting": item.logtraffic,
            "Comments": item.comments,
            "Status": _enabled_text(item.status),
            "SNAT Type": _nat_type(nat_item.translation_type) if nat_item else None,
            "SNAT Address": list(nat_item.translated_addresses) if nat_item else [],
            "IP Pool Name": list(nat_item.pool_names) if nat_item else [],
            "VDOM": item.vdom,
            "Source Explicit Fields": sorted(item.explicit_fields),
            "Additional Settings": _additional_settings(item),
        }
        _add_analysis_status(
            row,
            context,
            vdom=item.vdom,
            names=(item.name, item.policy_id),
            extra_reasons=(
                ("Policy-name normalization collision",)
                if name_info is not None and name_info.collision
                else ()
            ),
        )
        _overlay_raw(row, item.raw_extra, headers)
        _add_analysis_status(row, context, vdom=item.vdom, names=(item.name,))
        rows.append(row)
    return rows


def _ip_pool_rows(context: _ExcelContext, headers: Sequence[str]) -> list[dict[str, Any]]:
    return _model_rows(
        context,
        context.config.ip_pools,
        headers,
        {
            "Name": "name",
            "Type": "type",
            "Start IP": "startip",
            "End IP": "endip",
            "Source Start IP": "source_startip",
            "Source End IP": "source_endip",
            "Start Port": "startport",
            "End Port": "endport",
            "Associated Interface": "associated_interface",
            "ARP Reply": "arp_reply",
            "ARP Interface": "arp_intf",
            "Permit Any Host": "permit_any_host",
            "Excluded IPs": "exclude_ip",
            "NAT64": "nat64",
            "Add NAT64 Route": "add_nat64_route",
            "Description": "comments",
            "Comments": "comments",
            "Source Explicit Fields": "explicit_fields",
        },
    )


def _vip_rows(
    context: _ExcelContext,
    headers: Sequence[str],
) -> list[dict[str, Any]]:
    return _model_rows(
        context,
        context.config.vips,
        headers,
        {
            "Name": "name",
            "Type": "type",
            "Status": "status",
            "External IP": "extip",
            "External Address Objects": "extaddr",
            "External Interface": "extintf",
            "Mapped IPs": "mappedip",
            "Real Servers": lambda item: [
                ":".join(
                    str(value)
                    for value in (server.ip or server.address, server.port)
                    if value is not None
                )
                for server in item.realservers
            ],
            "Real Server Count": lambda item: len(item.realservers),
            "Port Forward": "portforward",
            "Protocol": "protocol",
            "External Port": "extport",
            "Mapped Port": "mappedport",
            "ARP Reply": "arp_reply",
            "NAT Source VIP": "nat_source_vip",
            "Services": "service",
            "Load Balance Method": "ldb_method",
            "Server Type": "server_type",
            "Monitors": "monitor",
            "Description": "comment",
            "VDOM": "vdom",
        },
    )


def _vip_real_server_rows(
    context: _ExcelContext,
    headers: Sequence[str],
) -> list[dict[str, Any]]:
    rows = []

    for vip in context.config.vips:
        for item in vip.realservers:
            row = {
                "VIP Name": vip.name,
                "Server ID": item.id,
                "IP": item.ip,
                "Address": item.address,
                "Port": item.port,
                "Status": item.status,
                "Weight": item.weight,
                "Monitors": item.monitor,
                "VDOM": vip.vdom,
                "Additional Settings": _additional_settings(item),
            }
            _add_analysis_status(
                row,
                context,
                vdom=vip.vdom,
                names=(vip.name, item.id),
            )
            _overlay_raw(
                row,
                item.raw_extra,
                headers,
            )
            rows.append(row)

    return rows


def _vip_group_rows(context: _ExcelContext, headers: Sequence[str]) -> list[dict[str, Any]]:
    return _model_rows(
        context,
        context.config.vip_groups,
        headers,
        {
            "Name": "name",
            "Interface": "interface",
            "Members": "members",
            "Comments": "comments",
            "VDOM": "vdom",
        },
    )


def _nat_rows(
    context: _ExcelContext,
    headers: Sequence[str],
) -> list[dict[str, Any]]:
    policies = {
        (policy.vdom, policy.policy_id): policy
        for policy in context.config.policies
    }

    rows: list[dict[str, Any]] = []

    for item in context.derived.nat:
        policy = policies.get(
            (item.vdom, item.policy_id)
        )

        row = {
            "Rule #": item.policy_id,
            "Policy Name": item.policy_name,
            "Source Interface": (
                policy.srcintf
                if policy is not None
                else []
            ),
            "Destination Interface": (
                policy.dstintf
                if policy is not None
                else list(item.egress_interfaces)
            ),
            "Source Addresses": (
                policy.srcaddr
                if policy is not None
                else []
            ),
            "Destination Addresses": (
                policy.dstaddr
                if policy is not None
                else []
            ),
            "Services": (
                policy.service
                if policy is not None
                else []
            ),
            "NAT Enabled": _enabled_text(policy.nat) if policy is not None else None,
            "SNAT Type": _nat_type(item.translation_type),
            "SNAT Address": list(item.translated_addresses),
            "IP Pool Name": list(item.pool_names),
            "Egress Interfaces": list(
                item.egress_interfaces
            ),
            "VDOM": item.vdom,
        }

        _add_analysis_status(
            row,
            context,
            vdom=item.vdom,
            names=(
                item.policy_name,
                item.policy_id,
            ),
            extra_reasons=item.issues,
        )

        rows.append(row)

    return rows


def _route_rows(
    context: _ExcelContext,
    headers: Sequence[str],
) -> list[dict[str, Any]]:
    return _model_rows(
        context,
        context.config.static_routes,
        headers,
        {
            "Route ID": "seq_num",
            "Destination": "dst",
            "Destination Address Object": "dstaddr",
            "Interface": "device",
            "Gateway": "gateway",
            "Distance": "distance",
            "Priority": "priority",
            "Status": "status",
            "SD-WAN Zone": "sdwan_zone",
            "Preferred Source": "preferred_source",
            "Source Prefix": "src",
            "Dynamic Gateway": "dynamic_gateway",
            "Blackhole": "blackhole",
            "Description": "comment",
            "Address Family": "address_family",
            "VDOM": "vdom",
        },
    )


def _vpn_phase1_rows(
    context: _ExcelContext,
    headers: Sequence[str],
) -> list[dict[str, Any]]:
    topology = {
        (item.vdom, item.name): item
        for item in context.derived.topology.vpns
    }

    rows: list[dict[str, Any]] = []

    for item in context.config.ipsec_phase1:
        top = topology.get(
            (item.vdom, item.name)
        )

        row = {
            "Name": item.name,
            "Local Interface": item.interface,
            "Remote Gateway IPv4": item.remote_gw,
            "Remote Gateway DDNS": item.remotegw_ddns,
            "Type": item.type,
            "IKE Version": item.ike_version,
            "Authentication Method": item.authmethod,
            "Proposal": item.proposal,
            "DH Groups": item.dhgrp,
            "Key Lifetime": item.keylife,
            "NAT Traversal": item.nattraversal,
            "DPD": item.dpd,
            "Local Gateway": item.local_gw,
            "Local ID": item.localid,
            "Peer ID": item.peerid,
            "Certificate": item.certificate,
            "Description": item.comments,
            "VDOM": item.vdom,
            "Aggregate": (
                top.aggregate
                if top
                else None
            ),
            "Attached Physical Interfaces": (
                list(top.physical_interfaces)
                if top
                else []
            ),
            "Topology Path": (
                list(top.path)
                if top
                else []
            ),
            "Topology Issues": (
                list(top.issues)
                if top
                else []
            ),
            "Source Explicit Fields": sorted(
                item.explicit_fields
            ),
            "Additional Settings": _additional_settings(item),
        }

        _add_analysis_status(
            row,
            context,
            vdom=item.vdom,
            names=(item.name,),
            domains=("ipsec_phase1", "vpn_topology"),
            extra_reasons=(
                top.issues
                if top
                else ()
            ),
        )

        _overlay_raw(
            row,
            item.raw_extra,
            headers,
        )

        rows.append(row)

    return rows


def _vpn_phase2_rows(
    context: _ExcelContext,
    headers: Sequence[str],
) -> list[dict[str, Any]]:
    normalized = {
        (item.vdom, item.name): item
        for item in context.derived.vpn.phase2
    }

    rows: list[dict[str, Any]] = []

    for item in context.config.ipsec_phase2:
        norm = normalized.get(
            (item.vdom, item.name)
        )

        row = {
            "Name": item.name,
            "Phase 1": item.phase1name,
            "Proposal": item.proposal,
            "PFS": item.pfs,
            "DH Groups": item.dhgrp,
            "Keylife Seconds": item.keylifeseconds,
            "Keylife KB": item.keylifekbs,
            "Source Range": (
                norm.source_range
                if norm
                else None
            ),
            "Destination Range": (
                norm.destination_range
                if norm
                else None
            ),
            "Auto Negotiate": item.auto_negotiate,
            "VDOM": item.vdom,
            "Comments": item.comments,
            "Source Explicit Fields": sorted(
                item.explicit_fields
            ),
            "Additional Settings": _additional_settings(item),
        }

        _add_analysis_status(
            row,
            context,
            vdom=item.vdom,
            names=(item.name,),
            domains=("vpn_phase2",),
        )

        _overlay_raw(
            row,
            item.raw_extra,
            headers,
        )

        rows.append(row)

    return rows


def _dhcp_server_rows(
    context: _ExcelContext,
    headers: Sequence[str],
) -> list[dict[str, Any]]:
    return _model_rows(
        context,
        context.config.dhcp_servers,
        headers,
        {
            "Server ID": "id",
            "Interface": "interface",
            "Status": "status",
            "Server Type": "server_type",
            "IP Mode": "ip_mode",
            "Default Gateway": "default_gateway",
            "Netmask": "netmask",
            "Lease Time": "lease_time",
            "DNS Service": "dns_service",
            "DNS Server 1": "dns_server1",
            "DNS Server 2": "dns_server2",
            "DNS Server 3": "dns_server3",
            "DNS Server 4": "dns_server4",
            "Timezone Option": "timezone_option",
            "Timezone": "timezone",
            "Relay Agent": "relay_agent",
            "VDOM": "vdom",
        },
    )


def _dhcp_ip_range_rows(context: _ExcelContext, headers: Sequence[str]) -> list[dict[str, Any]]:
    return _dhcp_child_rows(context, headers, "ip_ranges")


def _dhcp_exclude_range_rows(context: _ExcelContext, headers: Sequence[str]) -> list[dict[str, Any]]:
    return _dhcp_child_rows(context, headers, "exclude_ranges")


def _dhcp_reservation_rows(context: _ExcelContext, headers: Sequence[str]) -> list[dict[str, Any]]:
    rows = []
    for server in context.config.dhcp_servers:
        for item in server.reserved_addresses:
            row = {
                "Server ID": server.id,
                "Interface": server.interface,
                "Reservation ID": item.id,
                "IP Address": item.ip,
                "MAC Address": item.mac,
                "Description": item.description,
                "Action": item.action,
                "Type": item.type,
                "VDOM": server.vdom,
                "Extraction Status": "EXTRACTED",
                "Source Explicit Fields": sorted(item.explicit_fields),
                "Additional Settings": _additional_settings(item),
            }
            _overlay_raw(row, item.raw_extra, headers)
            _add_analysis_status(row, context, vdom=server.vdom, names=(server.id, item.id))
            rows.append(row)
    return rows


def _dhcp_child_rows(
    context: _ExcelContext,
    headers: Sequence[str],
    attribute: str,
) -> list[dict[str, Any]]:
    rows = []
    for server in context.config.dhcp_servers:
        for item in getattr(server, attribute):
            row = {
                "Server ID": server.id,
                "Interface": server.interface,
                "Range ID": item.id,
                "Source ID": item.id,
                "Start IP": item.start_ip,
                "End IP": item.end_ip,
                "Lease Time": item.lease_time,
                "VDOM": server.vdom,
                "Extraction Status": "EXTRACTED",
                "Manual Review": "No",
                "Source Explicit Fields": sorted(item.explicit_fields),
                "Additional Settings": _additional_settings(item),
            }
            _overlay_raw(row, item.raw_extra, headers)
            _add_analysis_status(row, context, vdom=server.vdom, names=(server.id, item.id))
            rows.append(row)
    return rows


def _sdwan_rows(context: _ExcelContext, headers: Sequence[str]) -> list[dict[str, Any]]:
    rows = []
    for item in context.config.sdwans:
        row = {
            "Status": item.status,
            "Load Balance Mode": item.load_balance_mode,
            "Extraction Status": "EXTRACTED",
            "Manual Review": "No",
            "Additional Settings": _additional_settings(item),
            "VDOM": item.vdom,
        }
        _overlay_raw(row, item.raw_extra, headers)
        _add_analysis_status(row, context, vdom=item.vdom, names=(item.vdom,))
        rows.append(row)
    return rows


def _sdwan_zone_rows(context: _ExcelContext, headers: Sequence[str]) -> list[dict[str, Any]]:
    rows = []
    for sdwan in context.config.sdwans:
        for item in sdwan.zones:
            row = {
                "Zone Name": item.name,
                "Additional Settings": _additional_settings(item),
                "VDOM": sdwan.vdom,
            }
            _overlay_raw(row, item.raw_extra, headers)
            _add_analysis_status(row, context, vdom=sdwan.vdom, names=(item.name,))
            rows.append(row)
    return rows


def _sdwan_member_rows(context: _ExcelContext, headers: Sequence[str]) -> list[dict[str, Any]]:
    topology = {(item.vdom, item.name): item for item in context.derived.topology.interfaces}
    rows = []
    for sdwan in context.config.sdwans:
        for item in sdwan.members:
            top = topology.get((sdwan.vdom, item.interface or ""))
            row = {
                "ID": item.seq_num,
                "Interface": item.interface,
                "Zone": item.zone,
                "Gateway": item.gateway,
                "Source": item.source,
                "Cost": item.cost,
                "Weight": item.weight,
                "Priority": item.priority,
                "Status": item.status,
                "Additional Settings": _additional_settings(item),
                "VDOM": sdwan.vdom,
                "Physical Interfaces": list(top.physical_interfaces) if top else [],
                "Aggregate": top.aggregate if top else None,
                "Source Explicit Fields": sorted(item.explicit_fields),
            }
            _overlay_raw(row, item.raw_extra, headers)
            _add_analysis_status(row, context, vdom=sdwan.vdom, names=(item.interface, item.seq_num))
            rows.append(row)
    return rows


def _sdwan_health_rows(context: _ExcelContext, headers: Sequence[str]) -> list[dict[str, Any]]:
    rows = []
    for sdwan in context.config.sdwans:
        for item in sdwan.health_checks:
            row = {
                "Name": item.name,
                "Server": item.server,
                "Members": item.members,
                "Protocol": item.protocol,
                "Interval": item.interval,
                "Fail Time": item.failtime,
                "Recovery Time": item.recoverytime,
                "Additional Settings": _additional_settings(item),
                "VDOM": sdwan.vdom,
            }
            _overlay_raw(row, item.raw_extra, headers)
            _add_analysis_status(row, context, vdom=sdwan.vdom, names=(item.name,))
            rows.append(row)
    return rows


def _sdwan_rule_rows(context: _ExcelContext, headers: Sequence[str]) -> list[dict[str, Any]]:
    rows = []
    for sdwan in context.config.sdwans:
        for item in sdwan.services:
            row = {
                "ID": item.id,
                "Name": item.name,
                "Status": item.status,
                "Mode": item.mode,
                "Source": item.src,
                "Destination": item.dst,
                "Priority Members": item.priority_members,
                "Health Checks": item.health_check,
                "Priority Zones": item.priority_zone,
                "Additional Settings": _additional_settings(item),
                "VDOM": sdwan.vdom,
            }
            _overlay_raw(row, item.raw_extra, headers)
            _add_analysis_status(row, context, vdom=sdwan.vdom, names=(item.id, item.name))
            rows.append(row)
    return rows


def _ssl_settings_rows(context: _ExcelContext, headers: Sequence[str]) -> list[dict[str, Any]]:
    return _model_rows(
        context,
        context.config.ssl_vpn_settings,
        headers,
        {
            "Status": "status",
            "Minimum Protocol": "ssl_min_proto_ver",
            "Maximum Protocol": "ssl_max_proto_ver",
            "Authentication Timeout": "auth_timeout",
            "Idle Timeout": "idle_timeout",
            "Port": "port",
            "DNS Server 1": "dns_server1",
            "DNS Server 2": "dns_server2",
            "Server Certificate": "servercert",
            "Source Interfaces": "source_interface",
            "Source Addresses": "source_address",
            "Tunnel IP Pools": "tunnel_ip_pools",
            "Default Portal": "default_portal",
            "VDOM": "vdom",
        },
    )


def _ssl_portal_rows(context: _ExcelContext, headers: Sequence[str]) -> list[dict[str, Any]]:
    return _model_rows(
        context,
        context.config.ssl_vpn_portals,
        headers,
        {
            "Name": "name",
            "Tunnel Mode": "tunnel_mode",
            "IPv6 Tunnel Mode": "ipv6_tunnel_mode",
            "IP Pools": "ip_pools",
            "IPv6 Pools": "ipv6_pools",
            "Split Tunneling": "split_tunneling",
            "Limit User Logins": "limit_user_logins",
            "FortiClient Download": "forticlient_download",
            "Split Tunneling Routing Addresses": "split_tunneling_routing_address",
            "VDOM": "vdom",
        },
    )


def _ssl_auth_rule_rows(context: _ExcelContext, headers: Sequence[str]) -> list[dict[str, Any]]:
    rows = []
    for settings in context.config.ssl_vpn_settings:
        for item in settings.authentication_rules:
            row = {
                "ID": item.id,
                "Auth": item.auth,
                "Cipher": item.cipher,
                "Client Certificate": item.client_cert,
                "Realm": item.realm,
                "Source Interfaces": item.source_interface,
                "Source Addresses": item.source_address,
                "Source Address Negate": item.source_address_negate,
                "Users": item.users,
                "User Peer": item.user_peer,
                "Groups": item.groups,
                "Portal": item.portal,
                "Extraction Status": "EXTRACTED",
                "Manual Review": "No",
                "Additional Settings": _additional_settings(item),
            }
            _overlay_raw(row, item.raw_extra, headers)
            _add_analysis_status(row, context, vdom=settings.vdom, names=(item.id, item.portal))
            rows.append(row)
    return rows


def _ssl_host_check_rows(context: _ExcelContext, headers: Sequence[str]) -> list[dict[str, Any]]:
    return _model_rows(
        context,
        context.config.ssl_vpn_host_check_software,
        headers,
        {
            "Name": "name",
            "Type": "type",
            "OS Type": "os_type",
            "Version": "version",
            "GUID": "guid",
            "Check Item Count": lambda item: len(item.check_items),
            "VDOM": "vdom",
        },
    )


def _ssl_host_check_item_rows(context: _ExcelContext, headers: Sequence[str]) -> list[dict[str, Any]]:
    rows = []
    for software in context.config.ssl_vpn_host_check_software:
        for item in software.check_items:
            row = {
                "Host Check": software.name,
                "ID": item.id,
                "Action": item.action,
                "Type": item.type,
                "Target": item.target,
                "MD5s": item.md5s,
                "Version": item.version,
                "Extraction Status": "EXTRACTED",
                "Manual Review": "No",
            "Additional Settings": _additional_settings(item),
        }
        _overlay_raw(row, item.raw_extra, headers)
        _add_analysis_status(row, context, vdom=item.vdom, names=(item.vdom,))
        rows.append(row)
    return rows


def _local_user_rows(context: _ExcelContext, headers: Sequence[str]) -> list[dict[str, Any]]:
    return _model_rows(
        context,
        context.config.local_users,
        headers,
        {
            "Name": "name",
            "ID": "id",
            "Status": "status",
            "Type": "type",
            "Password Configured": "password_configured",
            "Password Time": "passwd_time",
            "Two Factor": "two_factor",
            "Two Factor Authentication": "two_factor_authentication",
            "Two Factor Notification": "two_factor_notification",
            "FortiToken": "fortitoken",
            "Email": "email_to",
            "LDAP Server": "ldap_server",
            "RADIUS Server": "radius_server",
            "TACACS+ Server": "tacacs_server",
            "Auth Concurrent Override": "auth_concurrent_override",
            "Auth Concurrent Value": "auth_concurrent_value",
            "Authentication Timeout": "authtimeout",
            "Password Policy": "passwd_policy",
            "Workstation": "workstation",
            "Username Sensitivity": "username_sensitivity",
            "PPK Identity": "ppk_identity",
            "PPK Secret Configured": "ppk_secret_configured",
            "VDOM": "vdom",
        },
    )


def _user_group_rows(context: _ExcelContext, headers: Sequence[str]) -> list[dict[str, Any]]:
    return _model_rows(
        context,
        context.config.user_groups,
        headers,
        {
            "Name": "name",
            "ID": "id",
            "Type": "group_type",
            "Group Type": "group_type",
            "Members": "members",
            "Match Count": lambda item: len(item.matches),
            "Auth Concurrent Override": "auth_concurrent_override",
            "Auth Concurrent Value": "auth_concurrent_value",
            "Authentication Timeout": "authtimeout",
            "VDOM": "vdom",
        },
    )


def _user_group_match_rows(context: _ExcelContext, headers: Sequence[str]) -> list[dict[str, Any]]:
    rows = []
    for group in context.config.user_groups:
        for item in group.matches:
            row = {
                "User Group": group.name,
                "Match ID": item.id,
                "Server Name": item.server_name,
                "Group Name": item.group_name,
                "VDOM": group.vdom,
                "Additional Settings": _additional_settings(item),
            }
            _overlay_raw(row, item.raw_extra, headers)
            _add_analysis_status(row, context, vdom=group.vdom, names=(group.name, item.id))
            rows.append(row)
    return rows


def _user_group_guest_rows(context: _ExcelContext, headers: Sequence[str]) -> list[dict[str, Any]]:
    rows = []
    for group in context.config.user_groups:
        for item in group.guests:
            row = {
                "User Group": group.name,
                "Guest ID": item.id,
                "Name": item.name,
                "User ID": item.user_id,
                "Email": item.email,
                "Mobile Phone": item.mobile_phone,
                "Expiration": item.expiration,
                "Sponsor": item.sponsor,
                "Comment": item.comment,
                "Password Configured": item.password_configured,
                "VDOM": group.vdom,
                "Additional Settings": _additional_settings(item),
            }
            _overlay_raw(row, item.raw_extra, headers)
            _add_analysis_status(row, context, vdom=group.vdom, names=(group.name, item.id, item.name))
            rows.append(row)
    return rows


def _administrator_rows(context: _ExcelContext, headers: Sequence[str]) -> list[dict[str, Any]]:
    return _model_rows(
        context,
        context.config.administrators,
        headers,
        {
            "Name": "name",
            "Access Profile": "accprofile",
            "VDOMs": "vdoms",
            "IPv4 Trusted Hosts": "trusthosts",
            "Two Factor": "two_factor",
            "Two Factor Authentication": "two_factor_authentication",
            "Two Factor Notification": "two_factor_notification",
            "Remote Auth": "remote_auth",
            "Remote Group": "remote_group",
            "Credential Configured": lambda item: (
                bool(item.password_configured or item.ssh_key_configured)
            ),
            "FortiToken": "fortitoken",
            "Guest User Groups": "guest_usergroups",
            "Schedule": "schedule",
            "Peer Auth": "peer_auth",
            "Peer Group": "peer_group",
            "SSH Certificate": "ssh_certificate",
            "Additional Settings": "raw_extra",
        },
    )


def _admin_profile_rows(context: _ExcelContext, headers: Sequence[str]) -> list[dict[str, Any]]:
    return _model_rows(
        context,
        context.config.admin_profiles,
        headers,
        {
            "Name": "name",
            "Additional Settings": "raw_extra",
        },
    )


def _admin_permission_rows(context: _ExcelContext, headers: Sequence[str]) -> list[dict[str, Any]]:
    rows = []
    for profile in context.config.admin_profiles:
        for kind, permission in (
            ("Firewall", profile.firewall_permission),
            ("Log", profile.log_permission),
            ("Network", profile.network_permission),
            ("System", profile.system_permission),
            ("UTM", profile.utm_permission),
        ):
            if permission is None:
                continue
            values = permission.model_dump(mode="python", exclude={"raw_extra", "explicit_fields"})
            for setting, value in values.items():
                if value in (None, "", [], {}):
                    continue
                rows.append(
                    {
                        "Profile": profile.name,
                        "Permission Group": kind,
                        "Setting": setting,
                        "Value": value,
                        "Extraction Status": "EXTRACTED",
                        "Additional Settings": _additional_settings(permission),
                    }
                )
                _add_analysis_status(
                    rows[-1],
                    context,
                    vdom="global",
                    names=(profile.name, setting),
                )
    return rows


def _ips_sensor_rows(context: _ExcelContext, headers: Sequence[str]) -> list[dict[str, Any]]:
    return _model_rows(
        context,
        context.config.ips_sensors,
        headers,
        {
            "Name": "name",
            "Description": "comment",
            "Block Malicious URL": "block_malicious_url",
            "Scan Botnet Connections": "scan_botnet_connections",
            "Extended Log": "extended_log",
            "Replacement Message Group": "replacemsg_group",
            "Entry Count": lambda item: len(item.entries),
            "VDOM": "vdom",
        },
    )


def _ips_entry_rows(
    context: _ExcelContext,
    headers: Sequence[str],
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []

    for sensor in context.config.ips_sensors:
        for item in sensor.entries:
            row = {
                "Sensor": sensor.name,
                "Entry ID": item.id,
                "Signature IDs": item.rule,
                "CVEs": item.cve,
                "Applications": item.application,
                "OS": item.os,
                "Protocols": item.protocol,
                "Severities": item.severity,
                "Location": item.location,
                "Default Action": item.default_action,
                "Default Status": item.default_status,
                "Action": item.action,
                "Status": item.status,
                "Log": item.log,
                "Log Packet": item.log_packet,
                "Log Attack Context": item.log_attack_context,
                "Rate Count": item.rate_count,
                "Rate Duration": item.rate_duration,
                "Rate Mode": item.rate_mode,
                "Rate Track": item.rate_track,
                "Quarantine": item.quarantine,
                "Quarantine Expiry": item.quarantine_expiry,
                "Quarantine Log": item.quarantine_log,
                "Vulnerability Types": item.vuln_type,
                "VDOM": sensor.vdom,
                "Additional Settings": _additional_settings(item),
            }

            _overlay_raw(
                row,
                item.raw_extra,
                headers,
            )
            _add_analysis_status(row, context, vdom=sensor.vdom, names=(sensor.name, item.id))

            rows.append(row)

    return rows


def _ips_exempt_rows(context: _ExcelContext, headers: Sequence[str]) -> list[dict[str, Any]]:
    rows = []
    for sensor in context.config.ips_sensors:
        for entry in sensor.entries:
            for item in entry.exempt_ips:
                row = {
                        "Sensor": sensor.name,
                        "Entry ID": entry.id,
                        "Exempt IP ID": item.id,
                        "Source IP": item.src_ip,
                        "Destination IP": item.dst_ip,
                        "VDOM": sensor.vdom,
                        "Additional Settings": _additional_settings(item),
                    }
                _add_analysis_status(row, context, vdom=sensor.vdom, names=(sensor.name, entry.id, item.id))
                rows.append(row)
    return rows


def _security_profile_rows(
    context: _ExcelContext,
    headers: Sequence[str],
) -> list[dict[str, Any]]:
    return _model_rows(
        context,
        context.config.profile_groups,
        headers,
        {
            "Name": "name",
            "Antivirus": "av_profile",
            "IPS Sensor": "ips_sensor",
            "Application List": "application_list",
            "Web Filter": "webfilter_profile",
            "DNS Filter": "dnsfilter_profile",
            "File Filter": "file_filter_profile",
            "SSL/SSH Profile": "ssl_ssh_profile",
            "VDOM": "vdom",
        },
    )


def _external_resource_rows(context: _ExcelContext, headers: Sequence[str]) -> list[dict[str, Any]]:
    return _model_rows(
        context,
        context.config.external_resources,
        headers,
        {
            "Name": "name",
            "Resource": "resource",
            "Type": "type",
            "Refresh Rate": "refresh_rate",
            "Comments": "comments",
            "VDOM": "vdom",
        },
    )


def _policy_source_rows(context: _ExcelContext, headers: Sequence[str]) -> list[dict[str, Any]]:
    rows = []
    policy_names = {str(item.policy_id): item.name for item in context.config.policies if item.policy_id is not None}
    for record in context.source_by_path.get("firewall policy", ()):
        for command in record.commands:
            rows.append(
                {
                    "Source Policy ID": record.object_name,
                    "Policy Name": policy_names.get(str(record.object_name), ""),
                    "Operation": command.operation,
                    "Setting": command.key,
                    "Ordered Source Values": _safe_command_value(
                        command.key,
                        command.values,
                    ),
                }
            )
    return rows


def _interface_source_rows(context: _ExcelContext, headers: Sequence[str]) -> list[dict[str, Any]]:
    rows = []
    for record in context.source_by_path.get("system interface", ()):
        for command in record.commands:
            rows.append(
                {
                    "Interface": record.object_name,
                    "Setting": command.key,
                    "Value": _safe_command_value(
                        command.key,
                        command.values,
                    ),
                    "Extraction Status": "EXTRACTED",
                }
            )
    return rows


def _interface_nested_rows(context: _ExcelContext, headers: Sequence[str]) -> list[dict[str, Any]]:
    rows = []
    for record in context.extracted.source_objects:
        if not record.source_path.startswith("system interface "):
            continue
        interface = record.parent_objects[0] if record.parent_objects else None
        for command in record.commands:
            rows.append(
                {
                    "Interface": interface,
                    "Config Path": record.source_path,
                    "Node Type": "edit" if record.object_name is not None else "config",
                    "Object / Edit": record.object_name,
                    "Operation": command.operation,
                    "Setting": command.key,
                    "Value": _safe_command_value(command.key, command.values),
                    "Extraction Status": "EXTRACTED",
                    "Manual Review": "No",
                }
            )
    return rows


def _unresolved_reference_rows(context: _ExcelContext, headers: Sequence[str]) -> list[dict[str, Any]]:
    rows = []
    for item in context.derived.broken_references:
        rows.append(
            {
                "Source VDOM": item.source_vdom,
                "Source Type": item.source_kind,
                "Source Object": item.source_name,
                "Field": item.source_field,
                "Reference": item.reference,
                "Expected Type": [kind.value for kind in item.expected_kinds],
                "Reason": "Reference could not be resolved in the same VDOM.",
            }
        )
    return rows


def _warning_rows(
    context: _ExcelContext,
    headers: Sequence[str],
) -> list[dict[str, Any]]:
    del headers

    return [
        {
            "ID": index,
            "Category": issue.domain,
            "Message": issue.message,
        }
        for index, issue in enumerate(
            context.validation.warnings,
            start=1,
        )
    ]


def _unsupported_rows(context: _ExcelContext, headers: Sequence[str]) -> list[dict[str, Any]]:
    registered = set(registered_sections())
    counts: dict[str, int] = defaultdict(int)

    for record in context.extracted.source_objects:
        if record.source_path not in registered:
            counts[record.source_path] += 1

    return [
        {
            "Section": path,
            "Object Count": count,
            "Status": "SOURCE_ONLY",
            "Reason": "No dedicated typed FortiGate model is currently defined for this source section.",
            "Raw Capture Location": "FortiGate Source Inventory",
        }
        for path, count in sorted(counts.items())
    ]


def _fortigate_source_inventory_rows(context: _ExcelContext, headers: Sequence[str]) -> list[dict[str, Any]]:
    rows = []
    registered = set(registered_sections())

    for record in context.extracted.source_objects:
        status = "TYPED" if record.source_path in registered else "SOURCE_ONLY"
        base = {
            "Domain": _source_category(record.source_path),
            "VDOM": record.vdom,
            "Scope Type": "Object" if record.object_name is not None else "Config",
            "Source Path": record.source_path,
            "Object": record.object_name,
            "Parent / Subsection": list(record.parent_objects),
            "Extraction Status": status,
        }

        if not record.commands:
            rows.append(base)
            continue

        for command in record.commands:
            rows.append(
                {
                    **base,
                    "Operation": command.operation,
                    "Setting": command.key,
                    "Value": _safe_command_value(command.key, command.values),
                }
            )
    return rows


def _coverage_rows(
    context: _ExcelContext,
    headers: Sequence[str],
) -> list[dict[str, Any]]:
    del headers

    registered = set(
        registered_sections()
    )

    rows: list[dict[str, Any]] = []

    for path in sorted(
        context.source_by_path
    ):
        records = context.source_by_path[path]
        typed = path in registered

        rows.append(
            {
                "Source Section": path,
                "Found": "Yes",
                "Source Objects": len(records),
                "Parsed Objects": len(records),
                "Status": (
                    "TYPED"
                    if typed
                    else "SOURCE_ONLY"
                ),
                "Semantic Level": (
                    "typed-source"
                    if typed
                    else "raw-source"
                ),
                "Parser Handler": "command evaluator",
                "Line Start": min(
                    (
                        record.start_line_number
                        for record in records
                        if record.start_line_number is not None
                    ),
                    default=None,
                ),
                "Line End": max(
                    (
                        record.end_line_number
                        for record in records
                        if record.end_line_number is not None
                    ),
                    default=None,
                ),
                "Semantic Unknowns": (
                    0
                    if typed
                    else len(records)
                ),
                "Notes": (
                    "Typed primitive source extraction available."
                    if typed
                    else (
                        "Explicit source retained without a dedicated "
                        "typed source model."
                    )
                ),
            }
        )

    return rows


def _generic_source_rows(
    sheet_name: str,
    context: _ExcelContext,
    headers: Sequence[str],
) -> list[dict[str, Any]]:
    del headers

    paths = _SOURCE_PATHS_BY_SHEET.get(
        sheet_name,
        (),
    )

    rows: list[dict[str, Any]] = []

    for path in paths:
        for record in context.source_by_path.get(
            path,
            (),
        ):
            values = sanitize_source_attributes(
                record.values
            )

            if not values:
                row = {
                    "VDOM": record.vdom,
                    "Source Path": record.source_path,
                    "Object": record.object_name,
                    "Parent / Subsection": list(
                        record.parent_objects
                    ),
                    "Setting": None,
                    "Value": None,
                }

                _add_analysis_status(
                    row,
                    context,
                    vdom=record.vdom,
                    names=(record.object_name,),
                )

                rows.append(row)
                continue

            for setting, value in values.items():
                row = {
                    "VDOM": record.vdom,
                    "Source Path": record.source_path,
                    "Object": record.object_name,
                    "Parent / Subsection": list(
                        record.parent_objects
                    ),
                    "Setting": setting,
                    "Value": value,
                }

                _add_analysis_status(
                    row,
                    context,
                    vdom=record.vdom,
                    names=(record.object_name,),
                )

                rows.append(row)

    return rows


def _model_rows(
    context: _ExcelContext,
    objects: Iterable[Any],
    headers: Sequence[str],
    mapping: Mapping[str, str | None | Any],
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []

    for item in objects:
        data = item.model_dump(
            mode="python"
        )

        row: dict[str, Any] = {}

        for header, attribute in mapping.items():
            if header not in headers:
                continue

            if callable(attribute):
                row[header] = attribute(item)
            elif attribute:
                row[header] = data.get(
                    attribute
                )

        raw_extra = _additional_settings(item, row.values())

        _overlay_raw(
            row,
            raw_extra,
            headers,
        )

        if "Source Explicit Fields" in headers:
            row["Source Explicit Fields"] = sorted(
                data.get(
                    "explicit_fields",
                    [],
                )
            )

        if "Additional Settings" in headers:
            row["Additional Settings"] = raw_extra

        vdom = str(data.get("vdom") or "")

        name = (
            data.get("name")
            or data.get("policy_id")
            or data.get("seq_num")
            or data.get("id")
        )

        _add_analysis_status(
            row,
            context,
            vdom=vdom,
            names=(name,),
        )

        rows.append(row)

    return rows


def _additional_settings(
    item: Any,
    visible_values: Iterable[Any] = (),
) -> dict[str, Any]:
    data = item.model_dump(mode="python")
    visible = list(visible_values)
    settings = dict(data.get("raw_extra", {}))

    for field in data.get("explicit_fields", ()):
        if field in {"name", "vdom", "raw_extra", "explicit_fields"}:
            continue
        value = data.get(field)
        if value in (None, "", [], {}) or value in visible:
            continue
        settings[field] = value

    return sanitize_source_attributes(settings)


def _add_analysis_status(
    row: dict[str, Any],
    context: _ExcelContext,
    *,
    vdom: str,
    names: Iterable[Any],
    domains: Iterable[str] | None = None,
    extra_reasons: Iterable[str] = (),
) -> None:
    issues = context.issues_for(vdom=vdom, names=names, domains=domains)
    reasons = [issue.message for issue in issues]
    reasons.extend(str(reason) for reason in extra_reasons if reason)
    reasons = list(dict.fromkeys(reasons))

    row["Analysis Status"] = "REVIEW_REQUIRED" if reasons else "EXTRACTED"
    row["Review Reasons"] = reasons
    row["__review__"] = bool(reasons)
    row["__error__"] = any(
        issue.severity.value == "error"
        for issue in issues
    )


def _overlay_raw(
    row: dict[str, Any],
    raw: Mapping[str, Any],
    headers: Sequence[str],
) -> None:
    safe = sanitize_source_attributes(raw)

    for header in headers:
        if row.get(header) not in (None, "", [], {}):
            continue

        value, _ = _lookup_source_header(header, safe)
        if value not in (None, "", [], {}):
            row[header] = value


def _lookup_source_header(
    header: str,
    values: Mapping[str, Any],
) -> tuple[Any, str | None]:
    normalized_values = {
        _normalize_key(key): (key, value)
        for key, value in values.items()
    }

    candidates = [
        *_SOURCE_HEADER_ALIASES.get(
            header,
            (),
        ),
        header,
    ]

    for candidate in candidates:
        normalized = _normalize_key(
            candidate
        )

        if normalized not in normalized_values:
            continue

        source_key, value = normalized_values[
            normalized
        ]

        return value, source_key

    return None, None


def _normalize_key(value: Any) -> str:
    return str(value).strip().lower().replace("-", "_").replace(" ", "_")


def _enabled_text(value: Any) -> str | None:
    if value is None:
        return None
    normalized = str(value).strip().lower()
    if normalized in {"enable", "enabled", "up", "yes", "true", "1"}:
        return "Yes"
    if normalized in {"disable", "disabled", "down", "no", "false", "0"}:
        return "No"
    return str(value)


def _nat_type(value: str | None) -> str | None:
    return {
        "ip_pool": "IP Pool",
        "interface_address": "Interface Address",
    }.get(value, value)


def _disabled_text(value: Any) -> str | None:
    enabled = _enabled_text(value)
    if enabled == "Yes":
        return "No"
    if enabled == "No":
        return "Yes"
    return None


def _interface_source_values(
    context: _ExcelContext,
    *,
    vdom: str,
    name: str,
) -> dict[str, Any]:
    """
    Return safe explicit source values for one interface, including nested
    interface configuration.

    Nested values are exposed both by their raw key when unambiguous and by a
    namespaced key such as ipv6.ip6-address. This is report presentation only;
    FGInterface remains the typed source model.
    """

    values: dict[str, Any] = {}

    for record in context.extracted.source_objects:
        if record.vdom != vdom:
            continue

        direct = (
            record.source_path == "system interface"
            and record.object_name == name
        )

        nested = (
            record.source_path.startswith("system interface ")
            and bool(record.parent_objects)
            and record.parent_objects[0] == name
        )

        if not direct and not nested:
            continue

        suffix = (
            record.source_path.removeprefix("system interface ").strip()
            if nested
            else ""
        )

        for key, value in record.values.items():
            if key not in values:
                values[key] = value

            if suffix:
                values[f"{suffix}.{key}"] = value

    return sanitize_source_attributes(values)


def _source_secret_configured(
    context: _ExcelContext,
    *,
    vdom: str,
    name: str,
    keys: Iterable[str],
) -> str | None:
    wanted = {
        _normalize_key(key)
        for key in keys
    }

    for record in context.extracted.source_objects:
        if record.vdom != vdom:
            continue
        if (
            record.source_path != "system interface"
            or record.object_name != name
        ):
            continue

        present = {
            _normalize_key(key)
            for key in record.values
        }

        if present & wanted:
            return "Yes"

    return None


def _additional_source_settings(
    source_values: Mapping[str, Any],
    row: Mapping[str, Any],
    headers: Sequence[str],
) -> dict[str, Any]:
    consumed: set[str] = set()

    for header in headers:
        if row.get(header) in (None, "", [], {}):
            continue

        _, source_key = _lookup_source_header(
            header,
            source_values,
        )

        if source_key is not None:
            consumed.add(source_key)

    return {
        key: value
        for key, value in source_values.items()
        if key not in consumed
    }


def _source_category(path: str) -> str:
    parts = path.split()
    return " ".join(parts[:2]) if len(parts) >= 2 else path


def _sheet_for_domain(domain: str) -> str:
    return {
        "interface": "Interfaces",
        "interface_topology": "Interfaces",
        "zone": "Zones",
        "address": "Addresses",
        "address6": "Addresses",
        "address_group": "Address Groups",
        "address_group6": "Address Groups",
        "service": "Services",
        "service_group": "Service Groups",
        "policy": "Policies",
        "nat": "NAT Rules",
        "vip": "Virtual IPs",
        "vip_group": "VIP Groups",
        "ip_pool": "IP Pools",
        "ipsec_phase1": "VPN Tunnels",
        "vpn_topology": "VPN Tunnels",
        "vpn_phase2": "VPN Phase 2",
        "static_route": "Routes",
        "user_group": "User Groups",
        "ips_sensor": "IPS Sensors",
    }.get(domain, "FortiGate Source Inventory")


def _safe_command_value(
    key: str,
    values: Iterable[Any],
) -> Any:
    """
    Sanitize one raw CLI command before it reaches Excel.

    Source-command appendix sheets must never bypass the same secret-redaction
    rules used for raw_extra/source inventory.
    """

    normalized_key = str(key).lower().replace("-", "_")
    raw_values = list(values)
    raw_value: Any

    if not raw_values:
        raw_value = True
    elif len(raw_values) == 1:
        raw_value = raw_values[0]
    else:
        raw_value = raw_values

    sanitized = sanitize_source_attributes(
        {
            key: raw_value,
        }
    )

    return sanitized.get(
        normalized_key
    )


def _excel_safe(value: Any) -> Any:
    value = _plain_value(value)

    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value

    return value


def _plain_value(value: Any) -> Any:
    if value is None:
        return None

    if isinstance(value, Enum):
        return value.value

    if is_dataclass(value):
        return _plain_value(asdict(value))

    if isinstance(value, Mapping):
        return "\n".join(
            f"{key} = {_plain_value(item)}"
            for key, item in value.items()
        )

    if isinstance(value, (list, tuple, set, frozenset)):
        return "\n".join(
            str(_plain_value(item))
            for item in value
            if item not in (None, "", [], {})
        )

    if isinstance(value, bool):
        return "Yes" if value else "No"

    if isinstance(value, (str, int, float)):
        return value

    return str(value)
