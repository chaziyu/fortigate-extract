"""Source-vendor visibility and review usability policy for Excel workbooks.

This module intentionally operates only on the rendered workbook. It does not
parse vendor syntax, mutate canonical IR, or affect target generation.
"""

from __future__ import annotations

import io
from typing import Any

from fwmigrate.report.excel_exporter import IRExcelExporter as _BaseIRExcelExporter


_BASE_SHEET_ORDER = tuple(_BaseIRExcelExporter.SHEET_ORDER)


def _without_sheets(order: tuple[str, ...], excluded: frozenset[str]) -> tuple[str, ...]:
    return tuple(sheet_name for sheet_name in order if sheet_name not in excluded)


class VendorAwareIRExcelExporter(_BaseIRExcelExporter):
    """Apply vendor filtering plus a human-oriented workbook review layer."""

    REVIEW_SHEET = "Review Required"

    PALO_ALTO_ONLY_SHEETS = frozenset(
        {
            "Management Service Routes",
            "Security Profile Definitions",
            "Security Profile Rules",
            "Custom URL Categories",
            "GlobalProtect Portals",
            "GlobalProtect Gateways",
            "GlobalProtect Client Auth",
            "GlobalProtect Portal Configs",
            "GlobalProtect External Gateways",
            "GlobalProtect App Settings",
            "GlobalProtect Root CAs",
            "GlobalProtect Gateway Roles",
            "GlobalProtect Tunnel Configs",
            "GlobalProtect Network Gateways",
            "PAN Log Servers",
            "PAN Log Forwarding",
            "PAN Log Forward Matches",
            "PAN DNS Proxies",
            "PAN DNS Proxy Domains",
            "PAN Monitor Profiles",
            "PAN QoS Profiles",
            "PAN QoS Classes",
            "PAN High Availability",
            "PAN HA Monitoring",
            "PAN Device Settings",
            "PAN VSYS Settings",
            "PAN Botnet Report",
            "PAN Custom Reports",
            "PAN SD-WAN Interface Profiles",
            "PAN SD-WAN Link Settings",
            "PAN SD-WAN Path Quality",
            "PAN SD-WAN Traffic Distribution",
            "PAN SD-WAN Rules",
        }
    )

    CISCO_ONLY_SHEETS = frozenset({"Cisco ACP"})
    CHECKPOINT_ONLY_SHEETS = frozenset({"Checkpoint Access Rules"})

    # Keep this deliberately narrow. These sheets are explicitly FortiGate
    # source-detail views by name/implementation; shared canonical sheets stay
    # visible for other vendors.
    FORTIGATE_ONLY_SHEETS = frozenset(
        {
            "FortiGate Source Configuration",
            "Firewall Policy Source Settings",
            "Interface Nested Configuration",
            "VIP Nested Configuration",
        }
    )

    PALO_ALTO_SUMMARY_LABELS = frozenset(
        {
            "Management Service Routes",
            "Security Profile Definitions",
            "Security Profile Rules",
            "Custom URL Categories",
            "GlobalProtect Portals",
            "GlobalProtect Gateways",
            "GlobalProtect Network Gateways",
            "GlobalProtect Client Auth",
            "GlobalProtect Portal Configs",
            "GlobalProtect External Gateways",
            "GlobalProtect App Settings",
            "GlobalProtect Gateway Roles",
            "GlobalProtect Tunnel Configs",
            "PAN SD-WAN Interface Profiles",
            "PAN SD-WAN Link Settings",
            "PAN SD-WAN Path Quality",
            "PAN SD-WAN Traffic Distribution",
            "PAN SD-WAN Rules",
        }
    )

    # These worksheets are retained in the workbook for complete source evidence,
    # but they are hidden by default because they are subordinate/detail views.
    DETAIL_SHEETS = frozenset(
        {
            "Interface Secondary IPs",
            "Interface Source Settings",
            "Interface Nested Configuration",
            "DHCP IP Ranges",
            "DHCP Reservations",
            "Address Group Tags",
            "Firewall Policy Source Settings",
            "VIP Real Servers",
            "VIP Nested Configuration",
            "Session TTL Overrides",
            "SD-WAN Zones",
            "SD-WAN Members",
            "SD-WAN Health Checks",
            "SD-WAN SLAs",
            "SD-WAN Duplication",
            "SD-WAN Neighbors",
            "SD-WAN Rule SLAs",
            "Routing Protocol Settings",
            "Routing Dependencies",
            "Routing Dependency Settings",
            "SSL VPN Authentication Rules",
            "SSL VPN Host Check Items",
            "SSL VPN Portal Split DNS",
            "SSL VPN Portal MAC Rules",
            "SSL VPN Portal OS Checks",
            "SSL VPN Bookmark Groups",
            "SSL VPN Bookmarks",
            "SSL VPN Bookmark Form Data",
            "SSL VPN Landing Pages",
            "SSL VPN Landing Form Data",
            "RADIUS Accounting Servers",
            "FSSO AD Groups",
            "FSSO Polling",
            "User Group Matches",
            "User Group Guests",
            "Admin Profile Permissions",
            "Authentication Sequences",
            "Identity Server Endpoints",
            "GlobalProtect Client Auth",
            "GlobalProtect Portal Configs",
            "GlobalProtect External Gateways",
            "GlobalProtect App Settings",
            "GlobalProtect Root CAs",
            "GlobalProtect Gateway Roles",
            "GlobalProtect Tunnel Configs",
            "PAN Log Forward Matches",
            "PAN DNS Proxy Domains",
            "PAN QoS Classes",
            "PAN HA Monitoring",
            "Security Profile Definitions",
            "Security Profile Rules",
            "Source Security Profile Setting",
            "Security Identity Dependencies",
            "DoS Anomalies",
            "FortiGate Source Configuration",
        }
    )

    # Audit sheets remain visible even when empty because an empty audit result is
    # useful evidence. Other zero-record sheets are hidden rather than deleted so
    # existing consumers can still access the worksheet by name.
    ALWAYS_VISIBLE_SHEETS = frozenset(
        {
            "Summary",
            REVIEW_SHEET,
            "Extraction Coverage",
        }
    )

    # Preserve the columns that identify a record while horizontally scrolling.
    # Unlisted wide tables default to freezing the first identifier column.
    FREEZE_PANES = {
        REVIEW_SHEET: "B4",
        "Interfaces": "C4",
        "Addresses": "C4",
        "Address Groups": "C4",
        "Services": "C4",
        "Service Groups": "C4",
        "Policies": "E4",
        "Checkpoint Access Rules": "E4",
        "NAT Rules": "D4",
        "Routes": "D4",
        "VPN Tunnels": "C4",
        "VPN Phase 2": "C4",
        "LDAP Servers": "C4",
        "RADIUS Servers": "C4",
        "TACACS+ Servers": "C4",
        "SAML Servers": "C4",
        "Local Users": "C4",
        "User Groups": "C4",
        "Administrators": "C4",
        "Security Profiles": "C4",
        "SD-WAN Rules": "D4",
        "Virtual IPs": "C4",
        "IP Pools": "C4",
    }

    _VENDOR_ALIASES = {
        "fortigate": "fortigate",
        "fortinet": "fortigate",
        "fortios": "fortigate",
        "palo_alto": "palo_alto",
        "paloalto": "palo_alto",
        "panos": "palo_alto",
        "pan_os": "palo_alto",
        "cisco_asa": "cisco_asa",
        "cisco asa": "cisco_asa",
        "checkpoint": "checkpoint",
        "check_point": "checkpoint",
        "check point": "checkpoint",
        "juniper_srx": "juniper_srx",
        "juniper srx": "juniper_srx",
        "juniper": "juniper_srx",
    }
    _KNOWN_VENDORS = frozenset(_VENDOR_ALIASES.values())

    # Existing tests and callers historically inspect this class-level constant.
    # FortiGate is the legacy/default source vendor, so expose its complete active
    # order here with Review Required inserted after Summary. Visibility never
    # changes the physical/order contract.
    SHEET_ORDER = (
            "Summary",
            REVIEW_SHEET,
            *(
                sheet_name
                for sheet_name in _without_sheets(
                    _BASE_SHEET_ORDER,
                    PALO_ALTO_ONLY_SHEETS
                    | CISCO_ONLY_SHEETS
                    | CHECKPOINT_ONLY_SHEETS,
                )
                if sheet_name != "Summary"
            ),
    )

    def _source_vendor(self) -> str:
        extraction_vendor = getattr(self.extraction, "source_vendor", None)
        metadata = getattr(self.ir, "metadata", None)
        ir_vendor = getattr(metadata, "source_vendor", None)
        raw_vendor = extraction_vendor or ir_vendor or ""
        normalized = str(raw_vendor).strip().lower().replace("-", "_")
        return self._VENDOR_ALIASES.get(normalized, normalized)

    def _active_sheet_order(self) -> tuple[str, ...]:
        vendor = self._source_vendor()

        # Unknown/legacy vendor identifiers keep the historical full workbook.
        # This avoids hiding source evidence when vendor identity is ambiguous.
        if vendor not in self._KNOWN_VENDORS:
            return _BASE_SHEET_ORDER

        excluded: set[str] = set()
        if vendor != "palo_alto":
            excluded.update(self.PALO_ALTO_ONLY_SHEETS)
        if vendor != "fortigate":
            excluded.update(self.FORTIGATE_ONLY_SHEETS)
        if vendor != "cisco_asa":
            excluded.update(self.CISCO_ONLY_SHEETS)
        if vendor != "checkpoint":
            excluded.update(self.CHECKPOINT_ONLY_SHEETS)

        return tuple(
            sheet_name
            for sheet_name in _BASE_SHEET_ORDER
            if sheet_name not in excluded
        )

    def _workbook_sheet_order(
        self,
        active_order: tuple[str, ...],
    ) -> tuple[str, ...]:
        return (
            "Summary",
            self.REVIEW_SHEET,
            *(
                sheet_name
                for sheet_name in active_order
                if sheet_name != "Summary"
            ),
        )

    def generate(self) -> bytes:
        """Generate the inventory, then apply vendor and review usability policy."""
        # The base exporter currently builds every known worksheet before ordering.
        # Give it the historical complete order so its validation remains unchanged.
        self.SHEET_ORDER = _BASE_SHEET_ORDER
        workbook_bytes = super().generate()

        # Import only after base generation has confirmed the optional reports
        # dependency is available.
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(workbook_bytes))
        active_order = self._active_sheet_order()
        active_sheets = set(active_order)

        # Vendor-inapplicable sheets are still removed. Presentation-only hiding
        # below must never delete applicable worksheets or change their order.
        for worksheet in list(workbook.worksheets):
            if worksheet.title not in active_sheets:
                workbook.remove(worksheet)

        self._apply_review_usability(workbook)

        # Rebuild Summary after filtering/usability processing. Its navigation uses
        # the complete applicable order, including worksheets hidden by presentation
        # policy, so users can unhide and follow the original workbook contract.
        if "Summary" in workbook.sheetnames:
            workbook.remove(workbook["Summary"])

        self.SHEET_ORDER = self._workbook_sheet_order(active_order)
        self._build_summary(workbook)
        self._remove_inapplicable_summary_rows(workbook["Summary"])
        self._add_summary_visibility(workbook["Summary"], workbook)
        self._apply_sheet_view(workbook["Summary"])

        # Reuse the base ordering pass. It preserves the full logical order and
        # reapplies normal tab colors, including the rebuilt Summary sheet.
        self._order_sheets(workbook)

        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()

    def _apply_review_usability(self, workbook: Any) -> None:
        self._build_review_required(workbook)

        for sheet in workbook.worksheets:
            if sheet.title in {"Summary", self.REVIEW_SHEET}:
                continue

            record_count = self._record_count(sheet)
            if sheet.title in self.DETAIL_SHEETS or (
                record_count == 0 and sheet.title not in self.ALWAYS_VISIBLE_SHEETS
            ):
                sheet.sheet_state = "hidden"
            else:
                sheet.sheet_state = "visible"

            self._apply_sheet_view(sheet)

        review_sheet = workbook[self.REVIEW_SHEET]
        review_sheet.sheet_state = "visible"
        self._apply_sheet_view(review_sheet)

    def _apply_sheet_view(self, sheet: Any) -> None:
        sheet.sheet_view.zoomScale = 90
        sheet.sheet_view.zoomScaleNormal = 90

        if sheet.max_column > 8:
            sheet.freeze_panes = self.FREEZE_PANES.get(sheet.title, "B4")
        else:
            sheet.freeze_panes = self.FREEZE_PANES.get(sheet.title, "A4")

    @staticmethod
    def _record_count(sheet: Any) -> int:
        # Normal inventory tables use title, note, header in rows 1-3.
        return max(sheet.max_row - 3, 0)

    @staticmethod
    def _header_map(sheet: Any) -> dict[str, int]:
        return {
            str(sheet.cell(3, column).value or "").strip(): column
            for column in range(1, sheet.max_column + 1)
            if sheet.cell(3, column).value
        }

    @staticmethod
    def _truthy_review(value: Any) -> bool:
        return str(value or "").strip().lower() in {
            "yes",
            "true",
            "1",
            "manual",
            "required",
        }

    @staticmethod
    def _review_status(value: Any) -> bool:
        normalized = str(value or "").strip().upper().replace(" ", "_")
        return normalized in {
            "PARTIALLY_NORMALIZED",
            "UNSUPPORTED",
            "PARSE_ERROR",
            "MANUAL",
            "PARTIAL",
            "UNRESOLVED",
        }

    def _review_rows(self, workbook: Any) -> list[tuple[str, str, str, str, str, int]]:
        rows: list[tuple[str, str, str, str, str, int]] = []

        object_headers = (
            "Name",
            "Rule Name",
            "Object Name",
            "Item",
            "ID",
            "Source ID",
            "Section",
            "Interface",
            "Profile Name",
        )
        reason_headers = (
            "Review Reasons",
            "Review Reason",
            "Reason",
            "Message",
            "Notes",
            "Audit Note",
        )
        status_headers = (
            "Extraction Status",
            "Migration Status",
            "Status",
            "Confidence",
            "Result",
        )

        for sheet in workbook.worksheets:
            if sheet.title in {"Summary", self.REVIEW_SHEET} or sheet.max_row < 4:
                continue

            headers = self._header_map(sheet)
            manual_column = headers.get("Manual Review")
            status_columns = [headers[name] for name in status_headers if name in headers]
            reason_columns = [headers[name] for name in reason_headers if name in headers]
            object_columns = [headers[name] for name in object_headers if name in headers]

            for row_number in range(4, sheet.max_row + 1):
                manual_review = (
                    manual_column is not None
                    and self._truthy_review(sheet.cell(row_number, manual_column).value)
                )

                status_value = ""
                status_requires_review = False
                for column in status_columns:
                    candidate = sheet.cell(row_number, column).value
                    if candidate not in (None, "") and not status_value:
                        status_value = str(candidate)
                    status_requires_review = status_requires_review or self._review_status(candidate)

                if not (manual_review or status_requires_review):
                    continue

                object_value = ""
                for column in object_columns:
                    candidate = sheet.cell(row_number, column).value
                    if candidate not in (None, ""):
                        object_value = str(candidate)
                        break
                if not object_value:
                    object_value = f"Row {row_number}"

                issue = ""
                for column in reason_columns:
                    candidate = sheet.cell(row_number, column).value
                    if candidate not in (None, ""):
                        issue = str(candidate)
                        break
                if not issue:
                    issue = status_value or "Manual review required"

                try:
                    category = self._sheet_category(sheet.title)
                except Exception:
                    category = "Review"

                rows.append(
                    (
                        category,
                        object_value,
                        issue,
                        status_value or ("MANUAL" if manual_review else "REVIEW"),
                        sheet.title,
                        row_number,
                    )
                )

        return rows

    def _build_review_required(self, workbook: Any) -> None:
        if self.REVIEW_SHEET in workbook.sheetnames:
            workbook.remove(workbook[self.REVIEW_SHEET])

        rows = self._review_rows(workbook)
        sheet = self._table_sheet(
            workbook,
            self.REVIEW_SHEET,
            (
                "Category",
                "Object",
                "Issue / Review Reason",
                "Status",
                "Source Sheet",
                "Source Row",
            ),
            rows,
            empty_note="No items currently require manual review.",
            subtitle=(
                "Consolidated manual-review, unsupported, unresolved, partial, "
                "and parse-error findings. Source Sheet links open the detailed evidence."
            ),
        )

        from openpyxl.styles import PatternFill

        for row_number in range(4, sheet.max_row + 1):
            source_sheet = str(sheet.cell(row_number, 5).value or "")
            source_row = sheet.cell(row_number, 6).value
            if source_sheet and source_sheet in workbook.sheetnames and source_row:
                escaped = source_sheet.replace("'", "''")
                sheet.cell(row_number, 5).hyperlink = f"#'{escaped}'!A{source_row}"
                sheet.cell(row_number, 5).style = "Hyperlink"

            status_cell = sheet.cell(row_number, 4)
            status = str(status_cell.value or "").upper()
            if "UNSUPPORTED" in status or "PARSE_ERROR" in status:
                status_cell.fill = PatternFill("solid", fgColor=self._LIGHT_RED)
            elif "PARTIAL" in status or "MANUAL" in status or "UNRESOLVED" in status:
                status_cell.fill = PatternFill("solid", fgColor=self._LIGHT_AMBER)

    def _add_summary_visibility(self, summary: Any, workbook: Any) -> None:
        """Add visibility state without excluding hidden applicable worksheets."""
        from openpyxl.styles import Alignment, Font, PatternFill

        header_row = next(
            (
                row
                for row in range(1, summary.max_row + 1)
                if summary.cell(row, 1).value == "Category"
                and summary.cell(row, 2).value == "Sheet"
            ),
            None,
        )
        if header_row is None:
            return

        title_row = header_row - 1
        summary.unmerge_cells(
            start_row=title_row,
            start_column=1,
            end_row=title_row,
            end_column=5,
        )
        summary.merge_cells(
            start_row=title_row,
            start_column=1,
            end_row=title_row,
            end_column=6,
        )

        header_cell = summary.cell(header_row, 6, "Visibility")
        header_cell.font = Font(
            name="Aptos",
            bold=True,
            color=self._WHITE,
        )
        header_cell.fill = PatternFill("solid", fgColor=self._NAVY)
        header_cell.alignment = Alignment(
            wrap_text=True,
            vertical="center",
        )

        last_row = header_row
        for row in range(header_row + 1, summary.max_row + 1):
            sheet_name = summary.cell(row, 2).value
            if not sheet_name or sheet_name not in workbook.sheetnames:
                break

            last_row = row
            target_sheet = workbook[sheet_name]
            visibility_cell = summary.cell(
                row,
                6,
                "Visible" if target_sheet.sheet_state == "visible" else "Hidden",
            )
            visibility_cell.font = Font(
                name="Aptos",
                size=10,
                color=self._TEXT,
            )
            visibility_cell.alignment = Alignment(
                wrap_text=True,
                vertical="top",
            )

            if (row - header_row) % 2 == 0:
                visibility_cell.fill = PatternFill(
                    "solid",
                    fgColor="F8FAFC",
                )

        summary.auto_filter.ref = f"A{header_row}:F{last_row}"
        summary.column_dimensions["F"].width = 14

    def _remove_inapplicable_summary_rows(self, summary: Any) -> None:
        vendor = self._source_vendor()
        if vendor not in self._KNOWN_VENDORS:
            return

        hidden_labels: set[str] = set()
        if vendor != "palo_alto":
            hidden_labels.update(self.PALO_ALTO_SUMMARY_LABELS)

        if not hidden_labels:
            return

        # Delete bottom-up so row indexes remain stable. Only Inventory Counts
        # labels match this set; navigation uses sheet names in column B.
        for row in range(summary.max_row, 1, -1):
            if summary.cell(row, 1).value in hidden_labels:
                summary.delete_rows(row, 1)
