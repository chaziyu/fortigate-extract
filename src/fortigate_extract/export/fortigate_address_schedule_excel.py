"""Excel projection for audited FortiGate address/schedule corrections."""

from __future__ import annotations

from copy import copy
from typing import Any

from openpyxl.utils import get_column_letter

from fwmigrate.report.excel_readability import ReadableFortiGateExcelExporter


IP_POOL_GROUP_SHEET = "IP Pool Groups"


def _join_values(value: Any) -> Any:
    if isinstance(value, (list, tuple, set)):
        return "\n".join(str(item) for item in value)
    return value


class FortiGateAddressScheduleExcelExporter(ReadableFortiGateExcelExporter):
    """Expose corrected FortiGate address, schedule, and NAT inventory semantics."""

    def _active_sheet_order(self) -> tuple[str, ...]:
        order = list(super()._active_sheet_order())
        if self._source_vendor() not in {"fortigate", "fortinet"}:
            return tuple(
                sheet_name
                for sheet_name in order
                if sheet_name != IP_POOL_GROUP_SHEET
            )

        if IP_POOL_GROUP_SHEET not in order:
            insert_at = (
                order.index("IP Pools") + 1
                if "IP Pools" in order
                else len(order)
            )
            order.insert(insert_at, IP_POOL_GROUP_SHEET)
        return tuple(order)

    @staticmethod
    def _copy_cell_style(source: Any, target: Any) -> None:
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)
        target.number_format = source.number_format

    def _build_addresses(self, workbook: Any) -> None:
        super()._build_addresses(workbook)

        if "Addresses" in workbook.sheetnames:
            sheet = workbook["Addresses"]
            defaults_column = sheet.max_column + 1
            defaults_header = sheet.cell(3, defaults_column, "Effective Defaults")
            if defaults_column > 1:
                self._copy_cell_style(
                    sheet.cell(3, defaults_column - 1),
                    defaults_header,
                )

            for row_number, address in enumerate(self.ir.addresses, start=4):
                defaults_cell = sheet.cell(
                    row_number,
                    defaults_column,
                    self._format_settings(
                        getattr(address, "source_effective_defaults", {}) or {}
                    ),
                )
                if defaults_column > 1:
                    self._copy_cell_style(
                        sheet.cell(row_number, defaults_column - 1),
                        defaults_cell,
                    )

            sheet.column_dimensions[
                get_column_letter(defaults_column)
            ].width = 28
    def _build_ip_pools(self, workbook: Any) -> None:
        super()._build_ip_pools(workbook)
        self._build_ip_pool_groups(workbook)

    def _build_ip_pool_groups(self, workbook: Any) -> None:
        if self._source_vendor() not in {"fortigate", "fortinet"}:
            return
        groups = list(getattr(self.ir, "ip_pool_groups", []) or [])
        rows = [
            (
                item.name,
                item.source_context,
                item.members,
                item.unresolved_members,
                item.source_explicit_fields,
                item.migration_status,
                self._optional_bool_literal(item.requires_manual_review),
                item.review_reasons,
                self._format_settings(item.source_attributes),
            )
            for item in groups
        ]
        self._table_sheet(
            workbook,
            IP_POOL_GROUP_SHEET,
            (
                "Name",
                "Source Context",
                "Members",
                "Unresolved Members",
                "Source Explicit Fields",
                "Migration Status",
                "Manual Review",
                "Review Reasons",
                "Additional Settings",
            ),
            rows,
            empty_note="No FortiGate IP-pool groups were extracted.",
            subtitle=(
                "FortiGate IP-pool group membership is retained for inventory "
                "and NAT correlation; target translation remains withheld."
            ),
        )

    def _build_nat_rules(self, workbook: Any) -> None:
        super()._build_nat_rules(workbook)
        if (
            self._source_vendor() not in {"fortigate", "fortinet"}
            or "NAT Rules" not in workbook.sheetnames
        ):
            return

        sheet = workbook["NAT Rules"]
        column = sheet.max_column + 1
        header = sheet.cell(3, column, "Source Pool Group References")
        if column > 1:
            self._copy_cell_style(sheet.cell(3, column - 1), header)
        for row_number, rule in enumerate(self.ir.nat_rules, start=4):
            cell = sheet.cell(
                row_number,
                column,
                _join_values(
                    getattr(rule, "source_pool_group_references", []) or []
                ),
            )
            if column > 1:
                self._copy_cell_style(sheet.cell(row_number, column - 1), cell)
        sheet.column_dimensions[get_column_letter(column)].width = 28

    def _apply_sheet_view(self, sheet: Any) -> None:
        super()._apply_sheet_view(sheet)
        if sheet.title == "Addresses" and sheet.max_row >= 3:
            headers = {
                str(cell.value or ""): cell.column
                for cell in sheet[3]
            }
            for header in (
                "Effective Defaults",
                "IPv6 Template Reference",
                "Template Reference Resolved",
            ):
                column = headers.get(header)
                if column:
                    sheet.column_dimensions[
                        get_column_letter(column)
                    ].hidden = False
        if sheet.title == "NAT Rules" and sheet.max_row >= 3:
            headers = {
                str(cell.value or ""): cell.column
                for cell in sheet[3]
            }
            column = headers.get("Source Pool Group References")
            if column:
                sheet.column_dimensions[
                    get_column_letter(column)
                ].hidden = False

    def _build_schedule_groups(self, workbook: Any) -> None:
        rows = [
            (
                item.name,
                item.source_context,
                item.members,
                item.unresolved_members,
                item.source_attributes.get("color"),
                item.source_attributes.get("fabric_object"),
                item.migration_status,
                self._optional_bool_literal(item.requires_manual_review),
                self._format_settings({
                    key: value
                    for key, value in item.source_attributes.items()
                    if key not in {"color", "fabric_object"}
                }),
                item.description,
            )
            for item in self.ir.schedule_groups
        ]
        self._table_sheet(
            workbook,
            "Schedule Groups",
            (
                "Name",
                "Source Context",
                "Members",
                "Unresolved Members",
                "Source Color",
                "Source Fabric Object",
                "Migration Status",
                "Manual Review",
                "Additional Settings",
                "Description",
            ),
            rows,
            empty_note="No schedule groups were extracted.",
            subtitle="Ordered FortiGate schedule-group membership and source metadata retained for review.",
        )
