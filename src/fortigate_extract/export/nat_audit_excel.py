"""NAT audit visibility for Excel inventory output.

This report-only layer exposes source mechanism, effective FortiGate central-NAT
mode, and source-preserved Cisco NAT modifiers without changing canonical IR or
target-generation semantics.
"""

from __future__ import annotations

from copy import copy
from typing import Any

from openpyxl.utils import get_column_letter

from fwmigrate.report.excel_optimized import SinglePassIRExcelExporter


class NATAuditIRExcelExporter(SinglePassIRExcelExporter):
    """Expose NAT source authority needed for migration review."""

    NAT_SOURCE_MECHANISM_HEADER = "NAT Source Mechanism"
    EFFECTIVE_CENTRAL_NAT_MODE_HEADER = "Effective Central NAT Mode"
    CISCO_NAT_HEADERS = (
        "DNS Rewrite",
        "No Proxy ARP",
        "Route Lookup",
        "Unidirectional",
        "Net-to-Net",
        "PAT Options",
        "Raw NAT Options",
    )

    @staticmethod
    def _copy_cell_style(source: Any, target: Any) -> None:
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)
        target.number_format = source.number_format

    def _fortigate_central_nat_by_context(self) -> dict[str, Any]:
        return {
            str(getattr(context, "vdom", None) or "root"): getattr(
                context,
                "central_nat",
                None,
            )
            for context in getattr(self.ir, "execution_contexts", []) or []
        }

    def _format_cisco_nat_option(self, value: Any) -> str | None:
        if value in (None, "", [], {}, ()):
            return None
        if isinstance(value, dict):
            return self._format_settings(value)
        if isinstance(value, (list, tuple, set)):
            return " ".join(str(item) for item in value)
        return str(value)

    def _cisco_nat_audit_values(self, rule: Any) -> tuple[Any, ...]:
        attrs = getattr(rule, "source_attributes", {}) or {}
        pat_options = attrs.get("pat_pool_options")
        if pat_options in (None, "", [], {}):
            pat_options = attrs.get("fmc_pat_options")
        raw_options = attrs.get("raw_options")
        return (
            self._optional_bool_literal(attrs.get("dns")),
            self._optional_bool_literal(attrs.get("no_proxy_arp")),
            self._optional_bool_literal(attrs.get("route_lookup")),
            self._optional_bool_literal(attrs.get("unidirectional")),
            self._optional_bool_literal(attrs.get("net_to_net")),
            self._format_cisco_nat_option(pat_options),
            self._format_cisco_nat_option(raw_options),
        )

    def _build_nat_rules(self, workbook: Any) -> None:
        super()._build_nat_rules(workbook)

        sheet_names = tuple(
            getattr(self, "_partitioned_sheet_names", {}).get(
                "NAT Rules",
                ("NAT Rules",),
            )
        )
        sheet_names = tuple(
            name for name in sheet_names if name in workbook.sheetnames
        )
        if not sheet_names:
            return

        source_vendor = str(self._source_vendor() or "").strip().lower()
        is_fortigate = source_vendor in {"fortigate", "fortinet"}
        is_cisco = source_vendor in {"cisco_asa", "cisco_ftd"}
        central_nat_by_context = self._fortigate_central_nat_by_context()
        rule_index = 0

        for sheet_name in sheet_names:
            sheet = workbook[sheet_name]
            headers = {
                str(sheet.cell(3, column).value or "").strip(): column
                for column in range(1, sheet.max_column + 1)
            }
            mechanism_column = headers.get(self.NAT_SOURCE_MECHANISM_HEADER)
            central_mode_column = headers.get(self.EFFECTIVE_CENTRAL_NAT_MODE_HEADER)

            if mechanism_column is None:
                mechanism_column = sheet.max_column + 1
                header = sheet.cell(
                    3,
                    mechanism_column,
                    self.NAT_SOURCE_MECHANISM_HEADER,
                )
                if mechanism_column > 1:
                    self._copy_cell_style(
                        sheet.cell(3, mechanism_column - 1),
                        header,
                    )

            if central_mode_column is None:
                central_mode_column = sheet.max_column + 1
                header = sheet.cell(
                    3,
                    central_mode_column,
                    self.EFFECTIVE_CENTRAL_NAT_MODE_HEADER,
                )
                if central_mode_column > 1:
                    self._copy_cell_style(
                        sheet.cell(3, central_mode_column - 1),
                        header,
                    )

            cisco_columns: list[int] = []
            for header_name in self.CISCO_NAT_HEADERS:
                column = headers.get(header_name)
                if column is None:
                    column = sheet.max_column + 1
                    header = sheet.cell(3, column, header_name)
                    if column > 1:
                        self._copy_cell_style(sheet.cell(3, column - 1), header)
                cisco_columns.append(column)

            for row_number in range(4, sheet.max_row + 1):
                if rule_index >= len(self.ir.nat_rules):
                    break
                rule = self.ir.nat_rules[rule_index]
                rule_index += 1

                mechanism_cell = sheet.cell(
                    row_number,
                    mechanism_column,
                    getattr(rule, "source_origin", None),
                )

                central_mode = None
                if is_fortigate:
                    source_context = str(
                        getattr(rule, "source_context", None) or "root"
                    )
                    configured_mode = central_nat_by_context.get(source_context)
                    central_mode = (
                        str(configured_mode)
                        if configured_mode is not None
                        else "unknown"
                    )
                central_mode_cell = sheet.cell(
                    row_number,
                    central_mode_column,
                    central_mode,
                )

                if mechanism_column > 1:
                    self._copy_cell_style(
                        sheet.cell(row_number, mechanism_column - 1),
                        mechanism_cell,
                    )
                self._copy_cell_style(mechanism_cell, central_mode_cell)

                values = self._cisco_nat_audit_values(rule) if is_cisco else (None,) * len(cisco_columns)
                previous_cell = central_mode_cell
                for column, value in zip(cisco_columns, values):
                    cell = sheet.cell(row_number, column, value)
                    self._copy_cell_style(previous_cell, cell)
                    previous_cell = cell

            sheet.column_dimensions[
                get_column_letter(mechanism_column)
            ].width = 24
            sheet.column_dimensions[
                get_column_letter(central_mode_column)
            ].width = 28
            for column in cisco_columns:
                sheet.column_dimensions[get_column_letter(column)].width = 20

    def _apply_sheet_view(self, sheet: Any) -> None:
        super()._apply_sheet_view(sheet)
        if not sheet.title.startswith("NAT Rules") or sheet.max_row < 3:
            return

        visible_headers = {
            self.NAT_SOURCE_MECHANISM_HEADER,
            self.EFFECTIVE_CENTRAL_NAT_MODE_HEADER,
            *self.CISCO_NAT_HEADERS,
        }
        for column in range(1, sheet.max_column + 1):
            header = str(sheet.cell(3, column).value or "").strip()
            if header not in visible_headers:
                continue
            dimension = sheet.column_dimensions[get_column_letter(column)]
            dimension.hidden = False
            dimension.outlineLevel = 0
            dimension.collapsed = False
