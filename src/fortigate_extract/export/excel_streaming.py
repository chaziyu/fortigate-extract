"""Low-memory FAST XLSX export using openpyxl's write-only workbook."""

from __future__ import annotations

import io
import json
import logging
import time
from dataclasses import dataclass
from itertools import chain
from typing import Any, BinaryIO, Iterable, Sequence

import fwmigrate.report.excel_exporter as _excel_exporter
from fwmigrate.report.excel_audit import (
    ExcelAuditAccumulator,
    build_audit_classifier,
)
from fwmigrate.report.excel_exporter import (
    ExcelExportUnavailableError,
    IRExcelExporter,
)
from fwmigrate.report.excel_optimized import (
    _BUILDER_SPECS,
    _ExcelExportMetrics,
    _log_export_metrics,
    WorksheetExportMetric,
)
from fwmigrate.report.excel_options import ExcelExportProfile
from fwmigrate.report.excel_serialization import save_workbook_with_compression
from fwmigrate.report.fortigate_semantics_excel import FortiGateSemanticsExcelExporter
from fwmigrate.report.nat_audit_excel import NATAuditIRExcelExporter

try:
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover - exercised only without the reports dependency
    WriteOnlyCell = None
    get_column_letter = None


logger = logging.getLogger("fwmigrate.report.excel_optimized")


_STREAM_BUILD_ORDER = tuple(
    _BUILDER_SPECS[builder_name]
    for builder_name in (
    "_build_system_settings",
    "_build_ntp_settings",
    "_build_management_service_routes",
    "_build_interfaces",
    "_build_interface_secondary_ips",
    "_build_interface_source_settings",
    "_build_interface_nested_configuration",
    "_build_dhcp_servers",
    "_build_dhcp_ip_ranges",
    "_build_dhcp_reservations",
    "_build_zones",
    "_build_addresses",
    "_build_address_groups",
    "_build_address_group_tags",
    "_build_proxy_addresses",
    "_build_service_categories",
    "_build_services",
    "_build_service_groups",
    "_build_session_ttl_settings",
    "_build_session_ttl_overrides",
    "_build_schedules",
    "_build_schedule_groups",
    "_build_policies",
    "_build_firewall_filters",
    "_build_checkpoint_access_rule_sheet",
    "_build_cisco_acp",
    "_build_default_security_rules",
    "_build_pbf_rules",
    "_build_local_in_policies",
    "_build_security_policies",
    "_build_multicast_policies",
    "_build_firewall_policy_source_settings",
    "_build_ztna_providers",
    "_build_ip_pools",
    "_build_ipv6_eh_filter",
    "_build_virtual_ips",
    "_build_vip_real_servers",
    "_build_vip_nested_configuration",
    "_build_vip_groups",
    "_build_nat_rules",
    "_build_routes",
    "_build_policy_routes",
    "_build_cisco_pbr",
    "_build_vpn_tunnels",
    "_build_vpn_phase2",
    "_build_ssl_vpn",
    "_build_certificates",
    "_build_ssh_keys",
    "_build_routing_protocols",
    "_build_routing_dependencies",
    "_build_sdwan",
    "_build_security_profiles",
    "_build_security_profile_definitions",
    "_build_security_profile_rules",
    "_build_custom_url_categories",
    "_build_source_security_profiles",
    "_build_fortigate_source_configuration",
    "_build_identity_inventory",
    "_build_user_identity_settings",
    "_build_security_identity_dependencies",
    "_build_administrator_inventory",
    "_build_dos_inventory",
    "_build_firewall_sniffers",
    "_build_authentication_inventory",
    "_build_phase7_identity_sheets",
    "_build_globalprotect_sheets",
    "_build_pan_phase9_sheets",
    "_build_pan_sdwan_sheets",
    "_build_extraction_coverage",
    "_build_unresolved_references",
    )
)


_MANDATORY_EMPTY_SHEETS = frozenset(
    {
        "Summary",
        "Extraction Coverage",
        "Dependency Registry",
        "FortiGate Source Configuration",
        "Interface Nested Configuration",
        "VIP Nested Configuration",
    }
)

FAST_DEFAULT_COLUMN_WIDTH = 18
FAST_COLUMN_WIDTHS = {
    "Name": 28,
    "Migration Status": 22,
    "Extraction Status": 22,
    "Manual Review": 16,
    "Review Reasons": 50,
    "Review Reason": 50,
    "Issue / Review Reason": 50,
    "Source": 32,
    "Destination": 32,
    "Service": 28,
    "Source Attributes": 50,
    "Additional Settings": 50,
    "Description": 50,
    "Notes": 50,
    "Evidence / Note": 50,
    "Field": 28,
    "Value": 50,
    "Sheet": 32,
}


def _append_styled_row(
    sheet: Any,
    values: Sequence[Any],
    styles: dict[str, Any],
    kind: str,
    first_cell_only: bool = False,
) -> None:
    row = list(values)
    styled_count = 1 if first_cell_only else len(row)
    for index in range(styled_count):
        cell = WriteOnlyCell(sheet, value=row[index])
        cell.font = styles[f"{kind}_font"]
        cell.alignment = styles[f"{kind}_alignment"]
        if kind in {"title", "header"}:
            cell.fill = styles[f"{kind}_fill"]
        if kind == "header":
            cell.border = styles["header_border"]
        row[index] = cell
    sheet.append(row)


@dataclass
class _StreamingCell:
    value: Any = None


class _StreamingSheetProxy:
    """Compatibility surface for legacy style-only post-processing loops."""

    def __init__(self, sheet: Any, rows: int, columns: int):
        self._sheet = sheet
        self._rows = rows
        self._columns = columns

    @property
    def title(self) -> str:
        return self._sheet.title

    @property
    def max_row(self) -> int:
        return self._rows + 3

    @property
    def max_column(self) -> int:
        return self._columns

    def cell(self, row: int, column: int, value: Any = None) -> _StreamingCell:
        return _StreamingCell(value)


class StreamingFastExcelExporter(NATAuditIRExcelExporter):
    """Write FAST exports forward without retaining an in-memory workbook."""

    @staticmethod
    def _configure_fast_columns(sheet: Any, headers: Sequence[str]) -> None:
        for column, header in enumerate(headers, 1):
            sheet.column_dimensions[get_column_letter(column)].width = FAST_COLUMN_WIDTHS.get(
                str(header), FAST_DEFAULT_COLUMN_WIDTH
            )

    def _stream_proxy(self, rows: int = 0, columns: int = 0) -> _StreamingSheetProxy:
        return _StreamingSheetProxy(
            getattr(
                self,
                "_stream_last_sheet",
                type("Sheet", (), {"title": ""})(),
            ),
            rows,
            columns,
        )

    def _extra_columns(
        self,
        title: str,
    ) -> tuple[str, ...]:
        vendor = self._source_vendor()
        if title == "Addresses":
            return ("Effective Defaults",)
        if title == "Routes":
            return ("Device Index",)
        if title == "Policies" and vendor == "palo_alto":
            return ("Rule Type",)
        if title != "NAT Rules":
            return ()

        columns: list[str] = []
        if vendor == "palo_alto":
            columns.extend(
                (
                    "Service Matches",
                    "Destination Distribution",
                    "DNS Rewrite",
                    "Device Binding",
                    "Rulebase Position",
                    "Source Order",
                    "Effective Layer",
                    "Effective Rank",
                    "Effective Scope Chain",
                    "Effective Order Complete",
                    "Effective Order by Context",
                )
            )
        if vendor == "fortigate":
            columns.append("Source Pool Group References")
        columns.extend(
            (
                "NAT Source Mechanism",
                "Effective Central NAT Mode",
                "DNS Rewrite",
                "No Proxy ARP",
                "Route Lookup",
                "Unidirectional",
                "Net-to-Net",
                "PAT Options",
                "Raw NAT Options",
            )
        )
        return tuple(columns)

    def _extra_values(self, title: str, index: int) -> tuple[Any, ...]:
        vendor = self._source_vendor()
        if title == "Addresses":
            if index >= len(self.ir.addresses):
                return (None,)
            address = self.ir.addresses[index]
            return (
                self._format_settings(
                    getattr(address, "source_effective_defaults", {}) or {}
                ),
            )
        if title == "Routes":
            route = self.ir.routes[index] if index < len(self.ir.routes) else None
            return (
                dict(getattr(route, "source_attributes", {}) or {}).get("devindex")
                if route is not None
                else None,
            )
        if title == "Policies" and vendor == "palo_alto":
            policy = self.ir.policies[index] if index < len(self.ir.policies) else None
            settings = getattr(policy, "source_extra_settings", {}) or {}
            return (settings.get("pan_rule_type") if settings.get("pan_rule_type_valid") is True else None,)
        if title != "NAT Rules" or index >= len(self.ir.nat_rules):
            return ()

        rule = self.ir.nat_rules[index]
        values: list[Any] = []
        if vendor == "palo_alto":
            attrs = rule.source_attributes
            values.extend(
                (
                    json.dumps(
                        [match.model_dump(mode="json") for match in rule.service_matches],
                        sort_keys=True,
                    )
                    if rule.service_matches
                    else None,
                    rule.destination_translation_distribution.method
                    if rule.destination_translation_distribution
                    else None,
                    self._format_settings(rule.destination_dns_rewrite.model_dump(mode="json"))
                    if rule.destination_dns_rewrite
                    else None,
                    rule.source_device_binding,
                    attrs.get("pan_rulebase_position"),
                    attrs.get("pan_source_rule_index"),
                    attrs.get("effective_policy_layer"),
                    attrs.get("effective_policy_rank"),
                    self._format_settings({"scope": attrs.get("effective_scope_chain", [])}),
                    self._optional_bool_literal(attrs.get("effective_order_complete")),
                    self._format_settings(attrs.get("pan_effective_order_by_context", {}))
                    if attrs.get("pan_effective_order_by_context")
                    else None,
                )
            )
        if vendor == "fortigate":
            references = getattr(rule, "source_pool_group_references", []) or []
            values.append(
                "\n".join(str(item) for item in references)
                if isinstance(references, (list, tuple, set))
                else references
            )
        configured_mode = self._stream_central_nat_by_context.get(
            str(getattr(rule, "source_context", None) or "root")
        )
        values.extend(
            (
                getattr(rule, "source_origin", None),
                str(configured_mode)
                if vendor == "fortigate" and configured_mode is not None
                else ("unknown" if vendor == "fortigate" else None),
            )
        )
        values.extend(self._cisco_nat_audit_values(rule) if vendor in {"cisco_asa", "cisco_ftd"} else (None,) * 7)
        return tuple(values)

    def _table_sheet(
        self,
        workbook: Any,
        title: str,
        headers: Sequence[str],
        rows: Iterable[Sequence[Any]],
        empty_note: str = "No objects were represented in this IR collection.",
        subtitle: str = "Vendor-neutral IR inventory exported before migration optimization.",
    ) -> _StreamingSheetProxy:
        if title not in getattr(self, "_stream_active_sheets", set()):
            return self._stream_proxy(columns=len(headers))

        row_iterator = iter(rows)
        sentinel = object()
        first_row = next(row_iterator, sentinel)
        has_rows = first_row is not sentinel
        if not has_rows and title not in _MANDATORY_EMPTY_SHEETS:
            return self._stream_proxy(columns=len(headers))

        extra_headers = self._extra_columns(title)
        output_headers = tuple(headers) + extra_headers
        row_iterator = (
            iter(())
            if not has_rows
            else chain((first_row,), row_iterator)
        )

        def output_rows() -> Iterable[Sequence[Any]]:
            for index, values in enumerate(row_iterator):
                yield tuple(values) + self._extra_values(title, index)

        styles = self._stream_styles
        sheet_names = [title]
        sheets: list[tuple[Any, int, float, int]] = []
        limit = self._export_options.max_rows_per_sheet or self.MAX_ROWS_PER_DATA_SHEET
        current = workbook.create_sheet(title)
        self._stream_last_sheet = current
        build_started = time.perf_counter()

        def start_sheet(sheet: Any, note: str) -> None:
            width = len(output_headers)
            sheet.sheet_view.showGridLines = False
            self._configure_fast_columns(sheet, output_headers)
            sheet.freeze_panes = "A4"
            _append_styled_row(
                sheet,
                [title] + [""] * (width - 1),
                styles,
                "title",
                first_cell_only=True,
            )
            _append_styled_row(
                sheet,
                [self._safe_value(f"Back to Summary  |  {note}")] + [""] * (width - 1),
                styles,
                "subtitle",
                first_cell_only=True,
            )
            _append_styled_row(
                sheet,
                tuple(output_headers),
                styles,
                "header",
            )

        start_sheet(current, subtitle if has_rows else empty_note)
        classifiers = {}
        category = getattr(self, "_sheet_category", lambda _: "Inventory")
        if has_rows:
            classifiers[title] = build_audit_classifier(
                title,
                output_headers,
                category,
            )
        current_count = 0
        nonempty = 0
        for values in output_rows():
            if current_count >= limit:
                if len(sheet_names) == 1:
                    old_title = current.title
                    current.title = f"{title} 1"
                    sheet_names[0] = current.title
                    self._audit_accumulator.rename_sheet(old_title, current.title)
                    classifiers[current.title] = build_audit_classifier(
                        current.title,
                        output_headers,
                        category,
                    )
                    classifiers.pop(old_title, None)
                sheets.append((current, current_count, build_started, nonempty))
                next_title = f"{title} {len(sheet_names) + 1}"
                sheet_names.append(next_title)
                self._register_partitioned_sheets(title, sheet_names)
                current = workbook.create_sheet(next_title)
                self._stream_last_sheet = current
                start_sheet(current, subtitle)
                classifiers[next_title] = build_audit_classifier(
                    next_title,
                    output_headers,
                    category,
                )
                build_started = time.perf_counter()
                current_count = 0
                nonempty = 0

            current_count += 1
            safe_values = []
            for value in values:
                safe_value = self._safe_value(value)
                safe_values.append(safe_value)
                nonempty += safe_value not in (None, "")
            current.append(safe_values)
            classifier = classifiers[current.title]
            if classifier.can_produce_review or classifier.can_produce_evidence:
                self._audit_accumulator.add_row(
                    current.title,
                    output_headers,
                    safe_values,
                    current_count + 3,
                    category,
                    classifier,
                )

        sheets.append((current, current_count, build_started, nonempty))
        self._register_partitioned_sheets(title, sheet_names)
        metrics = getattr(self, "_last_export_metrics", None)
        for sheet, row_count, started, populated in sheets:
            last_column = get_column_letter(len(output_headers))
            sheet.auto_filter.ref = f"A3:{last_column}{max(3, row_count + 3)}"
            self._stream_sheet_counts[sheet.title] = row_count
            self._stream_sheet_columns[sheet.title] = len(output_headers)
            if metrics is not None:
                metrics.worksheet_metrics = tuple(
                    metric for metric in metrics.worksheet_metrics if metric.name != sheet.title
                ) + (
                    WorksheetExportMetric(
                        name=sheet.title,
                        rows=row_count,
                        columns=len(output_headers),
                        cells=row_count * len(output_headers),
                        nonempty_cells=populated,
                        build_seconds=time.perf_counter() - started,
                        sizing_seconds=0.0,
                    ),
                )
        self._stream_last_sheet = current
        return _StreamingSheetProxy(current, current_count, len(output_headers))

    def _build_addresses(self, workbook: Any) -> None:
        IRExcelExporter._build_addresses(self, workbook)

    def _build_policies(self, workbook: Any) -> None:
        IRExcelExporter._build_policies(self, workbook)

    def _build_nat_rules(self, workbook: Any) -> None:
        IRExcelExporter._build_nat_rules(self, workbook)

    def _build_routes(self, workbook: Any) -> None:
        IRExcelExporter._build_routes(self, workbook)

    def _write_fixed_sheet(
        self,
        sheet: Any,
        title: str,
        headers: Sequence[str],
        rows: Iterable[Sequence[Any]],
    ) -> None:
        width = len(headers)
        styles = self._stream_styles
        sheet.sheet_view.showGridLines = False
        self._configure_fast_columns(sheet, headers)
        sheet.freeze_panes = "A4"
        _append_styled_row(
            sheet,
            [title] + [""] * (width - 1),
            styles,
            "title",
            first_cell_only=True,
        )
        sheet.append([""] * width)
        _append_styled_row(sheet, headers, styles, "header")
        row_count = 0
        for row in rows:
            row_count += 1
            sheet.append([self._safe_value(value) for value in row])
        last_column = get_column_letter(width)
        sheet.auto_filter.ref = f"A3:{last_column}{max(3, row_count + 3)}"

    def _write_summary(self, sheet: Any) -> None:
        styles = self._stream_styles
        sheet.sheet_view.showGridLines = False
        self._configure_fast_columns(sheet, ("Field", "Value", "Sheet", "Rows", "Columns"))
        _append_styled_row(sheet, ["Firewall Source Inventory"], styles, "title")
        _append_styled_row(sheet, ["FAST streaming export"], styles, "subtitle")
        _append_styled_row(sheet, ["Field", "Value"], styles, "header")
        metadata = self.ir.metadata
        for label, value in (
            ("Source Vendor", metadata.source_vendor),
            ("Hostname", metadata.hostname),
            ("Input Type", metadata.input_type),
            ("Source Version", metadata.source_version),
            ("Source Context", metadata.source_context),
            ("Extracted At (UTC)", metadata.migration_timestamp),
        ):
            sheet.append([label, self._safe_value(value)])
        sheet.append([])
        _append_styled_row(sheet, ["Sheet", "Rows", "Columns"], styles, "header")
        for name in self.SHEET_ORDER:
            if name in {"Summary", "Review Required", "Extraction Evidence"}:
                continue
            count = self._stream_sheet_counts.get(name, 0)
            if count or name in _MANDATORY_EMPTY_SHEETS:
                sheet.append([name, count, self._stream_sheet_columns.get(name, 0)])

    def _write_audit_sheets(self, review_sheet: Any, evidence_sheet: Any) -> None:
        review_rows = []
        for category, obj, issue, status, source_sheet, source_row in self._audit_accumulator.review_rows:
            normalized = self._normalized_status(status)
            severity = "Critical" if normalized in {"UNSUPPORTED", "PARSE_ERROR"} else (
                "High" if normalized == "UNRESOLVED" else "Review"
            )
            review_rows.append(
                (severity, category, obj, issue, self._friendly_status(status), source_sheet, source_row)
            )
        self._write_fixed_sheet(
            review_sheet,
            "Review Required",
            ("Severity", "Category", "Object", "Issue / Review Reason", "Status", "Source Sheet", "Source Row"),
            review_rows,
        )
        self._write_fixed_sheet(
            evidence_sheet,
            "Extraction Evidence",
            ("Category", "Object", "Evidence / Note", "Status", "Source Sheet", "Source Row"),
            self._audit_accumulator.evidence_rows,
        )
        self._stream_sheet_counts["Review Required"] = len(review_rows)
        self._stream_sheet_columns["Review Required"] = 7
        self._stream_sheet_counts["Extraction Evidence"] = len(
            self._audit_accumulator.evidence_rows
        )
        self._stream_sheet_columns["Extraction Evidence"] = 6
        metrics = getattr(self, "_last_export_metrics", None)
        if metrics is not None:
            metrics.worksheet_metrics += (
                WorksheetExportMetric(
                    name="Review Required",
                    rows=len(review_rows),
                    columns=7,
                    cells=len(review_rows) * 7,
                    nonempty_cells=sum(
                        value not in (None, "")
                        for row in review_rows
                        for value in row
                    ),
                    build_seconds=0.0,
                    sizing_seconds=0.0,
                ),
                WorksheetExportMetric(
                    name="Extraction Evidence",
                    rows=len(self._audit_accumulator.evidence_rows),
                    columns=6,
                    cells=len(self._audit_accumulator.evidence_rows) * 6,
                    nonempty_cells=sum(
                        value not in (None, "")
                        for row in self._audit_accumulator.evidence_rows
                        for value in row
                    ),
                    build_seconds=0.0,
                    sizing_seconds=0.0,
                ),
            )

    def generate(self) -> bytes:
        output = io.BytesIO()
        self.generate_to(output)
        return output.getvalue()

    def generate_to(self, output: BinaryIO) -> None:
        if _excel_exporter.Workbook is None:
            raise ExcelExportUnavailableError(
                "Excel export requires openpyxl. Install the project with the reports extra."
            )
        if self._export_options.profile is not ExcelExportProfile.FAST:
            raise ValueError("StreamingFastExcelExporter requires the FAST profile")

        self._stream_styles = self._table_styles()
        debug_timings = logger.isEnabledFor(logging.DEBUG)
        metrics = _ExcelExportMetrics() if debug_timings else None
        self._last_export_metrics = metrics
        self._stream_sheet_counts: dict[str, int] = {}
        self._stream_sheet_columns: dict[str, int] = {}
        self._partitioned_sheet_names = {}
        self._audit_accumulator = ExcelAuditAccumulator()
        self._stream_central_nat_by_context = {
            str(getattr(context, "vdom", None) or "root"): getattr(
                context,
                "central_nat",
                None,
            )
            for context in getattr(self.ir, "execution_contexts", []) or []
        }
        active_order = self._active_sheet_order()
        self.SHEET_ORDER = self._workbook_sheet_order(active_order)
        self._stream_active_sheets = set(active_order)

        total_started = time.perf_counter()
        workbook = _excel_exporter.Workbook(write_only=True)
        workbook.properties.title = "Firewall Source Inventory"
        workbook.properties.subject = "Vendor-neutral firewall configuration extraction"
        workbook.properties.creator = "Firewall Migration Tool"
        summary_sheet = workbook.create_sheet("Summary")

        build_started = time.perf_counter()
        for spec in _STREAM_BUILD_ORDER:
            self._build_registered_if_active(
                workbook,
                self._stream_active_sheets,
                spec.method_name,
            )
        if metrics is not None:
            metrics.timings["streaming workbook construction"] = time.perf_counter() - build_started

        self._write_summary(summary_sheet)
        self._order_sheets(workbook)

        if metrics is not None:
            metrics.worksheet_count = len(metrics.worksheet_metrics)
            metrics.rows_written = sum(self._stream_sheet_counts.values())
            metrics.cells_written = sum(
                metric.cells for metric in metrics.worksheet_metrics
            )
            metrics.nonempty_cells = sum(
                metric.nonempty_cells for metric in metrics.worksheet_metrics
            )
            metrics.populated_cells = metrics.nonempty_cells
            metrics.total_rows = sum(
                metric.rows + 3 for metric in metrics.worksheet_metrics
            )
            metrics.largest_worksheets = tuple(
                (metric.name, metric.rows, metric.columns)
                for metric in sorted(
                    metrics.worksheet_metrics,
                    key=lambda item: (item.rows, item.columns),
                    reverse=True,
                )[:5]
            )

        serialization_started = time.perf_counter()
        save_workbook_with_compression(
            workbook,
            output,
            self._export_options.compression_level
            if self._export_options.compression_level is not None
            else 1,
        )
        if metrics is not None:
            metrics.timings["final XLSX serialization"] = time.perf_counter() - serialization_started
            metrics.output_bytes = output.tell()
            _log_export_metrics(metrics, time.perf_counter() - total_started)
        self._stream_styles = None
