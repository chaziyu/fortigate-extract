"""Vendor-neutral Excel inventory export for :class:`IRConfig`."""

from __future__ import annotations

import io
import json
import re
import time
from datetime import datetime, timezone
from enum import Enum
from typing import Any, Iterable, Sequence

from fwmigrate.extraction.sanitize import sanitize_source_attributes, sanitize_source_value
from fwmigrate.ir.address import IRAddress
from fwmigrate.ir import IRConfig
from fwmigrate.ir.enums import AddressType
from fwmigrate.ir.enums import MigrationConfidence
from fwmigrate.parsers.fortigate.coverage import (
    fortigate_semantic_support_level,
    fortigate_source_category,
)
from fwmigrate.report.excel_audit import (
    AuditSheetClassifier,
    build_audit_classifier,
)
from fwmigrate.report.excel_options import ExcelExportOptions, ExcelExportProfile

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover - exercised only without the reports dependency
    Workbook = None


XLSX_MIMETYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_FORMULA_PREFIXES = ("=", "+", "-", "@")
_ILLEGAL_XML_CHARS = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")
_MAX_CELL_TEXT = 32767
_REMOVED_SHEETS = frozenset({"Review Required", "Extraction Evidence"})


class ExcelExportUnavailableError(RuntimeError):
    """Raised when the optional workbook dependency is unavailable."""


class IRExcelExporter:
    """Render a source firewall inventory directly from vendor-neutral IR."""

    def _vendor_extension_value(self, field: str) -> Any:
        vendor = self.ir.metadata.source_vendor.casefold()
        containers = (
            (self.ir.vendor_extensions.panos, self.ir.vendor_extensions.checkpoint, self.ir.vendor_extensions.fortios)
            if vendor in {"palo_alto", "panos", "palo-alto"}
            else (self.ir.vendor_extensions.checkpoint, self.ir.vendor_extensions.fortios, self.ir.vendor_extensions.panos)
            if vendor in {"checkpoint", "check_point", "check-point"}
            else (self.ir.vendor_extensions.fortios, self.ir.vendor_extensions.checkpoint, self.ir.vendor_extensions.panos)
        )
        fallback = None
        for container in containers:
            value = getattr(container, field, None)
            if value is None:
                continue
            if fallback is None:
                fallback = value
            if not isinstance(value, (list, tuple, dict, set)) or value:
                return value
        return fallback

    @staticmethod
    def _object_extension_value(item: Any, field: str) -> Any:
        extension = getattr(item, "vendor_extension", None)
        return getattr(extension, field, None) if extension is not None else None

    OVERVIEW_SHEETS = (
        "Summary",
    )

    CORE_INVENTORY_SHEETS = (
        "System Settings",
        "DNS Settings",
        "NTP Settings",
        "Management Service Routes",
        "Interfaces",
        "Interface Secondary IPs",
        "Zones",
        "Addresses",
        "Address Groups",
        "Address Group Tags",
        "Service Categories",
        "Services",
        "Service Groups",
        "Schedules",
        "Schedule Groups",
        "Policies",
        "Firewall Filters",
        "Checkpoint Access Rules",
        "Cisco ACP",
        "Default Security Rules",
        "PBF Rules",
        "Local-In Policies",
        "NGFW Security Policies",
        "Multicast Policies",
        "IP Pools",
        "Virtual IPs",
        "VIP Real Servers",
        "VIP Nested Configuration",
        "VIP Groups",
        "NAT Rules",
        "Routes",
        "Policy Routes",
        "Cisco PBR",
        "VPN Tunnels",
        "VPN Phase 2",
    )

    NETWORK_ACCESS_SHEETS = (
        "DHCP Servers",
        "DHCP IP Ranges",
        "DHCP Reservations",
        "Session TTL Settings",
        "Session TTL Overrides",
        "SD-WAN",
        "SD-WAN Zones",
        "SD-WAN Members",
        "SD-WAN Health Checks",
        "SD-WAN SLAs",
        "SD-WAN Rules",
        "SD-WAN Duplication",
        "SD-WAN Neighbors",
        "SD-WAN Rule SLAs",
        "Routing Protocols",
        "Routing Protocol Settings",
        "Routing Dependencies",
        "Routing Dependency Settings",
        "SSL VPN Settings",
        "SSL VPN Portals",
        "SSL VPN Authentication Rules",
        "SSL VPN Host Checks",
        "SSL VPN Host Check Items",
        "SSL VPN Portal Split DNS",
        "SSL VPN Portal MAC Rules",
        "SSL VPN Portal OS Checks",
        "SSL VPN Bookmark Groups",
        "SSL VPN Bookmarks",
        "SSL VPN Bookmark Form Data",
        "SSL VPN Landing Pages",
        "SSL VPN Landing Form Data",
        "LDAP Servers",
        "RADIUS Servers",
        "RADIUS Accounting Servers",
        "TACACS+ Servers",
        "SAML Servers",
        "FSSO Servers",
        "FSSO AD Groups",
        "FSSO Polling",
        "Local Users",
        "User Groups",
        "User Group Matches",
        "User Group Guests",
        "User Authentication Settings",
        "User Quarantine",
        "Security Identity Dependencies",
        "Administrators",
        "Admin Profiles",
        "Admin Profile Permissions",
        "ZTNA Providers",
        "Authentication Schemes",
        "Authentication Rules",
        "Authentication Sequences",
        "Identity Server Endpoints",
        "Certificates",
        "SSL TLS Service Profiles",
        "GlobalProtect Portals",
        "GlobalProtect Gateways",
        "GlobalProtect Client Auth",
        "GlobalProtect Portal Configs",
        "GlobalProtect External Gateways",
        "GlobalProtect App Settings",
        "GlobalProtect Root CAs",
        "GlobalProtect Gateway Roles",
        "GlobalProtect Tunnel Configs",
        "GlobalProtect Network Gateways",
        "PAN Log Servers",
        "PAN Log Forwarding",
        "PAN Log Forward Matches",
        "PAN DNS Proxies",
        "PAN DNS Proxy Domains",
        "PAN Monitor Profiles",
        "PAN QoS Profiles",
        "PAN QoS Classes",
        "PAN High Availability",
        "PAN HA Monitoring",
        "PAN Device Settings",
        "PAN VSYS Settings",
        "PAN Botnet Report",
        "PAN Custom Reports",
        "PAN SD-WAN Interface Profiles",
        "PAN SD-WAN Link Settings",
        "PAN SD-WAN Path Quality",
        "PAN SD-WAN Traffic Distribution",
        "PAN SD-WAN Rules",
    )

    SOURCE_DETAIL_SHEETS = (
        "FortiGate Source Configuration",
        "Firewall Policy Source Settings",
        "Interface Source Settings",
        "Interface Nested Configuration",
        "Proxy Addresses",
        "SSH Keys",
        "Security Profiles",
        "Security Profile Definitions",
        "Security Profile Rules",
        "Custom URL Categories",
        "Source Security Profiles",
        "Source Security Profile Setting",
        "DoS Policies",
        "DoS Anomalies",
        "Firewall Sniffer",
        "IPv6 EH Filter",
    )

    AUDIT_SHEETS = ("Dependency Registry", "Extraction Coverage")

    SHEET_ORDER = (
        OVERVIEW_SHEETS
        + CORE_INVENTORY_SHEETS
        + NETWORK_ACCESS_SHEETS
        + SOURCE_DETAIL_SHEETS
        + AUDIT_SHEETS
    )
    # Excel's 1,048,576 row limit minus title, subtitle, and header rows.
    MAX_ROWS_PER_DATA_SHEET = 1_048_573
    LARGE_TABLE_ROW_THRESHOLD = 1000
    COLUMN_WIDTH_SAMPLE_ROWS = 200

    _NAVY = "17324D"
    _TEAL = "0F766E"
    _LIGHT_TEAL = "D7F0EC"
    _LIGHT_BLUE = "E8F0F7"
    _LIGHT_AMBER = "FEF3C7"
    _LIGHT_RED = "FEE2E2"
    _WHITE = "FFFFFF"
    _TEXT = "1F2937"
    _MUTED = "64748B"
    _BORDER = "CBD5E1"

    _FORTIGATE_ADDRESS_SECTIONS = frozenset({
        "firewall address",
        "firewall address6",
    })

    _FORTIGATE_DEDICATED_INVENTORY_PATHS = {
        "system global",
        "system dns",
        "system interface",
        "system zone",
        "system dhcp server",
        "system session-helper",
        "system session-ttl",
        "system sdwan",
        "endpoint-control fctems",
        "firewall address",
        "firewall address6",
        "firewall multicast-address",
        "firewall multicast-address6",
        "firewall addrgrp",
        "firewall wildcard-fqdn custom",
        "firewall service category",
        "firewall service custom",
        "firewall service group",
        "firewall schedule recurring",
        "firewall schedule onetime",
        "firewall schedule group",
        "firewall shaper traffic-shaper",
        "firewall proxy-address",
        "web-proxy global",
        "firewall policy",
        "firewall ippool",
        "firewall ipv6-eh-filter",
        "firewall vip",
        "firewall vip6",
        "firewall vipgrp",
        "firewall internet-service-name",
        "firewall internet-service-definition",
        "firewall internet-service-addition",
        "firewall internet-service-append",
        "firewall internet-service-custom",
        "firewall internet-service-custom-group",
        "firewall internet-service-extension",
        "firewall internet-service-group",
        "firewall DoS-policy",
        "firewall sniffer",
        "firewall ssh local-key",
        "firewall ssh local-ca",
        "router static",
        "router static6",
        "router rip",
        "router ripng",
        "router ospf",
        "router ospf6",
        "router bgp",
        "router isis",
        "router multicast",
        "vpn ipsec phase1-interface",
        "vpn ipsec phase2-interface",
        "vpn certificate remote",
        "vpn certificate local",
        "vpn certificate ca",
        "vpn ssl web portal",
        "vpn ssl settings",
        "ips sensor",
        "user ldap",
        "user fsso",
        "user adgrp",
        "user saml",
        "user local",
        "user group",
        "system admin",
        "system accprofile",
        "user fortitoken",
        "authentication scheme",
        "authentication rule",
        "user setting",
        "user quarantine",
    }

    def __init__(
        self,
        ir_config: IRConfig,
        extraction_result: Any = None,
        options: ExcelExportOptions | None = None,
    ):
        self.ir = ir_config
        self.extraction = extraction_result
        self._export_options = options or ExcelExportOptions()
        self._partitioned_sheet_names: dict[str, tuple[str, ...]] = {}
        self._audit_accumulator = None

    def generate(self) -> bytes:
        """Generate a complete ``.xlsx`` workbook and return its bytes."""
        if Workbook is None:
            raise ExcelExportUnavailableError(
                "Excel export requires openpyxl. Install the project with the reports extra."
            )

        workbook = Workbook()
        workbook.remove(workbook.active)

        workbook.properties.title = "Firewall Source Inventory"
        workbook.properties.subject = "Vendor-neutral firewall configuration extraction"
        workbook.properties.creator = "Firewall Migration Tool"

        # Build inventory sheets first.
        #
        # Summary is intentionally built last so its navigation section can derive
        # actual worksheet record counts and hyperlinks from the finished workbook.
        self._build_system_settings(workbook)
        self._build_ntp_settings(workbook)
        self._build_management_service_routes(workbook)

        self._build_interfaces(workbook)
        self._build_interface_secondary_ips(workbook)
        self._build_interface_source_settings(workbook)
        self._build_interface_nested_configuration(workbook)

        self._build_dhcp_servers(workbook)
        self._build_dhcp_ip_ranges(workbook)
        self._build_dhcp_reservations(workbook)

        self._build_zones(workbook)

        self._build_addresses(workbook)
        self._build_address_groups(workbook)
        self._build_address_group_tags(workbook)
        self._build_proxy_addresses(workbook)

        self._build_service_categories(workbook)
        self._build_services(workbook)
        self._build_service_groups(workbook)
        self._build_session_ttl_settings(workbook)
        self._build_session_ttl_overrides(workbook)

        self._build_schedules(workbook)
        self._build_schedule_groups(workbook)
        self._build_policies(workbook)
        self._build_firewall_filters(workbook)
        self._build_checkpoint_access_rule_sheet(workbook)
        self._build_cisco_acp(workbook)
        self._build_default_security_rules(workbook)
        self._build_local_in_policies(workbook)
        self._build_security_policies(workbook)
        self._build_multicast_policies(workbook)
        self._build_firewall_policy_source_settings(workbook)
        self._build_ztna_providers(workbook)

        self._build_ip_pools(workbook)
        self._build_ipv6_eh_filter(workbook)
        self._build_virtual_ips(workbook)
        self._build_vip_real_servers(workbook)
        self._build_vip_nested_configuration(workbook)
        self._build_vip_groups(workbook)
        self._build_nat_rules(workbook)
        self._build_pbf_rules(workbook)

        self._build_vpn_tunnels(workbook)
        self._build_vpn_phase2(workbook)
        self._build_ssl_vpn(workbook)
        self._build_certificates(workbook)
        self._build_ssh_keys(workbook)

        self._build_routes(workbook)
        self._build_policy_routes(workbook)
        self._build_cisco_pbr(workbook)
        self._build_routing_protocols(workbook)
        self._build_routing_dependencies(workbook)
        self._build_sdwan(workbook)


        self._build_security_profiles(workbook)
        self._build_security_profile_definitions(workbook)
        self._build_security_profile_rules(workbook)
        self._build_custom_url_categories(workbook)
        self._build_source_security_profiles(workbook)
        self._build_fortigate_source_configuration(workbook)

        self._build_identity_inventory(workbook)
        self._build_user_identity_settings(workbook)
        self._build_security_identity_dependencies(workbook)
        self._build_administrator_inventory(workbook)
        self._build_dos_inventory(workbook)
        self._build_firewall_sniffers(workbook)
        self._build_authentication_inventory(workbook)
        self._build_phase7_identity_sheets(workbook)
        self._build_globalprotect_sheets(workbook)
        self._build_pan_phase9_sheets(workbook)
        self._build_pan_sdwan_sheets(workbook)

        self._build_extraction_coverage(workbook)
        self._build_unresolved_references(workbook)

        # Summary is generated after all inventory sheets so navigation can use
        # actual generated worksheet counts.
        self._build_summary(workbook)

        self._order_sheets(workbook)

        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()

    @staticmethod
    def _format_source_command_values(
        key: str,
        values: Sequence[Any],
    ) -> str:
        values = sanitize_source_value(key, list(values))
        if not values:
            return ""

        if len(values) == 1:
            return str(values[0])

        return json.dumps(
            list(values),
            ensure_ascii=False,
        )

    def _interface_nested_config_rows(
        self,
    ) -> Iterable[tuple[Any, ...]]:
        def walk(
            interface_name: str,
            node: Any,
            parent_path: list[str],
        ) -> Iterable[tuple[Any, ...]]:
            if node.node_type == "config":
                config_path = [
                    *parent_path,
                    str(node.name),
                ]
                object_name = None
            else:
                config_path = list(
                    parent_path
                )
                object_name = str(node.name)

            if node.commands:
                for command in node.commands:
                    yield (
                        interface_name,
                        " / ".join(
                            config_path
                        ),
                        node.node_type,
                        object_name,
                        command.operation,
                        command.key,
                        self._format_source_command_values(
                            command.key,
                            command.values
                        ),
                        "EXTRACT_ONLY",
                        "Yes",
                    )

            elif not node.children:
                # Preserve the existence of an empty
                # nested config/edit block.
                yield (
                    interface_name,
                    " / ".join(
                        config_path
                    ),
                    node.node_type,
                    object_name,
                    None,
                    None,
                    None,
                    "EXTRACT_ONLY",
                    "Yes",
                )

            for child in node.children:
                child_parent = (
                    config_path
                    if node.node_type == "config"
                    else [
                        *config_path,
                        str(node.name),
                    ]
                )

                yield from walk(
                    interface_name,
                    child,
                    child_parent,
                )

        for interface in self.ir.interfaces:
            for root in (
                interface.nested_source_configs
            ):
                yield from walk(
                    interface.name,
                    root,
                    [],
                )

    def _build_interface_nested_configuration(
        self,
        workbook: Any,
    ) -> None:
        self._table_sheet(
            workbook,
            "Interface Nested Configuration",
            (
                "Interface",
                "Config Path",
                "Node Type",
                "Object / Edit",
                "Operation",
                "Setting",
                "Value",
                "Extraction Status",
                "Manual Review",
            ),
            self._interface_nested_config_rows(),
            empty_note=(
                "No nested interface configuration "
                "was extracted from the source firewall."
            ),
            subtitle=(
                "Nested FortiGate interface configuration "
                "retained as sanitized extraction-only "
                "source data. These settings are not "
                "consumed by target generators."
            ),
        )

    @classmethod
    def _has_dedicated_fortigate_inventory(cls, source_path: str) -> bool:
        if source_path.startswith("system sdwan"):
            return source_path in {
                "system sdwan",
                "system sdwan zone",
                "system sdwan members",
                "system sdwan health-check",
                "system sdwan health-check sla",
                "system sdwan service",
                "system sdwan service sla",
                "system sdwan duplication",
                "system sdwan neighbor",
            }
        return any(
            source_path == path or source_path.startswith(f"{path} ")
            for path in cls._FORTIGATE_DEDICATED_INVENTORY_PATHS
        )

    @classmethod
    def _is_source_address_backed_group(cls, item: Any) -> bool:
        return item.source_section in cls._FORTIGATE_ADDRESS_SECTIONS

    def _fortigate_source_inventory_items(self) -> list[Any]:
        if self.extraction is None:
            return []
        if str(self.ir.metadata.source_vendor).lower() not in {"fortigate", "fortinet"}:
            return []
        return [
            item
            for item in self.extraction.inventory_items
            if not self._has_dedicated_fortigate_inventory(item.source_path)
            and "structured-security-profile" not in item.notes
            and "structured-routing-protocol" not in item.notes
        ]

    @staticmethod
    def _flatten_fortigate_source_item(item: Any) -> Iterable[tuple[Any, ...]]:
        nested = {
            "nested-source-config",
            "interface-nested-config",
        }.intersection(item.notes)
        source_path = (
            item.source_path.rsplit(" ", 1)[0]
            if nested and " " in item.source_path
            else item.source_path
        )

        def walk(node: Any, hierarchy: list[str]) -> Iterable[tuple[Any, ...]]:
            for command in node.commands:
                yield (
                    fortigate_source_category(source_path),
                    source_path,
                    item.name,
                    item.source_id,
                    " / ".join(hierarchy),
                    command.operation,
                    command.key,
                    command.values,
                    item.status,
                    "Yes" if item.requires_manual_review else "No",
                )
            for child in node.children:
                yield from walk(child, [*hierarchy, str(child.name)])

        hierarchy = []
        if nested:
            hierarchy.append(str(item.source_path.rsplit(" ", 1)[-1]))
        yield from walk(item, hierarchy)

    def _build_fortigate_source_configuration(self, workbook: Any) -> None:
        items = self._fortigate_source_inventory_items()
        self._table_sheet(
            workbook,
            "FortiGate Source Configuration",
            (
                "Category", "Source Path", "Object", "Source ID",
                "Parent / Subsection", "Operation", "Setting", "Value",
                "Migration Status", "Manual Review",
            ),
            (
                row
                for item in items
                for row in self._flatten_fortigate_source_item(item)
            ),
            empty_note="No fallback FortiGate source configuration was retained.",
            subtitle=(
                "Sanitized source-only FortiGate configuration retained outside "
                "canonical migration IR. Dedicated inventory sections are omitted."
            ),
        )

    def _build_system_settings(self, workbook: Any) -> None:
        settings = self.ir.system_settings
        management = settings.management_plane if settings is not None else None
        self._table_sheet(
            workbook,
            "System Settings",
            (
                "Hostname", "Timezone", "Admin HTTPS Port", "Additional Settings",
                "Management IPv4 Address",
                "Management Netmask", "Management Default Gateway", "Management Address Type",
                "Management IPv6 Address", "Management IPv6 Default Gateway",
                "Management IPv6 Enabled", "Management IPv6 Address Type",
                "Management IPv6 Gateway Type", "Explicit Management Services",
                "System Permitted IPs",
            ),
            [] if settings is None else [(
                settings.hostname,
                settings.timezone,
                settings.admin_https_port,
                self._format_settings(settings.source_attributes),
                management.ipv4_address if management else None,
                management.netmask if management else None,
                management.default_gateway if management else None,
                management.address_type if management else None,
                management.ipv6_address if management else None,
                management.ipv6_default_gateway if management else None,
                management.ipv6_enabled if management else None,
                management.ipv6_address_type if management else None,
                management.ipv6_gateway_type if management else None,
                self._format_settings(management.services) if management and management.services else "",
                management.permitted_ips if management else [],
            )],
        )

        dns = self.ir.dns_settings
        self._table_sheet(
            workbook,
            "DNS Settings",
            ("Primary DNS", "Secondary DNS", "Additional Settings"),
            [] if dns is None else [(
                dns.primary,
                dns.secondary,
                self._format_settings(dns.source_attributes),
            )],
        )

    def _build_ntp_settings(self, workbook: Any) -> None:
        ntp = self.ir.ntp_settings
        self._table_sheet(
            workbook,
            "NTP Settings",
            (
                "Role", "Server Address", "Authentication Type", "Extraction Status",
                "Manual Review", "Additional Settings",
            ),
            [] if ntp is None else [(
                server.role,
                server.address,
                server.authentication_type,
                ntp.migration_status,
                self._optional_bool_literal(ntp.requires_manual_review),
                self._format_settings(server.source_attributes),
            ) for server in ntp.servers],
        )

    def _build_management_service_routes(self, workbook: Any) -> None:
        self._table_sheet(
            workbook,
            "Management Service Routes",
            (
                "Name", "Source Context", "Source Address", "Source Interface",
                "Extraction Status", "Manual Review", "Review Reasons", "Additional Settings",
            ),
            (
                (
                    route.name,
                    route.source_context,
                    route.source_address,
                    route.source_interface,
                    route.migration_status,
                    self._optional_bool_literal(route.requires_manual_review),
                    route.review_reasons,
                    self._format_settings(route.source_attributes),
                )
                for route in self.ir.management_service_routes
            ),
        )

    def _build_summary(self, workbook: Any) -> None:
        pan_ext = self.ir.vendor_extensions.panos
        forti_ext = self.ir.vendor_extensions.fortios
        sheet = workbook.create_sheet("Summary")
        sheet.sheet_view.showGridLines = False

        sheet.merge_cells("A1:E1")
        sheet["A1"] = "Firewall Source Inventory"
        sheet["A1"].font = Font(
            name="Aptos Display",
            size=20,
            bold=True,
            color=self._WHITE,
        )
        sheet["A1"].fill = PatternFill(
            "solid",
            fgColor=self._NAVY,
        )
        sheet["A1"].alignment = Alignment(
            vertical="center",
        )
        sheet.row_dimensions[1].height = 34

        sheet.merge_cells("A2:E2")
        sheet["A2"] = (
            "Vendor-neutral extraction generated before migration optimization"
        )
        sheet["A2"].font = Font(
            name="Aptos",
            size=10,
            italic=True,
            color=self._MUTED,
        )
        sheet["A2"].alignment = Alignment(
            vertical="center",
        )
        sheet.row_dimensions[2].height = 22

        if self.extraction is not None:
            unsupported_count = (
                len(self.extraction.unsupported_items)
                + sum(
                    1
                    for entry in self.ir.audit_entries
                    if entry.confidence == MigrationConfidence.UNSUPPORTED
                )
            )
        else:
            unsupported_count = sum(
                1
                for entry in self.ir.audit_entries
                if entry.confidence == MigrationConfidence.UNSUPPORTED
            )

        if unsupported_count:
            extraction_status = "COMPLETE_WITH_UNSUPPORTED_ITEMS"
        elif self.ir.audit_entries:
            extraction_status = "COMPLETE_WITH_WARNINGS"
        else:
            extraction_status = "COMPLETE"

        metadata_rows = [
            ("Source Vendor", self.ir.metadata.source_vendor),
            ("Hostname", self.ir.metadata.hostname),
            ("Input Type", self.ir.metadata.input_type),
            ("Source Version", self.ir.metadata.source_version),
            ("Source Context", self.ir.metadata.source_context),
            ("Extracted At (UTC)", self.ir.metadata.migration_timestamp),
            ("Extraction Status", extraction_status),
        ]

        self._summary_section(
            sheet,
            4,
            "Extraction Metadata",
            metadata_rows,
        )

        navigation_end_row = self._build_summary_navigation(
            sheet,
            workbook,
            start_row=14,
        )

        inventory_rows = [
            ("Interfaces", len(self.ir.interfaces)),
            ("System Settings", 1 if self.ir.system_settings is not None else 0),
            ("DNS Settings", 1 if self.ir.dns_settings is not None else 0),
            ("NTP Servers", len(self.ir.ntp_settings.servers) if self.ir.ntp_settings else 0),
            ("Management Service Routes", len(self.ir.management_service_routes)),
            (
                "Interface Secondary IPs",
                sum(
                    len(intf.secondary_ips) + len(intf.inactive_secondary_ips)
                    for intf in self.ir.interfaces
                ),
            ),
            ("DHCP Servers", len(self.ir.dhcp_servers)),
            (
                "DHCP IP Ranges",
                sum(
                    len(server.ip_ranges)
                    for server in self.ir.dhcp_servers
                ),
            ),
            (
                "DHCP Reservations",
                sum(
                    len(server.reservations)
                    for server in self.ir.dhcp_servers
                ),
            ),
            ("Zones", len(self.ir.zones)),
            (
                "Addresses",
                len(self.ir.addresses)
                + sum(
                    self._is_source_address_backed_group(item)
                    for item in self.ir.address_groups
                ),
            ),
            (
                "Address Groups",
                sum(
                    not self._is_source_address_backed_group(item)
                    for item in self.ir.address_groups
                ),
            ),
            ("Proxy Addresses", len(self.ir.proxy_addresses)),
            ("Service Categories", len(self.ir.service_categories)),
            ("Services", len(self.ir.services)),
            ("Service Groups", len(self.ir.service_groups)),
            (
                "Session TTL Settings",
                1 if forti_ext.session_ttl_settings is not None else 0,
            ),
            (
                "Session TTL Overrides",
                len(forti_ext.session_ttl_overrides),
            ),
            ("Schedules", len(self.ir.schedules)),
            ("Policies", len(self.ir.policies)),
            ("Multicast Policies", len(self.ir.multicast_policies)),
            ("ZTNA Providers", len(self.ir.ztna_providers)),
            ("IP Pools", len(self.ir.ip_pools)),
            ("Virtual IPs", len(self.ir.virtual_ips)),
            (
                "VIP Real Servers",
                sum(
                    len(vip.real_servers)
                    for vip in self.ir.virtual_ips
                ),
            ),
            (
                "VIP Nested Configuration",
                sum(len(vip.nested_source_configs) for vip in self.ir.virtual_ips),
            ),
            ("VIP Groups", len(self.ir.virtual_ip_groups)),
            ("NAT Rules", len(self.ir.nat_rules)),
            ("VPN Tunnels", len(self.ir.vpn_tunnels)),
            ("VPN Phase 2", len(self.ir.vpn_phase2)),
            ("GlobalProtect Portals", len(pan_ext.global_protect_portals)),
            ("GlobalProtect Gateways", len(pan_ext.global_protect_gateways)),
            ("GlobalProtect Network Gateways", len(pan_ext.global_protect_network_gateways)),
            ("GlobalProtect Client Auth", sum(len(item.client_authentication) for item in pan_ext.global_protect_portals + pan_ext.global_protect_gateways)),
            ("GlobalProtect Portal Configs", sum(len(item.client_configs) for item in pan_ext.global_protect_portals)),
            ("GlobalProtect External Gateways", sum(len(config.external_gateways) for item in pan_ext.global_protect_portals for config in item.client_configs)),
            ("GlobalProtect App Settings", sum(len(config.app_settings) for item in pan_ext.global_protect_portals for config in item.client_configs)),
            ("GlobalProtect Gateway Roles", sum(len(item.roles) for item in pan_ext.global_protect_gateways)),
            ("GlobalProtect Tunnel Configs", sum(len(item.remote_user_tunnel_configs) for item in pan_ext.global_protect_gateways)),
            ("SSL VPN Portals", len(self._vendor_extension_value("ssl_vpn_portals") or [])),
            ("SSL VPN Host Checks", len(self._vendor_extension_value("ssl_vpn_host_checks") or [])),
            (
                "SSL VPN Host Check Items",
                sum(len(item.check_items) for item in self._vendor_extension_value("ssl_vpn_host_checks") or []),
            ),
            (
                "SD-WAN Rules",
                sum(len(sdwan.rules) for sdwan in forti_ext.sdwans)
                + len(pan_ext.pan_sdwan_rules),
            ),
            ("LDAP Servers", len(self._vendor_extension_value("user_ldap_servers") or [])),
            ("RADIUS Servers", len(self._vendor_extension_value("user_radius_servers") or [])),
            ("TACACS+ Servers", len(self._vendor_extension_value("user_tacacs_servers") or [])),
            ("SAML Servers", len(self._vendor_extension_value("user_saml_servers") or [])),
            ("FSSO Servers", len(forti_ext.fsso_providers)),
            ("FSSO AD Groups", len(forti_ext.fsso_ad_groups)),
            ("FSSO Polling", len(forti_ext.fsso_polling)),
            ("Local Users", len(self._vendor_extension_value("local_users") or [])),
            ("User Groups", len(self._vendor_extension_value("user_groups") or [])),
            (
                "User Authentication Settings",
                1 if self._vendor_extension_value("user_authentication_settings") is not None else 0,
            ),
            (
                "User Quarantine",
                1 if self._vendor_extension_value("user_quarantine_settings") is not None else 0,
            ),
            ("DoS Policies", len(self._vendor_extension_value("dos_policies") or [])),
            ("Firewall Sniffers", len(self._vendor_extension_value("firewall_sniffers") or [])),
            ("Certificates", len(self.ir.certificates)),
            ("Routes", len(self.ir.routes)),
            (
                "Security Profiles",
                len(self.ir.security_profile_groups),
            ),
            ("Security Profile Definitions", len(self.ir.security_profile_definitions)),
            ("Security Profile Rules", sum(len(item.rules) for item in self.ir.security_profile_definitions)),
            ("Custom URL Categories", len(self.ir.custom_url_categories)),
        ]

        self._summary_section(
            sheet,
            navigation_end_row + 2,
            "Inventory Counts",
            inventory_rows,
        )

        sheet.column_dimensions["A"].width = 22
        sheet.column_dimensions["B"].width = 34
        sheet.column_dimensions["C"].width = 12
        sheet.column_dimensions["D"].width = 48
        sheet.column_dimensions["E"].width = 16

        sheet.freeze_panes = "A4"

    @staticmethod
    def _set_internal_link(cell: Any, sheet_name: str) -> None:
        escaped = sheet_name.replace("'", "''")
        cell.hyperlink = f"#'{escaped}'!A1"
        cell.style = "Hyperlink"

    def _summary_section(
        self, sheet: Any, start_row: int, title: str, rows: Sequence[tuple[str, Any]]
    ) -> None:
        sheet.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=2)
        title_cell = sheet.cell(start_row, 1, title)
        title_cell.font = Font(name="Aptos", size=11, bold=True, color=self._WHITE)
        title_cell.fill = PatternFill("solid", fgColor=self._TEAL)
        title_cell.alignment = Alignment(vertical="center")
        sheet.row_dimensions[start_row].height = 23
        for row_index, (label, value) in enumerate(rows, start_row + 1):
            sheet.cell(row_index, 1, self._safe_value(label))
            sheet.cell(row_index, 2, self._safe_value(value))
            sheet.cell(row_index, 1).font = Font(name="Aptos", bold=True, color=self._TEXT)
            sheet.cell(row_index, 1).fill = PatternFill("solid", fgColor=self._LIGHT_BLUE)
            sheet.cell(row_index, 2).font = Font(name="Aptos", color=self._TEXT)
            sheet.cell(row_index, 2).alignment = Alignment(wrap_text=True, vertical="top")

    def _build_summary_navigation(
        self,
        sheet: Any,
        workbook: Any,
        start_row: int,
    ) -> int:
        """Build workbook navigation using the already-generated worksheets."""
        sheet.merge_cells(
            start_row=start_row,
            start_column=1,
            end_row=start_row,
            end_column=5,
        )

        title_cell = sheet.cell(
            start_row,
            1,
            "Workbook Navigation",
        )
        title_cell.font = Font(
            name="Aptos",
            size=11,
            bold=True,
            color=self._WHITE,
        )
        title_cell.fill = PatternFill(
            "solid",
            fgColor=self._TEAL,
        )
        title_cell.alignment = Alignment(
            vertical="center",
        )
        sheet.row_dimensions[start_row].height = 23

        header_row = start_row + 1

        headers = (
            "Category",
            "Sheet",
            "Records",
            "Purpose",
            "Manual Review",
        )

        for column, header in enumerate(headers, 1):
            cell = sheet.cell(
                header_row,
                column,
                header,
            )
            cell.font = Font(
                name="Aptos",
                bold=True,
                color=self._WHITE,
            )
            cell.fill = PatternFill(
                "solid",
                fgColor=self._NAVY,
            )
            cell.alignment = Alignment(
                wrap_text=True,
                vertical="center",
            )

        generated_sheets = {
            worksheet.title: worksheet
            for worksheet in workbook.worksheets
            if worksheet.title != "Summary"
        }

        row_index = header_row

        for sheet_name in self.SHEET_ORDER:
            if sheet_name == "Summary" or sheet_name in _REMOVED_SHEETS:
                continue

            target_sheet = generated_sheets.get(sheet_name)
            if target_sheet is None:
                continue

            row_index += 1

            category = self._sheet_category(sheet_name)

            # All normal table worksheets use:
            # row 1 = title
            # row 2 = note
            # row 3 = headers
            #
            # Therefore max_row - 3 is the actual record count.
            record_count = max(
                target_sheet.max_row - 3,
                0,
            )

            values = (
                category,
                sheet_name,
                record_count,
                self._sheet_purpose(
                    sheet_name,
                    category,
                ),
                (
                    "Yes"
                    if self._sheet_requires_review(
                        sheet_name,
                        category,
                    )
                    else "No"
                ),
            )

            for column, value in enumerate(values, 1):
                cell = sheet.cell(
                    row_index,
                    column,
                    self._safe_value(value),
                )
                cell.font = Font(
                    name="Aptos",
                    size=10,
                    color=self._TEXT,
                )
                cell.alignment = Alignment(
                    wrap_text=True,
                    vertical="top",
                )

            self._set_internal_link(
                sheet.cell(row_index, 2),
                sheet_name,
            )

            if (row_index - header_row) % 2 == 0:
                for column in range(1, 6):
                    sheet.cell(
                        row_index,
                        column,
                    ).fill = PatternFill(
                        "solid",
                        fgColor="F8FAFC",
                    )

        sheet.auto_filter.ref = (
            f"A{header_row}:E{row_index}"
        )

        return row_index


    def _sheet_category(
        self,
        sheet_name: str,
    ) -> str:
        if sheet_name in self.CORE_INVENTORY_SHEETS:
            return "Core Inventory"

        if sheet_name in self.NETWORK_ACCESS_SHEETS:
            return "Network / Access"

        if sheet_name in self.SOURCE_DETAIL_SHEETS:
            return "Source Detail"

        if sheet_name in self.AUDIT_SHEETS:
            return "Audit"

        return "Overview"


    @staticmethod
    def _sheet_purpose(
        sheet_name: str,
        category: str,
    ) -> str:
        purposes = {
            "System Settings": "System-level firewall settings",
            "DNS Settings": "Configured DNS settings",
            "NTP Settings": "Configured NTP servers",
            "Management Service Routes": "PAN-OS management service source routes",
            "Interfaces": "Interface inventory",
            "Zones": "Security/interface zones",
            "Addresses": "Address objects",
            "Address Groups": "Address object groups",
            "Services": "Service and protocol objects",
            "Service Groups": "Service object groups",
            "Schedules": "Policy schedule objects",
            "Policies": "Firewall security policies",
            "Default Security Rules": "Configured PAN-OS default security-rule overrides",
            "Local-In Policies": "FortiGate control-plane local-in policies",
            "NGFW Security Policies": "FortiGate policy-based NGFW security policies",
            "Multicast Policies": "Canonical multicast policy inventory",
            "IP Pools": "Source NAT pools",
            "Virtual IPs": "Destination NAT/VIP objects",
            "VIP Real Servers": "VIP backend servers",
            "VIP Nested Configuration": "FortiGate VIP nested source evidence",
            "VIP Groups": "FortiGate VIP groups",
            "NAT Rules": "Normalized NAT inventory",
            "PBF Rules": "Policy-based forwarding and source route-table inventory",
            "Routes": "Static route inventory",
            "Policy Routes": "FortiGate policy-route source inventory",
            "Cisco PBR": "Cisco ASA/FTD policy-based routing inventory",
            "VPN Tunnels": "IPsec Phase 1 / tunnel inventory",
            "VPN Phase 2": "IPsec Phase 2 selectors and settings",
            "GlobalProtect Portals": "PAN-OS GlobalProtect portal source inventory",
            "GlobalProtect Gateways": "PAN-OS VSYS GlobalProtect gateway inventory",
            "GlobalProtect Client Auth": "GlobalProtect portal and gateway client authentication",
            "GlobalProtect Portal Configs": "GlobalProtect portal client configurations",
            "GlobalProtect External Gateways": "GlobalProtect external gateway definitions",
            "GlobalProtect App Settings": "Ordered GlobalProtect application settings",
            "GlobalProtect Root CAs": "GlobalProtect portal root CA references",
            "GlobalProtect Gateway Roles": "GlobalProtect gateway role/session settings",
            "GlobalProtect Tunnel Configs": "GlobalProtect remote-user tunnel configurations",
            "GlobalProtect Network Gateways": "Legacy network GlobalProtect gateways",
            "Interface Source Settings": (
                "Explicit FortiGate/source interface settings"
            ),
            "Source Security Profiles": (
                "Source security-profile inventory"
            ),
            "Source Security Profile Setting": (
                "Detailed source security-profile settings"
            ),
            "Extraction Coverage": (
                "Source-to-parser extraction coverage"
            ),
        }

        return purposes.get(
            sheet_name,
            f"{category}: {sheet_name}",
        )


    @staticmethod
    def _sheet_requires_review(
        sheet_name: str,
        category: str,
    ) -> bool:
        if category in {
            "Source Detail",
            "Audit",
        }:
            return True

        return sheet_name in {
            "NTP Settings",
            "Management Service Routes",
            "Addresses",
            "Interface Secondary IPs",
            "Policies",
            "Local-In Policies",
            "NGFW Security Policies",
            "Multicast Policies",
            "NAT Rules",
            "Routes",
            "Policy Routes",
            "VPN Tunnels",
            "VPN Phase 2",
            "GlobalProtect Portals",
            "GlobalProtect Gateways",
            "GlobalProtect Client Auth",
            "GlobalProtect Portal Configs",
            "GlobalProtect External Gateways",
            "GlobalProtect App Settings",
            "GlobalProtect Root CAs",
            "GlobalProtect Gateway Roles",
            "GlobalProtect Tunnel Configs",
            "GlobalProtect Network Gateways",
            "SSL VPN Settings",
            "SSL VPN Portals",
            "SSL VPN Authentication Rules",
            "SSL VPN Host Checks",
            "SSL VPN Host Check Items",
        }


    def _order_sheets(
        self,
        workbook: Any,
    ) -> None:
        """Validate and apply deterministic logical workbook ordering."""
        if len(self.SHEET_ORDER) != len(
            set(self.SHEET_ORDER)
        ):
            duplicates = sorted(
                {
                    sheet_name
                    for sheet_name in self.SHEET_ORDER
                    if self.SHEET_ORDER.count(
                        sheet_name
                    ) > 1
                }
            )

            raise ValueError(
                "Excel sheet order contains duplicate entries: "
                f"{duplicates}"
            )

        expected = set(self.SHEET_ORDER)
        for worksheet in list(workbook.worksheets):
            if worksheet.title in _REMOVED_SHEETS:
                workbook.remove(worksheet)

        actual = {
            worksheet.title
            for worksheet in workbook.worksheets
        }

        unknown = actual - expected

        if unknown:
            raise ValueError(
                "Excel sheet order missing entries for: "
                f"{sorted(unknown)}"
            )

        workbook._sheets.sort(
            key=lambda worksheet: (
                self.SHEET_ORDER.index(
                    worksheet.title
                )
            )
        )

        for worksheet in workbook.worksheets:
            worksheet.sheet_properties.tabColor = (
                self._sheet_tab_color(
                    worksheet.title
                )
            )


    def _sheet_tab_color(
        self,
        sheet_name: str,
    ) -> str:
        if sheet_name == "Summary":
            return self._NAVY

        if sheet_name in self.CORE_INVENTORY_SHEETS:
            return self._TEAL

        if sheet_name in self.NETWORK_ACCESS_SHEETS:
            return self._LIGHT_BLUE

        if sheet_name in self.SOURCE_DETAIL_SHEETS:
            return self._MUTED

        if sheet_name in self.AUDIT_SHEETS:
            return self._LIGHT_AMBER

        return self._LIGHT_BLUE
    def _build_interfaces(self, workbook: Any) -> None:
        headers = (
            "Name", "Source VDOM", "Zone", "VRF", "Virtual Router / Routing Instance",
            "Routing Instance Type", "IP / Prefix", "Additional IPv4 Addresses",
            "Additional IPv4 Source Addresses", "Remote IP / Prefix",
            "IPv6 Address", "IPv6 Source Address", "IPv6 Management Access",
            "Additional IPv6 Addresses", "Additional IPv6 Source Addresses",
            "IPv6 Prefix Advertisements", "IPv6 Delegated Prefixes", "DHCPv6 IA-PD",
            "IPv6 VRRP6",
            "IPv6 Mode", "IPv6 Send Adv", "IPv6 Manage Flag", "IPv6 Other Flag",
            "IPv6 Autoconf", "CLI IPv6 Connection Status", "DHCPv6 Client Options", "DHCPv6 Information Request",
            "DHCPv6 Prefix Delegation", "DHCPv6 Relay Interface ID", "DHCPv6 Relay IP",
            "DHCPv6 Relay Service", "DHCPv6 Relay Source Interface", "DHCPv6 Relay Source IP",
            "DHCPv6 Relay Type", "ICMPv6 Send Redirect", "IPv6 Interface Identifier",
            "IPv6 Default Life", "IPv6 Delegated Prefix IAID", "IPv6 DNS Server Override",
            "IPv6 Hop Limit", "IPv6 Link MTU", "IPv6 Max Interval", "IPv6 Min Interval",
            "IPv6 Prefix Mode", "IPv6 Reachable Time", "IPv6 Retransmit Time",
            "IPv6 Subnet", "IPv6 Upstream Interface",
            "Enabled", "MTU", "Link State", "Speed", "Duplex", "Media Type",
            "Bandwidth Monitoring",
            "Device Identification", "NetFlow Profile",
            "LLDP Enabled", "Interface Type", "Role", "Dedicated To", "Members", "Extraction Status",
            "IKE SAML Server", "IKE SAML Server Resolved", "Source IP Check",
            "Migration Status", "Manual Review", "Review Reasons", "Additional Settings",
            "Addressing Mode", "DHCP Client", "DNS Server Override", "Management Access", "Alias",
            "Parent / Underlay Interface",
            "Tag", "VLAN ID", "Management Profile", "PPPoE Mode", "PPPoE Username",
            "PPPoE Password Configured", "PPPoE Password Format",
            "Description",
        )
        rows = [
            (
                item.name, item.source_vdom, item.zone, item.source_vrf,
                item.source_routing_instance, item.source_routing_instance_type,
                item.ip,
                "\n".join(address.address or "" for address in item.additional_ipv4_addresses),
                "\n".join(address.source_address for address in item.additional_ipv4_addresses),
                item.remote_ip, item.ipv6_address,
                item.source_ipv6_address, item.source_ipv6_management_access,
                "\n".join(address.address or "" for address in item.additional_ipv6_addresses),
                "\n".join(address.source_address for address in item.additional_ipv6_addresses),
                "\n".join(
                    f"{prefix.prefix or ''} ({prefix.source_prefix})"
                    for prefix in item.ipv6_prefix_advertisements
                ),
                "\n".join(
                    f"{prefix.prefix_id}: {prefix.subnet or ''} ({prefix.upstream_interface or ''})"
                    for prefix in item.ipv6_delegated_prefixes
                ),
                "\n".join(
                    f"{iapd.source_iaid}: {iapd.prefix_hint or ''}"
                    for iapd in item.dhcp6_iapd
                ),
                "\n".join(
                    f"{vrrp.source_vrid}: {vrrp.vrip6 or ''}"
                    for vrrp in item.vrrp6
                ),
                item.source_ipv6_mode, item.source_ipv6_send_adv,
                item.source_ipv6_manage_flag, item.source_ipv6_other_flag,
                item.source_ipv6_autoconf,
                item.source_cli_conn6_status,
                "\n".join(item.source_dhcp6_client_options),
                item.source_dhcp6_information_request, item.source_dhcp6_prefix_delegation,
                item.source_dhcp6_relay_interface_id, "\n".join(item.source_dhcp6_relay_ip),
                item.source_dhcp6_relay_service, item.source_dhcp6_relay_source_interface,
                item.source_dhcp6_relay_source_ip, item.source_dhcp6_relay_type,
                item.source_icmp6_send_redirect, item.source_ipv6_interface_identifier,
                item.source_ip6_default_life, item.source_ip6_delegated_prefix_iaid,
                item.source_ip6_dns_server_override, item.source_ip6_hop_limit,
                item.source_ip6_link_mtu, item.source_ip6_max_interval, item.source_ip6_min_interval,
                item.source_ip6_prefix_mode, item.source_ip6_reachable_time,
                item.source_ip6_retrans_time, item.source_ip6_subnet,
                item.source_ip6_upstream_interface,
                item.status, item.mtu if item.mtu is not None else item.source_mtu, item.source_link_state,
                item.source_speed, item.source_duplex, item.source_media_type,
                self._optional_bool_literal(item.source_monitor_bandwidth),
                item.source_device_identification, item.source_netflow_profile,
                item.source_lldp_enabled,
                item.interface_type, item.role, item.source_dedicated_to,
                ", ".join(str(member) for member in item.members),
                item.migration_status,
                item.source_ike_saml_server,
                self._optional_bool_literal(item.source_ike_saml_server_resolved),
                self._optional_bool_literal(item.source_src_check),
                item.migration_status,
                self._optional_bool_literal(item.requires_manual_review),
                ", ".join(item.review_reasons),
                self._format_settings(item.source_attributes),
                item.addressing_mode, item.dhcp_client,
                self._optional_bool_literal(item.source_dns_server_override),
                item.management_access,
                item.alias, item.parent, item.tag, item.vlanid, item.management_profile,
                item.pppoe_mode, item.pppoe_username,
                self._optional_bool_literal(item.has_pppoe_password),
                item.pppoe_password_format, item.description,
            )
            for item in self.ir.interfaces
        ]
        self._table_sheet(workbook, "Interfaces", headers, rows)

    def _build_interface_secondary_ips(self, workbook: Any) -> None:
        headers = (
            "Interface",
            "Secondary IP Status",
            "Source ID",
            "Source IP",
            "IP / Prefix",
            "Management Access",
            "Extraction Status",
            "Manual Review",
            "Parse Error",
            "Additional Settings",
        )
        rows = []
        for intf in self.ir.interfaces:
            for sec in getattr(intf, "secondary_ips", []):
                if sec.parse_error:
                    status = "PARSE_ERROR"
                elif sec.requires_manual_review:
                    status = "PARTIALLY_NORMALIZED"
                else:
                    status = "NORMALIZED"

                rows.append(
                    (
                        intf.name,
                        "ACTIVE",
                        sec.source_id,
                        sec.source_ip,
                        sec.ip,
                        sec.management_access,
                        status,
                        self._optional_bool_literal(sec.requires_manual_review),
                        sec.parse_error,
                        self._format_settings(sec.source_attributes),
                    )
                )
            for sec in getattr(intf, "inactive_secondary_ips", []):
                parent_status = getattr(intf, "source_secondary_ip_status", None)
                secondary_status = (
                    "DISABLED" if parent_status == "disable" else "AMBIGUOUS"
                )
                rows.append(
                    (
                        intf.name,
                        secondary_status,
                        sec.source_id,
                        sec.source_ip,
                        None,
                        [],
                        "EXTRACT_ONLY",
                        "TRUE",
                        sec.parse_error,
                        self._format_settings(sec.source_attributes),
                    )
                )

        self._table_sheet(
            workbook,
            "Interface Secondary IPs",
            headers,
            rows,
            empty_note="No secondary interface IP addresses were extracted from the source firewall.",
            subtitle="Secondary interface IP configuration extracted for migration and inventory review.",
        )

    def _build_interface_source_settings(self, workbook: Any) -> None:
        """Expose every explicitly configured interface setting without reinterpreting it."""
        rows = []
        for item in self.ir.interfaces:
            for setting, value in sanitize_source_attributes(item.source_attributes).items():
                display_setting = str(setting).replace("_", "-")
                if isinstance(value, set):
                    display_value = json.dumps(sorted(value), ensure_ascii=False, default=str)
                elif isinstance(value, (dict, list, tuple)):
                    display_value = json.dumps(
                        value, ensure_ascii=False, sort_keys=True, default=str
                    )
                else:
                    display_value = value
                rows.append(
                    (
                        item.name,
                        self.ir.metadata.source_vendor,
                        display_setting,
                        display_value,
                        "EXTRACT_ONLY",
                    )
                )

        self._table_sheet(
            workbook,
            "Interface Source Settings",
            ("Interface", "Source Vendor", "Setting", "Value", "Extraction Status"),
            rows,
            empty_note="No explicit source-interface settings were retained by the parser.",
            subtitle=(
                "Explicit source settings retained for inventory. These values are extraction-only "
                "and are not consumed by target generators."
            ),
        )

    def _build_dhcp_servers(
        self,
        workbook: Any,
    ) -> None:
        rows = [
        (
            item.source_id,
            item.interface,
            item.enabled,
            item.default_gateway,
            item.netmask,
            item.lease_time_seconds,
            item.dns_service,
            item.dns_servers,
            item.timezone_option,
            item.migration_status,
            item.requires_manual_review,
            self._format_settings(
                item.source_attributes
            ),
        )
        for item in self.ir.dhcp_servers
        ]

        self._table_sheet(
            workbook,
            "DHCP Servers",
        (
            "Server ID",
            "Interface",
            "Enabled",
            "Default Gateway",
            "Netmask",
            "Lease Time (Seconds)",
            "DNS Service",
            "DNS Servers",
            "Timezone Option",
            "Extraction Status",
            "Manual Review",
            "Additional Settings",
        ),
            rows,
            empty_note=(
            "No DHCP server configuration was "
            "extracted from the source firewall."
        ),
            subtitle=(
            "DHCP server configuration retained for "
            "migration review."
        ),
        )


    def _build_dhcp_ip_ranges(
        self,
        workbook: Any,
    ) -> None:
        rows = [
        (
            server.source_id,
            server.interface,
            item.source_id,
            item.start_ip,
            item.end_ip,
            "EXTRACT_ONLY",
            True,
            self._format_settings(
                item.source_attributes
            ),
        )
        for server in self.ir.dhcp_servers
        for item in server.ip_ranges
        ]

        self._table_sheet(
            workbook,
            "DHCP IP Ranges",
        (
            "Server ID",
            "Interface",
            "Range ID",
            "Start IP",
            "End IP",
            "Extraction Status",
            "Manual Review",
            "Additional Settings",
        ),
            rows,
            empty_note=(
            "No DHCP IP ranges were extracted "
            "from the source firewall."
        ),
        )


    def _build_dhcp_reservations(
        self,
        workbook: Any,
    ) -> None:
        rows = [
        (
            server.source_id,
            server.interface,
            item.source_id,
            item.ip_address,
            item.mac_address,
            "EXTRACT_ONLY",
            True,
            self._format_settings(
                item.source_attributes
            ),
        )
        for server in self.ir.dhcp_servers
        for item in server.reservations
        ]

        self._table_sheet(
            workbook,
            "DHCP Reservations",
        (
            "Server ID",
            "Interface",
            "Reservation ID",
            "IP Address",
            "MAC Address",
            "Extraction Status",
            "Manual Review",
            "Additional Settings",
        ),
            rows,
            empty_note=(
            "No DHCP reservations were extracted "
            "from the source firewall."
        ),
        )

    def _build_zones(self, workbook: Any) -> None:
        rows = [
            (
                item.source_context,
                item.name,
                item.zone_type,
                item.interfaces,
                item.description,
                item.source_intrazone,
                item.source_effective_intrazone,
                item.source_path,
                self._optional_bool_literal(item.requires_manual_review),
                self._format_settings(item.source_attributes),
            )
            for item in self.ir.zones
        ]
        self._table_sheet(
            workbook,
            "Zones",
            (
                "VDOM",
                "Name",
                "Zone Type",
                "Members",
                "Description",
                "Configured Intrazone",
                "Effective Intrazone",
                "Source Path",
                "Manual Review",
                "Additional Settings",
            ),
            rows,
        )

    def _build_addresses(self, workbook: Any) -> None:
        from itertools import chain

        address_items = chain(
            self.ir.addresses,
            (
                IRAddress(
                name=item.name,
                type=AddressType.DYNAMIC,
                source_context=item.source_context,
                source_uuid=item.source_uuid,
                source_section=item.source_section,
                address_family=item.address_family,
                source_type="dynamic",
                dynamic_filter=item.dynamic_filter,
                tag_name=item.tags[0] if item.tags else None,
                source_sub_type=self._object_extension_value(item, "source_sub_type"),
                source_obj_tag=self._object_extension_value(item, "source_obj_tag"),
                source_tag_type=self._object_extension_value(item, "source_tag_type"),
                source_obj_type=self._object_extension_value(item, "source_obj_type"),
                source_dirty=self._object_extension_value(item, "source_dirty"),
                source_attributes=dict(item.source_attributes),
                migration_status=item.migration_status,
                requires_manual_review=item.requires_manual_review,
                audit_note=item.audit_note,
                description=item.description,
                )
                for item in self.ir.address_groups
                if self._is_source_address_backed_group(item)
            )
        )
        rows = (
            (
                item.name,
                item.source_uuid,
                item.type,
                item.value,
                "; ".join(
                    f"{entry.start}-{entry.end}" if entry.end else entry.start
                    for entry in item.mac_entries
                ) or None,
                len(item.mac_entries) or None,
                item.source_section,
                item.address_family,
                item.source_type,
                item.original_type,
                item.original_value,
                item.is_ipv6,
                item.is_multicast,
                item.associated_interface,
                self._optional_bool_literal(
                    item.allow_routing
                ),
                item.source_color,
                self._object_extension_value(item, "source_fsso_group"),
                self._object_extension_value(item, "source_hw_vendor"),
                self._object_extension_value(item, "source_hw_model"),
                item.source_interface,
                item.resolved_interface_subnet,
                self._optional_bool_literal(item.interface_reference_resolved),
                self._object_extension_value(item, "source_cache_ttl"),
                self._object_extension_value(item, "source_clearpass_spt"),
                self._object_extension_value(item, "source_epg_name"),
                self._object_extension_value(item, "source_fabric_object_setting"),
                item.dynamic_filter,
                self._object_extension_value(item, "source_sdn"),
                self._object_extension_value(item, "source_sdn"),
                self._object_extension_value(item, "source_sdn_addr_type"),
                self._object_extension_value(item, "source_sdn_tag"),
                self._object_extension_value(item, "source_organization"),
                self._object_extension_value(item, "source_os"),
                self._object_extension_value(item, "source_policy_group"),
                self._object_extension_value(item, "source_route_tag"),
                self._object_extension_value(item, "source_subnet_name"),
                self._object_extension_value(item, "source_sw_version"),
                self._object_extension_value(item, "source_tag_detection_level"),
                self._object_extension_value(item, "source_tenant"),
                self._optional_bool_literal(self._object_extension_value(item, "source_node_ip_only")),
                self._object_extension_value(item, "source_obj_id"),
                self._object_extension_value(item, "source_sub_type"),
                self._object_extension_value(item, "source_obj_tag"),
                self._object_extension_value(item, "source_tag_type"),
                self._object_extension_value(item, "source_obj_type"),
                self._object_extension_value(item, "source_dirty"),
                item.tags,
                item.source_list_entries,
                [
                    {"name": entry.name, "category": entry.category, "tags": entry.tags}
                    for entry in item.source_tagging_entries
                ],
                item.migration_status,
                item.requires_manual_review,
                item.audit_note,
                item.parse_error,
                self._format_settings(
                    {
                        key: value
                        for key, value in item.source_attributes.items()
                        if key not in {"cache_ttl", "clearpass_spt", "epg_name", "fabric_object", "fsso_group", "hw_vendor", "hw_model"}
                    }
                ),
                item.description,
            )
            for item in address_items
        )
        self._table_sheet(
            workbook,
            "Addresses",
            (
                "Name",
                "Source UUID",
                "Type",
                "Value",
                "MAC Entries",
                "MAC Count",
                "Source Section",
                "Address Family",
                "Source Type",
                "Original Type",
                "Original Value",
                "IPv6",
                "Multicast",
                "Associated Interface",
                "Allow Routing",
                "Source Color",
                "FSSO Group",
                "Hardware Vendor",
                "Hardware Model",
                "Source Interface",
                "Resolved Interface Subnet",
                "Interface Reference Resolved",
                "Cache TTL",
                "ClearPass SPT",
                "EPG Name",
                "Fabric Object",
                "Dynamic Filter",
                "SDN",
                "SDN Connector",
                "SDN Address Type",
                "SDN Tag",
                "Organization",
                "OS",
                "Policy Group",
                "Route Tag",
                "Subnet Name",
                "Software Version",
                "Tag Detection Level",
                "Tenant",
                "Node IP Only",
                "NSX Object ID",
                "EMS Sub-Type",
                "EMS Object Tag",
                "EMS Tag Type",
                "EMS Object Type",
                "EMS Dirty",
                "Tags",
                "IP List",
                "Object Tagging",
                "Migration Status",
                "Manual Review",
                "Audit Note",
                "Parse Error",
                "Additional Settings",
                "Description",
            ),
            rows,
        )

    def _build_address_groups(self, workbook: Any) -> None:
        rows = (
            (
                item.name,
                item.source_uuid,
                item.members,
                item.source_direct_members,
                item.source_nested_group_members,
                item.members,
                item.is_dynamic,
                item.dynamic_filter,
                self._optional_bool_literal(
                    item.allow_routing
                ),
                item.source_color,
                item.source_category,
                self._object_extension_value(item, "source_sub_type"),
                self._object_extension_value(item, "source_obj_tag"),
                self._object_extension_value(item, "source_tag_type"),
                self._object_extension_value(item, "source_obj_type"),
                self._object_extension_value(item, "source_dirty"),
                item.tags,
                item.source_section,
                item.address_family,
                item.exclusion_enabled,
                item.exclude_members,
                self._object_extension_value(item, "source_exclude_setting"),
                self._object_extension_value(item, "source_group_type"),
                self._object_extension_value(item, "source_fabric_object_setting"),
                item.migration_status,
                self._optional_bool_literal(item.requires_manual_review),
                item.audit_note,
                item.source_context,
                self._format_settings(
                    item.source_attributes
                ),
                item.description,
            )
            for item in self.ir.address_groups
            if not self._is_source_address_backed_group(item)
        )
        self._table_sheet(
            workbook,
            "Address Groups",
            (
                "Name",
                "Source UUID",
                "Members",
                "Direct Address Members",
                "Nested Address-Set Members",
                "Resolved Members",
                "Dynamic",
                "Dynamic Filter",
                "Allow Routing",
                "Source Color",
                "Source Category",
                "EMS Sub-Type",
                "EMS Object Tag",
                "EMS Tag Type",
                "EMS Object Type",
                "EMS Dirty",
                "Tags",
                "Source Section",
                "Address Family",
                "Exclusion Enabled",
                "Exclude Members",
                "Source Exclude Setting",
                "Group Type",
                "Fabric Object",
                "Migration Status",
                "Manual Review",
                "Audit Note",
                "Source VDOM",
                "Additional Settings",
                "Description",
            ),
            rows,
        )

    def _build_service_categories(self, workbook: Any) -> None:
        rows = [
            (
                item.name,
                item.description,
                item.source_fabric_object,
                item.migration_status,
                self._format_settings(
                    item.source_attributes
                ),
            )
            for item in self.ir.service_categories
        ]
        self._table_sheet(
            workbook,
            "Service Categories",
            (
                "Name",
                "Description",
                "Fabric Object",
                "Extraction Status",
                "Additional Settings",
            ),
            rows,
        )

    def _build_services(self, workbook: Any) -> None:
        rows = (
            (
                item.name,
                item.source_uuid,
                item.source_category,
                item.source_protocol_configured,
                item.source_protocol,
                item.source_protocol_number,
                [
                    self._format_port(port)
                    for port in item.ports
                ],
                [
                    port.source_port
                    for port in item.ports
                    if port.source_port is not None
                ],
                self._optional_bool_literal(
                    item.source_proxy
                ),
                item.source_color,
                item.source_fabric_object,
                item.source_unmodeled_semantic_settings,
                item.migration_status,
                self._optional_bool_literal(
                    item.requires_manual_review
                ),
                item.audit_note,
                self._format_settings(
                    item.source_attributes
                ),
                item.description,
            )
            for item in self.ir.services
        )
        self._table_sheet(
            workbook,
            "Services",
            (
                "Name",
                "Source UUID",
                "Category",
                "Configured Protocol",
                "Effective Protocol",
                "Source Protocol Number",
                "Protocol / Destination Port",
                "Source Port Constraint",
                "Proxy",
                "Source Color",
                "Fabric Object",
                "Unmodeled Semantic Settings",
                "Migration Status",
                "Manual Review",
                "Audit Note",
                "Additional Settings",
                "Description",
            ),
            rows,
        )

    def _build_service_groups(
    self,
    workbook: Any,
    ) -> None:
        rows = [
            (
                item.name,
                item.source_uuid,
                item.members,
                item.unsafe_members,
                self._optional_bool_literal(item.source_proxy),
                item.source_color,
                item.source_fabric_object,
                item.migration_status,
                self._optional_bool_literal(item.requires_manual_review),
                item.audit_note,
                self._format_settings(
                    item.source_attributes
                ),
                item.description,
            )
            for item in self.ir.service_groups
        ]

        self._table_sheet(
            workbook,
            "Service Groups",
            (
                "Name",
                "Source UUID",
                "Members",
                "Unsafe Members",
                "Proxy",
                "Source Color",
                "Fabric Object",
                "Migration Status",
                "Manual Review",
                "Audit Note",
                "Additional Settings",
                "Description",
            ),
            rows,
        )

    def _build_session_helpers(
        self,
        workbook: Any,
    ) -> None:
        """
        Export FortiGate session-helper / ALG inventory.

        Session helpers influence protocol handling but should not be
        converted into normal firewall service objects.
        """

        rows = [
            (
                item.source_id,
                item.name,
                item.protocol_name,
                item.protocol_number,
                item.port,
                item.classification,
                item.migration_status,
                item.requires_manual_review,
                self._format_settings(
                    item.source_attributes
                ),
            )
            for item in self._vendor_extension_value("session_helpers") or []
        ]

        sheet = self._table_sheet(
            workbook,
            "Session Helpers",
            (
                "Source ID",
                "Name",
                "Protocol",
                "Protocol Number",
                "Port",
                "Classification",
                "Extraction Status",
                "Manual Review",
                "Additional Settings",
            ),
            rows,
            empty_note=(
                "No FortiGate session-helper entries were "
                "extracted from the source configuration."
            ),
            subtitle=(
                "FortiGate protocol/session helpers retained for "
                "traffic-behavior inventory. DEFAULT entries match "
                "the known FortiOS baseline. CUSTOM, CUSTOMIZED, "
                "or UNKNOWN entries require target-platform review. "
                "Session helpers are not converted into service objects."
            ),
        )

        for row in range(4, sheet.max_row + 1):
            classification = str(
                sheet.cell(row, 6).value or ""
            ).upper()

            if classification in {
                "CUSTOM",
                "CUSTOMIZED",
                "UNKNOWN",
            }:
                for column in range(1, 10):
                    sheet.cell(
                        row,
                        column,
                    ).fill = PatternFill(
                        "solid",
                        fgColor=self._LIGHT_AMBER,
                    )

    def _build_session_ttl_settings(self, workbook: Any) -> None:
        settings = self._vendor_extension_value("session_ttl_settings")
        if settings is None:
            self._table_sheet(workbook, "Session TTL Settings", (
                "Default TTL", "Default Never", "Extraction Status",
                "Manual Review", "Additional Settings",
            ), ())
            return
        self._table_sheet(
            workbook,
            "Session TTL Settings",
            (
                "Default TTL",
                "Default Never",
                "Extraction Status",
                "Manual Review",
                "Additional Settings",
            ),
            (
                (
                    "never" if settings.default_never else settings.default_timeout_seconds,
                    settings.default_never,
                    settings.migration_status,
                    settings.requires_manual_review,
                    self._format_settings(settings.source_attributes),
                ),
            ),
        )

    def _build_session_ttl_overrides(
        self,
        workbook: Any,
    ) -> None:
        """
        Export explicit FortiGate session timeout overrides.

        These settings affect actual session behaviour and require
        target-platform review.
        """

        rows = [
            (
                item.source_id,
                item.protocol_name,
                item.protocol_number,
                item.start_port,
                item.end_port,
                "never" if item.timeout_never else item.timeout_seconds,
                item.refresh_direction,
                item.migration_status,
                item.requires_manual_review,
                self._format_settings(
                    item.source_attributes
                ),
            )
            for item in self._vendor_extension_value("session_ttl_overrides") or []
        ]

        sheet = self._table_sheet(
            workbook,
            "Session TTL Overrides",
            (
                "Source ID",
                "Protocol",
                "Protocol Number",
                "Start Port",
                "End Port",
                "Timeout",
                "Refresh Direction",
                "Extraction Status",
                "Manual Review",
                "Additional Settings",
            ),
            rows,
            empty_note=(
                "No explicit session TTL port overrides were "
                "extracted from the source configuration."
            ),
            subtitle=(
                "Explicit source session timeout overrides retained "
                "for traffic-behavior migration review. These settings "
                "are target-platform dependent and are not automatically "
                "converted into service or policy objects."
            ),
        )

        for row in range(4, sheet.max_row + 1):
            for column in range(1, 11):
                sheet.cell(
                    row,
                    column,
                ).fill = PatternFill(
                    "solid",
                    fgColor=self._LIGHT_AMBER,
                )

    def _build_schedules(
        self,
        workbook: Any,
    ) -> None:
        rows = (
            (
                item.name,
                item.source_context,
                item.schedule_type,
                item.start,
                item.end,
                item.start_utc,
                item.end_utc,
                item.days,
                item.source_color,
                item.expiration_days,
                item.source_fabric_object,
                item.migration_status,
                self._optional_bool_literal(item.requires_manual_review),
                item.review_reasons,
                self._format_settings(item.source_attributes),
            )
            for item in self.ir.schedules
        )

        self._table_sheet(
            workbook,
            "Schedules",
            (
                "Name",
                "Source Context",
                "Type",
                "Start",
                "End",
                "Start UTC",
                "End UTC",
                "Days",
                "Color",
                "Expiration Days",
                "Source Fabric Object",
                "Migration Status",
                "Manual Review",
                "Review Reasons",
                "Additional Settings",
            ),
            rows,
        )

    def _build_schedule_groups(self, workbook: Any) -> None:
        rows = (
            (
                item.name,
                item.source_context,
                item.members,
                item.unresolved_members,
                item.migration_status,
                self._optional_bool_literal(item.requires_manual_review),
                self._format_settings(item.source_attributes),
                item.description,
            )
            for item in self.ir.schedule_groups
        )
        self._table_sheet(
            workbook,
            "Schedule Groups",
            (
                "Name",
                "Source Context",
                "Members",
                "Unresolved Members",
                "Migration Status",
                "Manual Review",
                "Additional Settings",
                "Description",
            ),
            rows,
            empty_note="No schedule groups were extracted.",
            subtitle="Ordered schedule-group membership retained for source review.",
        )

    def _build_traffic_shapers(self, workbook: Any) -> None:
        rows = [
            (
                item.name,
                item.guaranteed_bandwidth,
                item.maximum_bandwidth,
                item.source_bandwidth_unit,
                item.priority,
                self._optional_bool_literal(item.per_policy),
                item.migration_status,
                self._optional_bool_literal(item.requires_manual_review),
                self._format_settings(item.source_attributes),
            )
            for item in self._vendor_extension_value("traffic_shapers") or []
        ]
        self._table_sheet(
            workbook,
            "Traffic Shapers",
            (
                "Name",
                "Guaranteed Bandwidth",
                "Maximum Bandwidth",
                "Source Bandwidth Unit",
                "Priority",
                "Per Policy",
                "Extraction Status",
                "Manual Review",
                "Additional Settings",
            ),
            rows,
            subtitle=(
                "FortiGate shaping inventory; exact target QoS behavior requires manual review."
            ),
        )

    def _build_proxy_addresses(self, workbook: Any) -> None:
        rows = [
            (
                item.name,
                item.source_uuid,
                item.proxy_address_type,
                item.host,
                item.host_regex,
                item.path,
                item.query,
                item.migration_status,
                self._optional_bool_literal(item.requires_manual_review),
                self._format_settings(item.source_attributes),
            )
            for item in self.ir.proxy_addresses
        ]
        self._table_sheet(
            workbook,
            "Proxy Addresses",
            (
                "Name",
                "Source UUID",
                "Type",
                "Host",
                "Host Regex",
                "Path",
                "Query",
                "Extraction Status",
                "Manual Review",
                "Additional Settings",
            ),
            rows,
            subtitle=(
                "Source proxy-address inventory retained without conversion to firewall addresses."
            ),
        )

    def _build_policies(self, workbook: Any) -> None:
        rows = (
            (
                index, item.source_rule_id, item.source_uuid, item.name, item.source_from_interfaces,
                item.from_zone, item.source_to_interfaces, item.to_zone,
                item.source_address_references, item.source,
                item.source_address_negate_setting,
                item.source_ipv6_address_references,
                item.source_ipv6_address_negate_setting,
                item.destination_address_references, item.destination,
                item.destination_address_negate_setting,
                item.destination_ipv6_address_references,
                item.destination_ipv6_address_negate_setting,
                item.source_user_groups, item.source_users,
                item.unresolved_user_groups, item.unresolved_users,
                self._optional_bool_literal(item.identity_dependency_review),
                item.source_service_references, item.service,
                item.source_service_negate_setting,
                item.source_action, item.action, item.source_schedule, item.schedule,
                item.disabled,
                self._object_extension_value(item, "source_vpn_tunnel"),
                item.source_log_setting, item.source_log_start_setting,
                self._object_extension_value(item, "source_utm_status"),
                self._object_extension_value(item, "source_effective_utm_status"),
                item.log_start, item.log_end,
                self._optional_bool_literal(item.nat_enabled),
                self._optional_bool_literal(item.nat_pool_enabled), item.nat_pool_names,
                item.nat_pool_names6,
                item.applications,
                self._object_extension_value(item, "source_internet_service_status"),
                item.internet_service, item.security_profile_group,
                item.antivirus, item.ips_sensor, item.webfilter, item.application_list,
                item.ssl_ssh_profile, item.source_profile_type,
                 item.source_profile_group, item.source_profile_protocol_options,
                 item.unresolved_security_profiles,
                 item.source_security_profile_references,
                 item.security_profile_reference_statuses,
                 item.unresolved_security_profile_references,
                 self._optional_bool_literal(item.security_profile_semantics_review),
                 item.url_categories,
                 item.security_profile_groups,
                 item.antivirus_profiles,
                 item.vulnerability_profiles,
                 item.antispyware_profiles,
                 item.url_filtering_profiles,
                 item.file_blocking_profiles,
                 item.wildfire_analysis_profiles,
                 item.data_filtering_profiles,
                 item.source_extra_settings.get("pan_unresolved_source_zone", []),
                 item.source_extra_settings.get("pan_unresolved_destination_zone", []),
                self._object_extension_value(item, "source_inspection_mode"),
                self._object_extension_value(item, "source_effective_inspection_mode"),
                self._object_extension_value(item, "source_ztna_status"),
                self._object_extension_value(item, "source_effective_ztna_status"),
                self._object_extension_value(item, "source_ztna_ems_tags"),
                self._object_extension_value(item, "source_timeout_send_rst"),
                self._object_extension_value(item, "source_effective_timeout_send_rst"),
                self._object_extension_value(item, "source_auto_asic_offload"),
                self._object_extension_value(item, "source_effective_auto_asic_offload"),
                self._object_extension_value(item, "source_np_acceleration"),
                self._object_extension_value(item, "source_effective_np_acceleration"),
                self._object_extension_value(item, "source_port_preserve"),
                self._object_extension_value(item, "source_effective_port_preserve"),
                self._object_extension_value(item, "source_policy_expiry"),
                self._object_extension_value(item, "source_effective_policy_expiry"),
                self._object_extension_value(item, "source_policy_expiry_date"),
                self._object_extension_value(item, "source_policy_expiry_date_utc"),
                self._object_extension_value(item, "source_schedule_timeout"),
                self._object_extension_value(item, "source_effective_schedule_timeout"),
                self._object_extension_value(item, "source_reputation_direction"),
                self._object_extension_value(item, "source_effective_reputation_direction"),
                self._object_extension_value(item, "source_reputation_direction6"),
                self._object_extension_value(item, "source_effective_reputation_direction6"),
                self._object_extension_value(item, "source_reputation_minimum"),
                self._object_extension_value(item, "source_effective_reputation_minimum"),
                self._object_extension_value(item, "source_reputation_minimum6"),
                self._object_extension_value(item, "source_effective_reputation_minimum6"),
                self._object_extension_value(item, "source_match_vip"),
                self._object_extension_value(item, "source_effective_match_vip"),
                self._object_extension_value(item, "source_match_vip_only"),
                self._object_extension_value(item, "source_effective_match_vip_only"),
                 self._format_settings(self._policy_source_settings(item)),
                 item.migration_status,
                 self._optional_bool_literal(item.requires_manual_review),
                 item.review_reasons,
                 item.description,
            )
            for index, item in enumerate(self.ir.policies, 1)
        )
        sheet = self._table_sheet(
            workbook,
            "Policies",
            (
                "Rule #",
                "Source Policy ID",
                "Source UUID",
                "Name",
                "Source Interface",
                "From Zone",
                "Destination Interface",
                "To Zone",
                "Source Address (Original)",
                "Source Address (Normalized)",
                "Source Address Negate",
                "Source IPv6 Address",
                "Source IPv6 Address Negate",
                "Destination Address (Original)",
                "Destination Address (Normalized)",
                "Destination Address Negate",
                "Destination IPv6 Address",
                "Destination IPv6 Address Negate",
                "User Groups",
                "Users",
                "Unresolved User Groups",
                "Unresolved Users",
                "Identity Dependency Review",
                "Service (Original)",
                "Service (Normalized)",
                "Service Negate",
                "Action (Original)",
                "Action (Normalized)",
                "Schedule (Original)",
                "Schedule (Normalized)",
                "Disabled",
                "VPN Tunnel",
                "Log Setting",
                "Log Start Setting",
                "UTM Status",
                "Effective UTM Status",
                "Log Start",
                "Log End",
                "NAT Enabled",
                "IP Pool Enabled",
                "NAT Pool",
                "NAT Pool IPv6",
                "Applications",
                "Internet Service Status",
                "Internet Services",
                "Security Profile Group",
                "Antivirus",
                "IPS Sensor",
                "Web Filter",
                "Application List",
                "SSL/SSH Profile",
                "Source Profile Type",
                "Source Profile Group",
                "Profile Protocol Options",
                "Unresolved Security Profiles",
                "Security Profile References",
                "Security Profile Reference Statuses",
                 "Unresolved Security Profile References",
                 "Security Profile Semantics Review",
                 "URL Categories",
                 "Security Profile Groups",
                 "Antivirus Profiles",
                 "Vulnerability Profiles",
                 "Antispyware Profiles",
                 "URL Filtering Profiles",
                 "File Blocking Profiles",
                 "Wildfire Analysis Profiles",
                 "Data Filtering Profiles",
                 "Unresolved Source Zones",
                 "Unresolved Destination Zones",
                 "Inspection Mode",
                "Effective Inspection Mode",
                "ZTNA Status",
                "Effective ZTNA Status",
                "ZTNA EMS Tags",
                "Timeout Send RST",
                "Effective Timeout Send RST",
                "Auto ASIC Offload",
                "Effective Auto ASIC Offload",
                "NP Acceleration",
                "Effective NP Acceleration",
                "Port Preserve",
                "Effective Port Preserve",
                "Policy Expiry", "Effective Policy Expiry", "Policy Expiry Date", "Policy Expiry Date UTC",
                "Schedule Timeout", "Effective Schedule Timeout",
                "IPv4 Reputation Direction", "Effective IPv4 Reputation Direction",
                "IPv6 Reputation Direction", "Effective IPv6 Reputation Direction",
                "IPv4 Reputation Minimum", "Effective IPv4 Reputation Minimum",
                "IPv6 Reputation Minimum", "Effective IPv6 Reputation Minimum",
                "Match VIP", "Effective Match VIP", "Match VIP Only", "Effective Match VIP Only",
                "Additional Settings",
                "Extraction Status",
                 "Manual Review",
                 "Review Reasons",
                 "Description",
            ),
            rows,
        )

        # Keep title/note/header visible and retain the four most useful
        # identifier columns while scrolling horizontally.
        sheet.freeze_panes = "E4"

    def _build_cisco_acp(self, workbook: Any) -> None:
        self._table_sheet(
            workbook,
            "Cisco ACP",
            (
                "Rule #", "Name", "Source Rule ID", "Source UUID", "From Zone", "To Zone",
                "Source", "Destination", "Destination Services", "Source Ports", "VLAN Criteria",
                "URL Categories", "Applications", "Users", "IPS Profiles", "File Policies",
                "Variable Sets", "Action", "Enabled", "Migration Status", "Manual Review",
                "Review Reasons", "Additional Settings",
            ),
            (
                (
                    index, item.name, item.source_rule_id, item.source_uuid, item.from_zone, item.to_zone,
                    item.source, item.destination, item.service, item.source_ports, item.vlan_criteria,
                    item.url_categories, item.applications, item.source_users,
                    item.vulnerability_profiles, item.file_blocking_profiles, item.variable_sets,
                    item.source_action, self._optional_bool_literal(not bool(item.disabled)),
                    item.migration_status, self._optional_bool_literal(item.requires_manual_review),
                    item.review_reasons, self._format_settings(self._policy_source_settings(item)),
                )
                for index, item in enumerate(self.ir.policies, 1)
                if item.source_context and str(item.source_context).startswith("fmc:")
            ),
            empty_note="No FMC Access Control Policy rules were extracted.",
            subtitle="Cisco FMC ACP fields are exported separately so source-port, VLAN, URL, and inspection references remain independent.",
        )

    def _build_default_security_rules(self, workbook: Any) -> None:
        rows = [
            (
                index, rule.source_rule_id, rule.name, rule.source_context,
                rule.rulebase_position, rule.source_order, rule.action,
                self._optional_bool_literal(rule.disabled),
                self._optional_bool_literal(rule.log_start),
                self._optional_bool_literal(rule.log_end), rule.log_setting,
                rule.schedule, rule.security_profile_groups,
                rule.antivirus_profiles, rule.vulnerability_profiles,
                rule.antispyware_profiles, rule.url_filtering_profiles,
                rule.file_blocking_profiles, rule.wildfire_analysis_profiles,
                rule.data_filtering_profiles, rule.tags, rule.group_tag,
                rule.source_user, rule.source_hip, rule.destination_hip,
                rule.icmp_unreachable, rule.negate_source, rule.negate_destination,
                self._format_settings(rule.source_options), rule.description,
                rule.migration_status,
                self._optional_bool_literal(rule.requires_manual_review),
                rule.review_reasons,
            )
            for index, rule in enumerate(self.ir.default_security_rules, 1)
        ]
        self._table_sheet(
            workbook,
            "Default Security Rules",
            (
                "Rule #", "Source Rule ID", "Name", "Source Context",
                "Rulebase Position", "Source Order", "Action", "Disabled",
                "Log Start", "Log End", "Log Setting", "Schedule",
                "Security Profile Groups", "Antivirus Profiles",
                "Vulnerability Profiles", "Antispyware Profiles",
                "URL Filtering Profiles", "File Blocking Profiles",
                "Wildfire Analysis Profiles", "Data Filtering Profiles", "Tags",
                "Group Tag", "Source User", "Source HIP", "Destination HIP",
                "ICMP Unreachable", "Negate Source", "Negate Destination", "Options",
                "Description", "Migration Status", "Manual Review", "Review Reasons",
            ),
            rows,
            empty_note="No configured PAN-OS default security-rule overrides were extracted.",
            subtitle="Configured PAN-OS default-rule overrides; untouched built-ins remain in source evidence.",
        )

    def _build_local_in_policies(self, workbook: Any) -> None:
        rows = []
        items = getattr(self.ir, "local_in_policies", None)
        if items is None:
            items = self._vendor_extension_value("local_in_policies") or []
        for item in sorted(items, key=lambda value: value.source_order):
            attrs = item.source_attributes
            rows.append((
                item.source_context,
                item.source_id,
                attrs.get("address_family") or ("ipv6" if item.family.endswith("ipv6") else "ipv4"),
                item.source_order,
                item.enabled,
                attrs.get("action"),
                item.effective_action,
                getattr(item, "interface", None) or attrs.get("intf"),
                getattr(item, "source", None) or attrs.get("srcaddr"),
                getattr(item, "destination", None) or attrs.get("dstaddr"),
                getattr(item, "service", None) or attrs.get("service"),
                getattr(item, "service_exclusions", None) or attrs.get("service_exclusions"),
                attrs.get("schedule"),
                getattr(item, "protocol", None) or attrs.get("protocol"),
                getattr(item, "protocol_exclusions", None) or attrs.get("protocol_exclusions"),
                attrs.get("junos_host_inbound_scope"),
                attrs.get("junos_host_inbound_kind"),
                attrs.get("comments"),
                item.migration_status,
                item.requires_manual_review,
                item.review_reasons,
                self._format_settings(attrs),
            ))
        self._table_sheet(
            workbook,
            "Local-In Policies",
            (
                "Source VDOM", "Rule ID", "Address Family", "Source Order", "Enabled",
                "Configured Action", "Effective Action", "Interface", "Source Address",
                "Destination Address", "Service", "Service Exclusions", "Schedule", "Protocol", "Protocol Exclusions",
                "Host-Inbound Scope", "Host-Inbound Kind", "Comments",
                "Migration Status", "Manual Review", "Review Reasons", "Additional Settings",
            ),
            rows,
            subtitle="Typed FortiGate control-plane local-in policy inventory; not forwarding policy intent.",
        )

    def _build_security_policies(self, workbook: Any) -> None:
        profile_fields = (
            "av_profile", "casb_profile", "cifs_profile", "diameter_filter_profile",
            "dlp_profile", "dnsfilter_profile",
            "emailfilter_profile", "file_filter_profile", "icap_profile", "ips_sensor",
            "ips_voip_filter", "webfilter_profile", "videofilter_profile", "voip_profile",
            "sctp_filter_profile", "ssh_filter_profile", "virtual_patch_profile",
            "profile_group", "profile_type",
        )
        dedicated_fields = (
            ("Schedule", "schedule"),
            ("Users", "users"),
            ("Groups", "groups"),
            ("FSSO Groups", "fsso_groups"),
            ("Source Address Negate", "srcaddr_negate"),
            ("Destination Address Negate", "dstaddr_negate"),
            ("IPv6 Source Address Negate", "srcaddr6_negate"),
            ("IPv6 Destination Address Negate", "dstaddr6_negate"),
            ("Service Negate", "service_negate"),
            ("Profile Type", "profile_type"),
            ("Profile Group", "profile_group"),
            ("Profile Protocol Options", "profile_protocol_options"),
            ("Internet Service Enable/Status", "internet_service"),
            ("Internet Service Negate", "internet_service_negate"),
            ("Internet Service Custom", "internet_service_custom"),
            ("Internet Service Custom Groups", "internet_service_custom_group"),
            ("Internet Service Groups", "internet_service_group"),
            ("Internet Service Names", "internet_service_name"),
            ("Internet Service Source Enable/Status", "internet_service_src"),
            ("Internet Service Source Negate", "internet_service_src_negate"),
            ("Internet Service Source Custom", "internet_service_src_custom"),
            ("Internet Service Source Custom Groups", "internet_service_src_custom_group"),
            ("Internet Service Source Groups", "internet_service_src_group"),
            ("Internet Service Source Names", "internet_service_src_name"),
            ("IPv6 Internet Service Enable/Status", "internet_service6"),
            ("IPv6 Internet Service Negate", "internet_service6_negate"),
            ("IPv6 Internet Service Custom", "internet_service6_custom"),
            ("IPv6 Internet Service Custom Groups", "internet_service6_custom_group"),
            ("IPv6 Internet Service Groups", "internet_service6_group"),
            ("IPv6 Internet Service Names", "internet_service6_name"),
            ("IPv6 Internet Service Source Enable/Status", "internet_service6_src"),
            ("IPv6 Internet Service Source Negate", "internet_service6_src_negate"),
            ("IPv6 Internet Service Source Custom", "internet_service6_src_custom"),
            ("IPv6 Internet Service Source Custom Groups", "internet_service6_src_custom_group"),
            ("IPv6 Internet Service Source Groups", "internet_service6_src_group"),
            ("IPv6 Internet Service Source Names", "internet_service6_src_name"),
        )
        rows = []
        for item in sorted(
            self.ir.vendor_extensions.fortios.security_policies,
            key=lambda value: value.source_order,
        ):
            attrs = item.source_attributes
            profiles = {
                key: attrs[key]
                for key in profile_fields
                if attrs.get(key) not in (None, "", [])
            }
            rows.append((
                item.source_context,
                item.source_id,
                item.name,
                item.source_order,
                item.enabled,
                attrs.get("action"),
                item.effective_action,
                attrs.get("srcintf"),
                attrs.get("dstintf"),
                attrs.get("srcaddr"),
                attrs.get("srcaddr6"),
                attrs.get("dstaddr"),
                attrs.get("dstaddr6"),
                attrs.get("service"),
                attrs.get("application"),
                attrs.get("app_category"),
                attrs.get("app_group"),
                attrs.get("application_list"),
                self._format_settings(profiles),
                attrs.get("ssl_ssh_profile"),
                attrs.get("logtraffic"),
                attrs.get("nat46"),
                attrs.get("nat64"),
                attrs.get("comments"),
                *(attrs.get(field) for _, field in dedicated_fields),
                item.migration_status,
                item.requires_manual_review,
                item.review_reasons,
                self._format_settings(attrs),
            ))
        self._table_sheet(
            workbook,
            "NGFW Security Policies",
            (
                "Source VDOM", "Rule ID", "Name", "Source Order", "Enabled", "Action", "Effective Action",
                "Source Interface", "Destination Interface", "Source Address", "Source IPv6 Address",
                "Destination Address", "Destination IPv6 Address", "Services", "Applications",
                "Application Categories", "Application Groups", "Application List",
                "Security Profile References", "SSL Inspection Reference", "Logging",
                "NAT46", "NAT64", "Comments",
                *(header for header, _ in dedicated_fields),
                "Migration Status", "Manual Review",
                "Review Reasons", "Additional Settings",
            ),
            rows,
            subtitle="Typed FortiGate policy-based NGFW security-policy inventory; not portable firewall policy intent.",
        )

    def _build_firewall_policy_source_settings(self, workbook: Any) -> None:
        items = (
            []
            if self.extraction is None
            else [
                item
                for item in self.extraction.inventory_items
                if item.source_path == "firewall policy"
            ]
        )

        self._table_sheet(
            workbook,
            "Firewall Policy Source Settings",
            (
                "Source Policy ID",
                "Policy Name",
                "Operation",
                "Setting",
                "Ordered Source Values",
            ),
            (
                (
                    item.source_id,
                    item.name,
                    command.operation,
                    command.key,
                    json.dumps(list(command.values), ensure_ascii=False),
                )
                for item in items
                for command in item.commands
            ),
            empty_note="No FortiGate firewall policy source commands were retained.",
            subtitle=(
                "Sanitized, ordered FortiGate policy commands retained for audit. "
                "This extraction-only detail is not consumed by target generators."
            ),
        )

    def _build_address_group_tags(self, workbook: Any) -> None:
        rows = [
            (group.name, group.source_section, group.address_family, entry.name,
             entry.category, entry.tags, entry.migration_status,
             self._optional_bool_literal(entry.requires_manual_review),
             self._format_settings(entry.source_attributes))
            for group in self.ir.address_groups
            for entry in group.source_tagging_entries
        ]
        self._table_sheet(workbook, "Address Group Tags", (
            "Group Name", "Source Section", "Address Family", "Tag Entry",
            "Category", "Tags", "Extraction Status", "Manual Review",
            "Additional Settings",
        ), rows)

    def _build_ztna_providers(self, workbook: Any) -> None:
        """
        Export ZTNA / endpoint-posture provider dependencies.

        Policy IDs and EMS tags shown here are observed elsewhere in the
        same source configuration. They are not asserted to belong to a
        specific provider unless explicit source correlation exists.
        """

        # Collect all policies that contain ZTNA intent or EMS tag references.
        ztna_policies = [
            policy
            for policy in self.ir.policies
            if (
                policy.source_ztna_status == "enable"
                or policy.source_ztna_ems_tags
            )
        ]

        # Preserve policy order while removing duplicates.
        observed_policy_ids = list(
            dict.fromkeys(
                policy.source_rule_id or policy.name
                for policy in ztna_policies
                if policy.source_rule_id or policy.name
            )
        )

        # Collect all ZTNA EMS tags referenced by those policies.
        observed_ems_tags = list(
            dict.fromkeys(
                tag
                for policy in ztna_policies
                for tag in policy.source_ztna_ems_tags
                if tag
            )
        )

        rows = [
            (
                item.name,
                item.source_vendor or self.ir.metadata.source_vendor,
                item.source_id,
                item.provider_type,
                self._optional_bool_literal(item.enabled),
                item.source_serial,
                item.source_tenant_id,
                self._optional_bool_literal(
                    item.source_cloud_authentication
                ),
                item.verifying_ca,
                item.verified_cn,
                item.capabilities,
                observed_policy_ids,
                observed_ems_tags,
                item.migration_status,
                self._optional_bool_literal(
                    item.requires_manual_review
                ),
                item.migration_instruction,
                self._format_settings(
                    item.source_attributes
                ),
            )
            for item in self.ir.ztna_providers
        ]

        self._table_sheet(
            workbook,
            "ZTNA Providers",
            (
                "Name",
                "Source Vendor",
                "Source ID",
                "Provider Type",
                "Enabled",
                "Source Serial",
                "Tenant ID",
                "Cloud Authentication",
                "Verifying CA",
                "Verified CN",
                "Capabilities",
                "ZTNA Policy IDs (Observed)",
                "ZTNA EMS Tags (Observed)",
                "Extraction Status",
                "Manual Review",
                "Migration Instruction",
                "Additional Settings",
            ),
            rows,
            empty_note=(
                "No meaningful ZTNA / endpoint-posture providers "
                "were extracted from the source configuration."
            ),
            subtitle=(
                "Source ZTNA and endpoint-posture dependencies retained "
                "for migration review. Policy IDs and EMS tags are observed "
                "in the same source configuration and are not automatically "
                "claimed as an exact mapping to an individual provider. "
                "Provider-specific configuration is not consumed by target "
                "generators."
            ),
        )

    def _build_ip_pools(self, workbook: Any) -> None:
        rows = [
            (
                item.name,
                item.address_family,
                item.routing_instance,
                item.pool_type,
                item.addresses,
                self._format_ip_pool_ranges(item.address_ranges),
                item.start_ip,
                item.end_ip,
                item.source_start_ip,
                item.source_end_ip,
                item.source_prefix6,
                item.start_port,
                item.end_port,
                item.associated_interface,
                self._optional_bool_literal(self._object_extension_value(item, "arp_reply")),
                self._object_extension_value(item, "arp_interface"),
                self._optional_bool_literal(item.permit_any_host),
                item.excluded_ips,
                self._object_extension_value(item, "block_size"),
                self._object_extension_value(item, "blocks_per_user"),
                self._object_extension_value(item, "pba_timeout"),
                self._object_extension_value(item, "pba_interim_log"),
                self._object_extension_value(item, "ports_per_user"),
                self._optional_bool_literal(self._object_extension_value(item, "privileged_port_use_pba")),
                self._optional_bool_literal(self._object_extension_value(item, "nat64")),
                self._optional_bool_literal(self._object_extension_value(item, "add_nat64_route")),
                self._optional_bool_literal(self._object_extension_value(item, "nat46")),
                self._optional_bool_literal(self._object_extension_value(item, "add_nat46_route")),
                self._object_extension_value(item, "client_prefix_length"),
                self._optional_bool_literal(self._object_extension_value(item, "include_subnet_broadcast")),
                self._object_extension_value(item, "cgn_block_size"),
                self._object_extension_value(item, "cgn_client_start_ip"),
                self._object_extension_value(item, "cgn_client_end_ip"),
                self._object_extension_value(item, "cgn_client_ipv6_shift"),
                self._optional_bool_literal(self._object_extension_value(item, "cgn_fixed_allocation")),
                self._optional_bool_literal(self._object_extension_value(item, "cgn_overload")),
                self._object_extension_value(item, "cgn_port_start"),
                self._object_extension_value(item, "cgn_port_end"),
                self._optional_bool_literal(self._object_extension_value(item, "cgn_spa")),
                self._object_extension_value(item, "utilization_alarm_clear"),
                self._object_extension_value(item, "utilization_alarm_raise"),
                self._object_extension_value(item, "tcp_session_quota"),
                self._object_extension_value(item, "udp_session_quota"),
                self._object_extension_value(item, "icmp_session_quota"),
                ", ".join(item.source_explicit_fields),
                self._format_settings(item.source_effective_settings),
                item.migration_status,
                self._optional_bool_literal(item.requires_manual_review),
                item.audit_note,
                self._format_settings(item.source_attributes),
                item.description,
                item.source_uuid, item.source_origin, item.checkpoint_pool_object_type,
                item.checkpoint_network_references, item.checkpoint_network_group_references,
                item.checkpoint_address_range_references, item.checkpoint_gateway_references,
                item.checkpoint_member_assignments, item.checkpoint_applicability,
                item.checkpoint_precedence, item.checkpoint_vpn_scope,
                self._optional_bool_literal(item.checkpoint_mep),
                item.source_context,
            )
            for item in self.ir.ip_pools
        ]
        self._table_sheet(
            workbook,
            "IP Pools",
            (
                "Name", "Address Family", "Routing Instance", "Type", "Addresses", "Address Ranges",
                "Start IP", "End IP", "Source Start IP",
                "Source End IP", "Source Prefix6", "Start Port", "End Port", "Associated Interface",
                "ARP Reply", "ARP Interface", "Permit Any Host", "Excluded IPs",
                "Block Size", "Blocks Per User", "PBA Timeout", "PBA Interim Log",
                "Ports Per User", "Privileged Port Uses PBA", "NAT64", "Add NAT64 Route",
                "NAT46", "Add NAT46 Route", "Client Prefix Length",
                "Include Subnet/Broadcast", "CGN Block Size", "CGN Client Start IP",
                "CGN Client End IP", "CGN Client IPv6 Shift", "CGN Fixed Allocation",
                "CGN Overload", "CGN Port Start", "CGN Port End", "CGN SPA",
                "Utilization Alarm Clear", "Utilization Alarm Raise", "TCP Session Quota",
                "UDP Session Quota", "ICMP Session Quota", "Source Explicit Fields",
                "Effective Source Settings", "Extraction Status",
                "Manual Review", "Review Reason", "Additional Settings", "Description",
                "Source UUID", "Source Origin", "Check Point Pool Object Type",
                "Check Point Networks", "Check Point Network Groups", "Check Point Address Ranges",
                "Check Point Gateways", "Check Point Member Assignments", "Check Point Applicability",
                "Check Point Precedence", "Check Point VPN Scope", "Check Point MEP", "Source VDOM",
            ),
            rows,
        )

    def _build_ipv6_eh_filter(self, workbook: Any) -> None:
        items = [
            item for item in (self.extraction.inventory_items if self.extraction else [])
            if item.source_path == "firewall ipv6-eh-filter"
        ]
        rows = []
        for item in items:
            attrs = item.source_attributes
            rows.append((
                item.source_context,
                attrs.get("auth"),
                attrs.get("dest_opt"),
                attrs.get("fragment"),
                attrs.get("hop_opt"),
                attrs.get("no_next"),
                attrs.get("routing"),
                attrs.get("hdopt_type", []),
                attrs.get("routing_type"),
                ", ".join(attrs.get("source_explicit_fields", [])),
                self._format_settings(attrs.get("source_effective_settings", {})),
                item.status.value,
                self._optional_bool_literal(item.requires_manual_review),
                "; ".join(attrs.get("review_reasons", [])),
                self._format_settings(attrs.get("additional_settings", {})),
            ))
        self._table_sheet(
            workbook,
            "IPv6 EH Filter",
            (
                "Source Context", "Authentication Header Blocking",
                "Destination Options Blocking", "Fragment Header Blocking",
                "Hop-by-Hop Blocking", "No Next Header Blocking",
                "Routing Header Blocking", "Hop/Destination Option Types",
                "Routing Types", "Source Explicit Fields",
                "Effective Source Settings", "Extraction Status", "Manual Review",
                "Review Reason", "Additional Settings",
            ),
            rows,
            empty_note="No IPv6 extension-header filter was extracted.",
            subtitle=(
                "Typed FortiOS IPv6 extension-header blocking settings retained as "
                "source-only inventory; enable means blocking and is not portable "
                "target policy intent."
            ),
        )

    def _build_multicast_policies(self, workbook: Any) -> None:
        rows = [
            (
                index, item.source_id, item.source_uuid, item.name,
                item.address_family, item.source_context, item.source_order,
                self._optional_bool_literal(item.enabled), item.action,
                item.source_interface, item.destination_interface,
                item.source_addresses, item.destination_addresses,
                item.protocol_number, item.destination_port_start,
                item.destination_port_end, item.utm_status, item.ips_sensor,
                item.logtraffic, item.traffic_shaper, item.auto_asic_offload,
                item.source_snat, item.source_snat_ip, item.source_dnat,
                item.comments,
                item.migration_status,
                self._optional_bool_literal(item.requires_manual_review),
                item.review_reasons, self._format_settings(item.source_attributes),
            )
            for index, item in enumerate(self.ir.multicast_policies, 1)
        ]
        self._table_sheet(
            workbook,
            "Multicast Policies",
            (
                "Rule #", "Source Policy ID", "Source UUID", "Name",
                "Address Family", "Source Context", "Source Order", "Enabled",
                "Action", "Source Interface", "Destination Interface",
                "Source Addresses", "Destination Addresses", "Protocol Number",
                "Destination Start Port", "Destination End Port", "UTM Status",
                "IPS Sensor", "Log Traffic", "Traffic Shaper", "Auto ASIC Offload",
                "Source SNAT", "Source SNAT IP", "Source DNAT", "Comments",
                "Migration Status", "Manual Review", "Review Reasons",
                "Additional Source Settings",
            ),
            rows,
        )

    def _build_checkpoint_access_rule_sheet(self, workbook: Any) -> None:
        self._table_sheet(
            workbook,
            "Checkpoint Access Rules",
            (
                "Name", "UID", "Rule Number", "Source Context", "Domain", "Package", "Layer",
                "Section Path", "Enabled", "Source", "Destination", "VPN", "Services",
                "Applications", "Access Roles", "Action", "Track", "Time", "Install On",
                "Source Negated", "Destination Negated", "Service Negated", "Content",
                "Content Negated", "Inline Layer", "Parent Layer", "Parent Rule UID", "Comments",
                "Migration Status", "Manual Review", "Review Reasons", "Additional Settings",
            ),
            (
                (
                    rule.name, rule.source_uuid, rule.rule_number, rule.source_context, rule.domain,
                    rule.package, rule.layer, rule.section_path,
                    self._optional_bool_literal(rule.enabled), rule.source, rule.destination,
                    rule.vpn, rule.services, rule.applications, rule.access_roles, rule.action,
                    self._format_settings(rule.track) if isinstance(rule.track, dict) else rule.track,
                    rule.time, rule.install_on, self._optional_bool_literal(rule.source_negated),
                    self._optional_bool_literal(rule.destination_negated),
                    self._optional_bool_literal(rule.service_negated), rule.content,
                    self._optional_bool_literal(rule.content_negated), rule.inline_layer_reference,
                    rule.parent_layer, rule.parent_rule_uid, rule.comments, rule.migration_status,
                    self._optional_bool_literal(rule.requires_manual_review), rule.review_reasons,
                    self._format_settings(rule.source_attributes),
                )
                for rule in self._vendor_extension_value("checkpoint_access_rules") or []
            ),
            empty_note="No Check Point access rules were extracted.",
            subtitle="Complete Check Point rule dimensions retained separately from portable IRPolicy semantics.",
        )

    def _build_nat_rules(self, workbook: Any) -> None:
        rows = (
            (
                index, item.sequence, self._format_nat_source_section(item), item.source_rule_id,
                item.name, item.source_rule_set, item.sequence,
                item.type, item.source_origin, item.nat_family,
                item.original_address_family, item.translated_address_family,
                f"{item.protocol_name or ''}/{item.protocol_number or ''}".strip("/"),
                self._format_nat_ports(item.original_source_ports),
                self._format_nat_ports(item.original_destination_ports),
                self._format_nat_ports(item.translated_source_ports),
                self._format_nat_ports(item.translated_destination_ports),
                item.source_port_behavior, item.install_translation_route,
                self._optional_bool_literal(getattr(item.runtime_behavior, "pcp_inbound", None)),
                self._optional_bool_literal(getattr(item.runtime_behavior, "pcp_outbound", None)),
                getattr(item.runtime_behavior, "pcp_pool_names", []),
                self._optional_bool_literal(getattr(item.runtime_behavior, "permit_stun_host", None)),
                self._optional_bool_literal(getattr(item.runtime_behavior, "rtp_nat", None)),
                getattr(item.runtime_behavior, "rtp_addresses", []),
                item.source_policy_reference,
                item.source_policy_uuid, self._optional_bool_literal(item.enabled),
                item.source_from_interfaces,
                self._format_nat_attachment_values(
                    self._object_extension_value(item, "source_attachments"), "name"
                ),
                self._format_nat_attachment_values(
                    self._object_extension_value(item, "source_attachments"), "kind"
                ),
                item.from_zone,
                item.from_zone,
                item.source_to_interfaces,
                self._format_nat_attachment_values(
                    self._object_extension_value(item, "destination_attachments"), "name"
                ),
                self._format_nat_attachment_values(
                    self._object_extension_value(item, "destination_attachments"), "kind"
                ),
                item.to_zone,
                item.to_zone,
                self._nat_attachment_resolution_status(item),
                item.source, item.destination,
                 item.from_routing_instances, item.to_routing_instances,
                 self._format_nat_address_ranges(item.address_range_mappings, "original_start"),
                 self._format_nat_address_ranges(item.address_range_mappings, "original_end"),
                 self._format_nat_address_ranges(item.address_range_mappings, "translated_start"),
                 self._format_nat_address_ranges(item.address_range_mappings, "translated_end"),
                 item.services,
                 item.internet_services, item.source_translation_mode,
                 item.destination_translation_mode,
                self._optional_bool_literal(item.source_translation_bidirectional),
                self._format_settings(item.source_translation_fallback.model_dump(mode="json"))
                if item.source_translation_fallback else None,
                item.translated_source_address_references,
                item.translated_destination_address_references,
                item.source_translation_address_selection.address_source
                if item.source_translation_address_selection else None,
                item.source_translation_address_selection.interface
                if item.source_translation_address_selection else None,
                item.source_translation_address_selection.ipv4_addresses
                if item.source_translation_address_selection else [],
                item.source_translation_address_selection.ipv6_addresses
                if item.source_translation_address_selection else [],
                item.source_translation_address_selection.floating_ips
                if item.source_translation_address_selection else [],
                item.source_pool_references, item.translated_sources,
                self._object_extension_value(item, "source_pool_type"),
                self._object_extension_value(item, "source_pool_excluded_ips"),
                self._optional_bool_literal(self._object_extension_value(item, "source_pool_permit_any_host")),
                self._object_extension_value(item, "source_pool_original_start_ip"),
                self._object_extension_value(item, "source_pool_original_end_ip"),
                self._object_extension_value(item, "source_vip_reference"),
                self._object_extension_value(item, "source_vip_group_reference"),
                self._object_extension_value(item, "source_vip_type"),
                self._optional_bool_literal(self._object_extension_value(item, "source_vip_enabled")),
                self._optional_bool_literal(self._object_extension_value(item, "source_vip_nat_source_vip")),
                self._object_extension_value(item, "source_vip_filters"),
                self._object_extension_value(item, "source_vip_interface_filters"),
                self._object_extension_value(item, "source_vip_services"),
                self._object_extension_value(item, "source_vip_port_mapping_type"),
                item.translated_destinations, item.original_destination_port,
                item.destination_protocol, item.translated_port,
                self._object_extension_value(item, "source_policy_fixed_port"),
                self._object_extension_value(item, "source_policy_nat46"),
                self._object_extension_value(item, "source_policy_nat64"),
                self._object_extension_value(item, "source_policy_nat_inbound"),
                self._object_extension_value(item, "source_policy_nat_outbound"),
                self._object_extension_value(item, "source_policy_nat_ip"),
                self._object_extension_value(item, "source_policy_match_vip"),
                self._object_extension_value(item, "source_policy_match_vip_only"),
                 item.migration_status,
                 self._optional_bool_literal(item.requires_manual_review),
                 item.review_reasons,
                 item.description, self._optional_bool_literal(item.identity),
                 self._optional_bool_literal(item.exemption),
                 item.source_rule_id, item.source_policy_uuid, item.services,
                 item.translated_services, item.source_attributes.get("install-on", item.source_attributes.get("install_on")),
                 item.source_attributes.get("method", item.source_attributes.get("hide-behind", item.source_attributes.get("hide_behind"))),
                 item.source_attributes.get("checkpoint-nat-origin"),
                 (item.source_attributes.get("checkpoint-provenance") or {}).get("section-path")
                 if isinstance(item.source_attributes.get("checkpoint-provenance"), dict) else None,
                 self._format_settings(item.source_attributes.get("checkpoint-source-nat-method-resolution") or {}),
                 self._optional_bool_literal(item.source_attributes.get("checkpoint-ordering-barrier")),
                self._format_settings(item.source_attributes),
                item.source_context,
            )
            for index, item in enumerate(self.ir.nat_rules, 1)
        )
        self._table_sheet(
            workbook,
            "NAT Rules",
            (
                "Rule #", "Sequence", "Source Section", "Source Rule ID", "Name", "Rule Set", "Rule Position",
                "Type", "Source Origin", "NAT Family",
                "Original Address Family", "Translated Address Family", "Protocol / Number",
                "Original Source Port", "Original Destination Port", "Translated Source Port",
                "Translated Destination Port", "Source Port Behavior", "Install Translation Route",
                "PCP Inbound", "PCP Outbound", "PCP Pools", "STUN Any Host", "RTP NAT",
                "RTP Addresses", "Source Policy ID", "Source Policy UUID",
                 "Enabled", "Source Interface", "Source Attachments", "Source Attachment Types",
                 "Canonical Source Zones", "From Zone", "Destination Interface", "Destination Attachments",
                 "Destination Attachment Types", "Canonical Destination Zones", "To Zone",
                 "Attachment Resolution Status",
                 "Original Source", "Original Destination",
                 "From Routing Instance", "To Routing Instance",
                 "Original Range Start", "Original Range End", "Translated Range Start",
                 "Translated Range End", "Services",
                 "Internet Services", "Source Translation Mode", "Destination Translation Mode",
                 "Static NAT Bi-directional", "Source Translation Fallback",
                 "Translated Source References", "Translated Destination References",
                 "Source Translation Address Type", "Source Translation Interface",
                 "Source Translation Interface IPv4", "Source Translation Interface IPv6",
                 "Source Translation Floating IPs", "IP Pool",
                "Translated Source", "IP Pool Type", "Pool Excluded IPs", "Pool Full Cone",
                "Pool Source Start IP", "Pool Source End IP", "VIP", "VIP Group",
                "VIP Type", "VIP Enabled", "VIP NAT Source VIP", "VIP Source Filters",
                "VIP Interface Filters", "VIP Services", "VIP Port Mapping Type",
                "Translated Destination", "Legacy Original Destination Port", "Legacy Destination Protocol",
                "Legacy Translated Port", "Policy Fixed Port", "Policy NAT46", "Policy NAT64",
                "Policy NAT Inbound", "Policy NAT Outbound", "Policy NAT IP",
                 "Policy Match VIP", "Policy Match VIP Only", "Migration Status",
                 "Manual Review", "Review Reasons", "Description", "Identity", "Exemption",
                 "Source Rule Number", "Source Rule UID", "Original Service", "Translated Service",
                 "Install On", "Source Translation Method", "Check Point NAT Origin", "Check Point Section",
                 "Check Point Source NAT Evidence", "Check Point Ordering Barrier", "Additional Settings", "Source VDOM",
            ),
            rows,
        )

    def _build_pbf_rules(self, workbook: Any) -> None:
        self._table_sheet(
            workbook,
            "PBF Rules",
            (
                "Rule #", "Name", "Source Context", "Rulebase Position", "Source Order",
                "From Zones", "From Interfaces", "To", "Source", "Destination",
                "Source User", "Applications", "Services", "Action", "Forward to VSYS",
                "Egress Interface", "Next Hop Type", "Next Hop", "Next VR", "Monitor Profile",
                "Schedule", "Negate Source", "Negate Destination", "Symmetric Return",
                "Symmetric Return Next Hops",
                "Monitor IP", "Monitor Enabled", "Disable if Unreachable", "Enabled",
                "Migration Status", "Manual Review", "Review Reasons", "Description",
                "Priority", "Protocol", "Destination Port", "Routing Table",
                "Table Routes", "Table Next Hop", "Table Output Interface",
            ),
            (
                (
                    index, rule.name, rule.source_context, rule.rulebase_position,
                    rule.source_order, rule.from_zone, rule.from_interface, rule.to,
                    rule.source, rule.destination, rule.source_user, rule.application,
                    rule.service, rule.action, rule.forward_to_vsys, rule.egress_interface,
                    rule.next_hop_type, rule.next_hop, rule.next_vr, rule.monitor_profile,
                    rule.schedule, self._optional_bool_literal(rule.source_negated),
                    self._optional_bool_literal(rule.destination_negated),
                    self._optional_bool_literal(
                        rule.symmetric_return.enabled
                        if rule.symmetric_return is not None
                        else rule.enforce_symmetric_return
                    ),
                    rule.symmetric_return.next_hop_addresses
                    if rule.symmetric_return is not None else [],
                    rule.monitor_ip, self._optional_bool_literal(rule.monitor_enabled),
                    self._optional_bool_literal(rule.disable_if_unreachable),
                    self._optional_bool_literal(rule.enabled), rule.migration_status,
                    self._optional_bool_literal(rule.requires_manual_review),
                    rule.review_reasons, rule.description, rule.priority, rule.protocol,
                    rule.destination_port, rule.routing_table,
                    self._format_pbr_table_routes(rule.source_attributes.get("table_routes", [])),
                    rule.table_next_hop,
                    rule.table_output_interface,
                )
                for index, rule in enumerate(self.ir.pbf_rules, 1)
            ),
            empty_note="No policy-based forwarding rules were extracted.",
            subtitle="Policy-based forwarding inventory; source route tables remain separate from static routes.",
        )

    def _build_virtual_ips(self, workbook: Any) -> None:
        rows = [
            (
                item.name,
                item.address_family,
                item.source_id,
                item.source_uuid,
                self._object_extension_value(item, "vip_type"),
                self._optional_bool_literal(item.enabled),
                item.external_ip,
                item.external_addresses,
                item.external_interface,
                item.mapped_ips,
                item.mapped_address,
                self._optional_bool_literal(item.port_forward),
                item.protocol,
                item.external_port,
                item.mapped_port,
                self._object_extension_value(item, "port_mapping_type"),
                self._optional_bool_literal(self._object_extension_value(item, "arp_reply")),
                self._object_extension_value(item, "gratuitous_arp_interval"),
                self._optional_bool_literal(self._object_extension_value(item, "nat_source_vip")),
                self._optional_bool_literal(self._object_extension_value(item, "nat44")),
                self._optional_bool_literal(self._object_extension_value(item, "nat46")),
                self._optional_bool_literal(self._object_extension_value(item, "nat64")),
                self._optional_bool_literal(self._object_extension_value(item, "nat66")),
                self._optional_bool_literal(self._object_extension_value(item, "add_nat46_route")),
                self._optional_bool_literal(self._object_extension_value(item, "add_nat64_route")),
                self._optional_bool_literal(self._object_extension_value(item, "ndp_reply")),
                self._object_extension_value(item, "ipv6_mapped_ip"),
                self._object_extension_value(item, "ipv6_mapped_port"),
                self._object_extension_value(item, "ipv4_mapped_ip"),
                self._object_extension_value(item, "ipv4_mapped_port"),
                self._object_extension_value(item, "embedded_ipv4_address"),
                item.source_filters,
                item.source_interface_filters,
                item.services,
                item.load_balance_method,
                self._object_extension_value(item, "server_type"),
                item.persistence,
                self._optional_bool_literal(self._object_extension_value(item, "http_redirect")),
                item.monitors,
                self._object_extension_value(item, "max_embryonic_connections"),
                item.color,
                item.description,
                self._object_extension_value(item, "h2_support"),
                self._object_extension_value(item, "h3_support"),
                self._object_extension_value(item, "http_multiplex"),
                self._object_extension_value(item, "ssl_mode"),
                item.ssl_certificate,
                self._object_extension_value(item, "ssl_algorithm"),
                self._object_extension_value(item, "ssl_min_version"),
                self._object_extension_value(item, "ssl_max_version"),
                self._object_extension_value(item, "ssl_server_algorithm"),
                self._object_extension_value(item, "ssl_server_min_version"),
                self._object_extension_value(item, "ssl_server_max_version"),
                self._object_extension_value(item, "ssl_pfs"),
                self._object_extension_value(item, "gslb_domain_name"),
                self._object_extension_value(item, "gslb_hostname"),
                ", ".join(item.source_explicit_fields),
                self._format_settings(item.source_effective_settings),
                self._format_settings(item.extra_settings),
                item.migration_status,
                self._optional_bool_literal(item.requires_manual_review),
                item.audit_note,
                item.source_context,
            )
            for item in self.ir.virtual_ips
        ]
        self._table_sheet(
            workbook,
            "Virtual IPs",
            (
                "Name", "Address Family", "Source ID", "Source UUID", "Type", "Enabled", "External IP",
                "External Address Objects", "External Interface", "Mapped IPs",
                "Mapped Address", "Port Forward", "Protocol", "External Port",
                "Mapped Port", "Port Mapping Type", "ARP Reply", "Gratuitous ARP Interval",
                "NAT Source VIP",
                "NAT44", "NAT46", "NAT64", "NAT66", "Add NAT46 Route",
                "Add NAT64 Route", "NDP Reply", "IPv6 Mapped IP", "IPv6 Mapped Port",
                "IPv4 Mapped IP", "IPv4 Mapped Port", "Embedded IPv4 Address",
                "Source Filters", "Source Interface Filters", "Services",
                "Load Balance Method", "Server Type", "Persistence", "HTTP Redirect",
                "Monitors", "Max Embryonic Connections", "Color", "Description",
                "H2 Support", "H3 Support", "HTTP Multiplex", "SSL Mode",
                "SSL Certificate", "SSL Algorithm", "SSL Min Version", "SSL Max Version",
                "SSL Server Algorithm", "SSL Server Min Version", "SSL Server Max Version",
                "SSL PFS", "GSLB Domain Name", "GSLB Hostname", "Source Explicit Fields",
                "Effective Source Settings", "Additional Settings", "Extraction Status",
                "Manual Review", "Review Reason", "Source VDOM",
            ),
            rows,
        )

    def _vip_nested_config_rows(self) -> Iterable[tuple[Any, ...]]:
        def walk(
            vip: Any,
            node: Any,
            parent_path: list[str],
        ) -> Iterable[tuple[Any, ...]]:
            if node.node_type == "config":
                config_path = [*parent_path, str(node.name)]
                object_name = None
            else:
                config_path = list(parent_path)
                object_name = str(node.name)

            if node.commands:
                for command in node.commands:
                    yield (
                        vip.name,
                        vip.address_family,
                        vip.source_context,
                        " / ".join(config_path),
                        node.node_type,
                        object_name,
                        command.operation,
                        command.key,
                        self._format_source_command_values(command.key, command.values),
                        "EXTRACT_ONLY",
                        "Yes",
                    )
            elif not node.children:
                yield (
                    vip.name,
                    vip.address_family,
                    vip.source_context,
                    " / ".join(config_path),
                    node.node_type,
                    object_name,
                    None,
                    None,
                    None,
                    "EXTRACT_ONLY",
                    "Yes",
                )

            for child in node.children:
                child_parent = (
                    config_path
                    if node.node_type == "config"
                    else [*config_path, str(node.name)]
                )
                yield from walk(vip, child, child_parent)

        for vip in self.ir.virtual_ips:
            for root in vip.nested_source_configs:
                yield from walk(vip, root, [])

    def _build_vip_nested_configuration(self, workbook: Any) -> None:
        self._table_sheet(
            workbook,
            "VIP Nested Configuration",
            (
                "VIP Name", "Address Family", "Source Context", "Config Path",
                "Node Type", "Object Name", "Operation", "Key", "Values",
                "Extraction Status", "Manual Review",
            ),
            self._vip_nested_config_rows(),
            empty_note="No nested VIP configuration was extracted from the source firewall.",
            subtitle=(
                "Nested FortiGate VIP configuration retained as sanitized "
                "extraction-only source data. These settings are not consumed "
                "by target generators."
            ),
        )

    def _build_vip_real_servers(self, workbook: Any) -> None:
        rows = [
            (
                vip.name,
                server.id,
                server.address_type,
                server.ip_address,
                server.address_reference,
                server.port,
                server.status,
                server.weight,
                server.holddown_interval,
                server.healthcheck,
                server.http_host,
                server.translate_host,
                server.max_connections,
                server.monitors,
                server.client_ip,
                server.migration_status,
                self._optional_bool_literal(server.requires_manual_review),
                server.audit_note,
                self._format_settings(server.source_attributes),
            )
            for vip in self.ir.virtual_ips
            for server in vip.real_servers
        ]
        self._table_sheet(
            workbook,
            "VIP Real Servers",
            (
                "VIP Name", "Server ID", "Address Type", "IP", "Address Object",
                "Port", "Status", "Weight", "Holddown Interval", "Health Check",
                "HTTP Host", "Translate Host", "Max Connections", "Monitors", "Client IP",
                "Extraction Status", "Manual Review", "Review Reason", "Additional Settings",
            ),
            rows,
        )

    def _build_vpn_tunnels(self, workbook: Any) -> None:
        rows = [
            (
                item.name,
                item.source_type,
                item.peer_address,
                item.local_interface,
                item.source_local_gateway_ipv4 or item.source_local_gateway,
                item.source_local_gateway_ipv6,
                item.source_remote_gateway_ipv4,
                item.source_remote_gateway_ipv6,
                item.source_remote_gateway_ddns,
                item.ike_version,
                item.source_mode,
                item.source_peer_type,
                item.source_auth_method,
                item.source_remote_auth_method,
                item.source_certificates,
                item.source_dh_groups,
                item.source_key_lifetime,
                item.source_nat_traversal,
                item.source_dpd_mode,
                item.source_dpd_retry_count,
                item.source_dpd_retry_interval,
                item.source_xauth_type,
                item.source_peer_id,
                item.source_local_id,
                item.source_local_id_type,
                self._optional_bool_literal(item.source_net_device),
                item.source_proposals,
                self._optional_bool_literal(item.source_mode_config),
                item.source_mode_config_allow_client_selector,
                self._optional_bool_literal(item.source_eap),
                item.source_eap_identity,
                item.source_auth_user,
                item.source_auth_user_group,
                item.unresolved_auth_user_groups,
                item.unresolved_interfaces,
                item.unresolved_certificates,
                item.source_client_ip_start,
                item.source_client_ip_end,
                (
                    f"{item.source_client_ip_start} - {item.source_client_ip_end}"
                    if item.source_client_ip_start and item.source_client_ip_end
                    else None
                ),
                item.source_dns_mode,
                item.source_split_include,
                item.source_split_exclude,
                item.source_ipv6_split_include,
                item.source_ipv6_split_exclude,
                item.source_backup_gateways,
                item.source_rekey,
                item.source_reauth,
                item.source_signature_hash_algorithms,
                (
                    "Configured / Redacted"
                    if item.has_psk or item.psk
                    else "Not configured"
                ),
                item.migration_status,
                self._optional_bool_literal(item.requires_manual_review),
                item.review_reasons,
                self._format_settings(item.source_attributes),
                item.ike_crypto_profile,
                item.ipsec_crypto_profile,
                item.description,
            )
            for item in self.ir.vpn_tunnels
        ]
        self._table_sheet(
            workbook,
            "VPN Tunnels",
            (
                "Name",
                "Type",
                "Peer Address",
                "Local Interface",
                "Local Gateway",
                "Local Gateway IPv6",
                "Remote Gateway IPv4",
                "Remote Gateway IPv6",
                "Remote Gateway DDNS",
                "IKE Version",
                "Mode",
                "Peer Type",
                "Authentication Method",
                "Remote Authentication Method",
                "Certificates",
                "DH Groups",
                "Key Lifetime",
                "NAT Traversal",
                "DPD Mode",
                "DPD Retry Count",
                "DPD Retry Interval",
                "XAuth Type",
                "Peer ID",
                "Local ID",
                "Local ID Type",
                "Net Device",
                "IKE Proposal",
                "Mode Config",
                "Mode Config Allow Client Selector",
                "EAP",
                "EAP Identity",
                "Auth User",
                "Auth User Group",
                "Unresolved Auth User Groups",
                "Unresolved Interfaces",
                "Unresolved Certificates",
                "Client IP Start",
                "Client IP End",
                "Client IP Range",
                "DNS Mode",
                "Split Include",
                "Split Exclude",
                "IPv6 Split Include",
                "IPv6 Split Exclude",
                "Backup Gateways",
                "Rekey",
                "Reauth",
                "Signature Hash Algorithms",
                "PSK",
                "Extraction Status",
                "Manual Review",
                "Review Reasons",
                "Additional Settings",
                "IKE Crypto Profile",
                "IPsec Crypto Profile",
                "Description",
            ),
            rows,
        )

    def _build_certificates(self, workbook: Any) -> None:
        extraction_timestamp = self.ir.metadata.migration_timestamp
        if extraction_timestamp.tzinfo is None:
            extraction_timestamp = extraction_timestamp.replace(
                tzinfo=timezone.utc
            )

        rows = []
        for item in self.ir.certificates:
            expired = None
            if item.valid_until is not None:
                valid_until = item.valid_until
                if valid_until.tzinfo is None:
                    valid_until = valid_until.replace(tzinfo=timezone.utc)
                expired = valid_until < extraction_timestamp

            rows.append(
                (
                    item.name,
                    item.certificate_type,
                    item.source_range,
                    item.source_origin,
                    item.subject,
                    item.issuer,
                    item.serial_number,
                    item.valid_from,
                    item.valid_until,
                    expired,
                    item.public_key_algorithm,
                    item.public_key_size,
                    item.signature_algorithm,
                    item.sha256_fingerprint,
                    item.is_ca,
                    item.is_self_signed,
                    item.has_certificate,
                    item.has_private_key,
                    item.private_key_encrypted,
                    item.has_password,
                    item.source_last_updated,
                    item.migration_status,
                    item.requires_manual_review,
                    item.parse_error,
                    self._format_settings(item.source_attributes),
                    item.description,
                )
            )

        self._table_sheet(
            workbook,
            "Certificates",
            (
                "Name",
                "Certificate Type",
                "Range",
                "Source",
                "Subject",
                "Issuer",
                "Serial Number",
                "Valid From",
                "Valid Until",
                "Expired",
                "Public Key Algorithm",
                "Key Size",
                "Signature Algorithm",
                "SHA-256 Fingerprint",
                "CA Certificate",
                "Self Signed",
                "Has Certificate",
                "Has Private Key",
                "Private Key Encrypted",
                "Has Password",
                "Last Updated",
                "Extraction Status",
                "Manual Review",
                "Parse Error",
                "Additional Settings",
                "Description",
            ),
            rows,
            empty_note="No remote, local, or CA certificates were extracted.",
            subtitle=(
                "Non-secret certificate inventory. Public certificate PEM, "
                "private keys, and passwords are intentionally excluded."
            ),
        )

    def _build_ssh_keys(self, workbook: Any) -> None:
        self._table_sheet(
            workbook,
            "SSH Keys",
            (
                "Name", "Type", "Source", "Has Public Key",
                "Has Private Key", "Has Password", "Extraction Status",
                "Manual Review", "Additional Settings",
            ),
            (
                (
                    item.name,
                    item.key_type,
                    item.source_origin,
                    bool(item.public_key),
                    item.has_private_key,
                    item.has_password,
                    item.migration_status,
                    item.requires_manual_review,
                    self._format_settings(item.source_attributes),
                )
                for item in self.ir.ssh_keys
            ),
            subtitle="Public-key presence is shown; private-key and password contents are never exported.",
        )

    def _build_routes(self, workbook: Any) -> None:
        rows = (
            (
                item.name,
                item.source_route_id,
                item.address_family,
                item.destination,
                item.destination,
                item.source_destination,
                item.source_destination_reference,
                item.source_prefix,
                item.source_preferred_source,
                item.interface,
                item.next_hop,
                item.next_hops,
                item.next_hop_type,
                item.route_type,
                item.rank,
                self._optional_bool_literal(item.scope_local),
                item.monitoring,
                item.administrative_distance,
                item.metric,
                item.priority,
                item.weight,
                item.blackhole,
                item.enabled,
                getattr(item, "installation", None),
                item.sdwan_zone,
                item.sdwan_zones,
                item.dynamic_gateway,
                item.link_monitor_exempt,
                item.bfd,
                item.vrf,
                item.route_tag,
                item.internet_service,
                item.internet_service_custom,
                item.migration_status,
                item.requires_manual_review,
                item.review_reasons,
                item.description,
                self._format_settings(item.source_attributes),
                item.parse_error,
            )
            for item in self.ir.routes
        )
        self._table_sheet(
            workbook,
            "Routes",
            (
                "Name",
                "Source Route ID",
                "Address Family",
                "Destination",
                "Destination Prefix (Normalized)",
                "Source Destination",
                "Destination Object / Group",
                "Source Prefix",
                "Preferred Source",
                "Interface",
                "Next Hop",
                "Next Hops",
                "Next Hop Type",
                "Route Type",
                "Rank",
                "Scope Local",
                "Monitoring",
                "Administrative Distance",
                "Metric",
                "Priority",
                "Weight",
                "Blackhole",
                "Enabled",
                "Installation",
                "SD-WAN Zone",
                "SD-WAN Zones",
                "Dynamic Gateway",
                "Link Monitor Exempt",
                "BFD",
                "VRF",
                "Route Tag",
                "Internet Service",
                "Internet Service Custom",
                "Migration Status",
                "Manual Review",
                "Review Reasons",
                "Description",
                "Additional Settings",
                "Parse Error",
            ),
            rows,
        )

    def _build_policy_routes(self, workbook: Any) -> None:
        rows = [
            (
                item.source_context,
                item.source_id,
                item.address_family,
                item.source_order,
                item.enabled,
                item.source_action,
                item.effective_action,
                item.input_devices,
                item.source_networks,
                item.source_addresses,
                item.destination_networks,
                item.destination_addresses,
                item.protocol,
                item.effective_protocol,
                item.source_port_start,
                item.source_port_end,
                item.destination_port_start,
                item.destination_port_end,
                item.effective_source_port_start,
                item.effective_source_port_end,
                item.effective_destination_port_start,
                item.effective_destination_port_end,
                item.gateway,
                item.output_device,
                item.comments,
                item.migration_status,
                item.requires_manual_review,
                item.review_reasons,
                item.source_explicit_fields,
                self._format_settings(item.source_attributes),
            )
            for item in sorted(self._vendor_extension_value("policy_routes") or [], key=lambda value: value.source_order)
        ]
        self._table_sheet(
            workbook,
            "Policy Routes",
            (
                "Source VDOM", "Rule ID", "Address Family", "Source Order", "Enabled",
                "Configured Action", "Effective Action", "Input Devices", "Source Networks",
                "Source Address References", "Destination Networks", "Destination Address References",
                "Protocol", "Effective Protocol", "Source Port Start", "Source Port End",
                "Destination Port Start", "Destination Port End", "Effective Source Port Start",
                "Effective Source Port End", "Effective Destination Port Start",
                "Effective Destination Port End", "Gateway", "Output Device", "Comments",
                "Migration Status", "Manual Review", "Review Reasons", "Source Explicit Fields",
                "Additional Settings",
            ),
            rows,
            subtitle="Typed FortiGate PBR inventory with configured selectors kept separate from FortiOS effective defaults.",
        )

    def _build_cisco_pbr(self, workbook: Any) -> None:
        self._table_sheet(
            workbook,
            "Cisco PBR",
            (
                "Name", "Source Context", "Source Rule ID", "Sequence", "Action",
                "Match ACL", "Match ACLs", "Resolved Match Criteria", "Match Evidence",
                "Ingress Interface", "Next Hop", "Next Hops", "Next Interface", "Output Interface", "Output Interfaces", "Enabled", "Migration Status", "Manual Review",
                "Review Reasons", "Additional Settings",
            ),
            (
                (
                    item.name, item.source_context, item.source_rule_id, item.source_order,
                    item.action, item.match_acl, item.match_acls, item.resolved_match_criteria,
                    item.match_evidence, item.ingress_interface, item.next_hop, item.next_hops,
                    item.next_interface, item.output_interface, item.output_interfaces,
                    self._optional_bool_literal(item.enabled), item.migration_status,
                    self._optional_bool_literal(item.requires_manual_review), item.review_reasons,
                    self._format_settings(item.source_attributes),
                )
                for item in sorted(self.ir.policy_route_rules, key=lambda value: value.source_order)
            ),
            empty_note="No Cisco ASA/FTD policy-based routing rules were extracted.",
            subtitle="Typed policy-route inventory; these rules are not static routes or firewall policies.",
        )

    def _build_firewall_filters(self, workbook: Any) -> None:
        rows = []
        for item in self.ir.firewall_filters:
            attachments = "; ".join(
                f"{entry.get('interface', '')} ({entry.get('direction', '')})"
                for entry in item.attachments
            )
            for term in item.terms:
                rows.append((
                    item.name,
                    item.source_context,
                    item.family,
                    attachments,
                    term.name,
                    term.source_order,
                    term.matches,
                    term.from_conditions,
                    term.actions,
                    term.migration_status,
                    term.requires_manual_review,
                    term.review_reasons,
                    self._format_settings(term.source_attributes),
                ))
        self._table_sheet(
            workbook,
            "Firewall Filters",
            (
                "Name", "Source Context", "Family", "Interface Attachments", "Term",
                "Source Order", "Matches", "From Conditions", "Actions", "Migration Status", "Manual Review",
                "Review Reasons", "Additional Settings",
            ),
            rows,
            empty_note="No Junos stateless firewall filters were extracted.",
            subtitle="Typed Junos stateless firewall-filter inventory; separate from SRX security policies and policy routes.",
        )

    def _build_vpn_phase2(self, workbook: Any) -> None:
        rows = [
            (
                item.name,
                item.phase1_name,
                item.proposals,
                item.source_address_type,
                item.source_names,
                item.source_names6,
                item.destination_address_type,
                item.destination_names,
                item.destination_names6,
                item.source_subnet,
                item.source_subnet6,
                item.destination_subnet,
                item.destination_subnet6,
                item.source_range_start,
                item.source_range_end,
                item.destination_range_start,
                item.destination_range_end,
                item.source_range_start6,
                item.source_range_end6,
                item.destination_range_start6,
                item.destination_range_end6,
                self._optional_bool_literal(item.pfs),
                item.dh_groups,
                item.key_lifetime,
                item.keylife_type,
                item.keylife_seconds,
                item.keylife_kilobytes,
                item.replay,
                item.protocol,
                item.source_port,
                item.destination_port,
                self._optional_bool_literal(item.auto_negotiate),
                self._optional_bool_literal(item.keepalive),
                item.migration_status,
                self._optional_bool_literal(item.requires_manual_review),
                item.review_reasons,
                self._format_settings(item.source_attributes),
                item.description,
            )
            for item in self.ir.vpn_phase2
        ]
        self._table_sheet(
            workbook,
            "VPN Phase 2",
            (
                "Name",
                "Phase 1",
                "Proposal",
                "Source Address Type",
                "Source Selector",
                "Source IPv6 Selector",
                "Destination Address Type",
                "Destination Selector",
                "Destination IPv6 Selector",
                "Source Subnet",
                "Source Subnet IPv6",
                "Destination Subnet",
                "Destination Subnet IPv6",
                "Source Range Start",
                "Source Range End",
                "Destination Range Start",
                "Destination Range End",
                "Source IPv6 Range Start",
                "Source IPv6 Range End",
                "Destination IPv6 Range Start",
                "Destination IPv6 Range End",
                "PFS",
                "DH Groups",
                "Key Lifetime",
                "Keylife Type",
                "Keylife Seconds",
                "Keylife Kilobytes",
                "Replay",
                "Protocol",
                "Source Port",
                "Destination Port",
                "Auto Negotiate",
                "Keepalive",
                "Extraction Status",
                "Manual Review",
                "Review Reasons",
                "Additional Settings",
                "Description",
            ),
            rows,
        )

    @staticmethod
    def _routing_protocol_label(source_path: str) -> str:
        return {
            "router rip": "RIP",
            "router ripng": "RIPng",
            "router ospf": "OSPF",
            "router ospf6": "OSPFv3",
            "router bgp": "BGP",
            "router isis": "ISIS",
            "router multicast": "Multicast Routing",
        }.get(source_path, source_path)

    def _structured_routing_items(self) -> list[Any]:
        if self.extraction is None:
            return []
        return [
            item for item in self.extraction.inventory_items
            if "structured-routing-protocol" in item.notes
        ]

    def _build_routing_protocols(self, workbook: Any) -> None:
        items = self._structured_routing_items()
        self._table_sheet(
            workbook,
            "Routing Protocols",
            (
                "Protocol", "Name or Instance", "Source Block Present",
                "Configured", "Extraction Status", "Manual Review",
            ),
            (
                (
                    self._routing_protocol_label(item.source_path),
                    item.name,
                    "Yes",
                    (
                        "Yes"
                        if self._flatten_source_profile_settings(item)
                        else "No"
                    ),
                    item.status,
                    self._optional_bool_literal(item.requires_manual_review),
                )
                for item in items
            ),
        )
        self._table_sheet(
            workbook,
            "Routing Protocol Settings",
            ("Protocol", "Object / Instance", "Subsection", "Entry", "Operation", "Setting", "Value"),
            (
                (self._routing_protocol_label(row[0]), *row[1:])
                for item in items
                for row in self._flatten_source_profile_settings(item)
            ),
        )

    def _build_security_profiles(self, workbook: Any) -> None:
        rows = (
            (
                item.name, item.antivirus, item.vulnerability, item.anti_spyware,
                item.url_filtering, item.file_blocking, item.wildfire,
                item.ssl_decryption, item.description, item.support_level,
                self._format_settings(item.source_attributes),
            )
            for item in self.ir.security_profile_groups
        )
        self._table_sheet(
            workbook,
            "Security Profiles",
            (
                "Name", "Antivirus", "Vulnerability", "Anti-Spyware", "URL Filtering",
                "File Blocking", "WildFire", "SSL Decryption", "Description",
                "Support Level", "Additional Settings",
            ),
            rows,
        )

    def _build_security_profile_definitions(self, workbook: Any) -> None:
        rows = [(
            item.name, item.family, item.source_family, item.source_context, item.description,
            len(item.rules), ", ".join(item.allow_categories), ", ".join(item.alert_categories),
            ", ".join(item.block_categories), ", ".join(item.continue_categories),
            ", ".join(item.override_categories),
            item.credential_enforcement.mode if item.credential_enforcement else None,
            item.credential_enforcement.log_severity if item.credential_enforcement else None,
            ", ".join(item.credential_enforcement.block_categories) if item.credential_enforcement else None,
            self._optional_bool_literal(item.log_http_hdr_xff),
            self._optional_bool_literal(item.log_http_hdr_user_agent), item.support_level,
            item.migration_status, item.requires_manual_review, ", ".join(item.review_reasons),
            self._format_settings(item.source_attributes),
        ) for item in self.ir.security_profile_definitions]
        self._table_sheet(workbook, "Security Profile Definitions", (
            "Name", "Family", "Source Family", "Source Context", "Description", "Rule Count",
            "Allow Categories", "Alert Categories", "Block Categories", "Continue Categories",
            "Override Categories", "Credential Mode", "Credential Log Severity",
            "Credential Block Categories", "Log HTTP XFF", "Log HTTP User Agent", "Support Level",
            "Extraction Status", "Manual Review", "Review Reasons", "Additional Settings",
        ), rows, empty_note="No PAN-OS security profile definitions were extracted.")

    def _build_security_profile_rules(self, workbook: Any) -> None:
        rows = [(
            profile.name, profile.family, profile.source_context, rule.name, rule.action,
            ", ".join(rule.applications), ", ".join(rule.file_types), rule.direction,
            ", ".join(rule.severities), ", ".join(rule.vendor_ids), ", ".join(rule.cves),
            rule.threat_name, rule.host, rule.category, rule.packet_capture,
            self._format_settings(rule.source_attributes),
        ) for profile in self.ir.security_profile_definitions for rule in profile.rules]
        self._table_sheet(workbook, "Security Profile Rules", (
            "Profile Name", "Profile Family", "Source Context", "Rule Name", "Action", "Applications",
            "File Types", "Direction", "Severities", "Vendor IDs", "CVEs", "Threat Name", "Host",
            "Category", "Packet Capture", "Additional Settings",
        ), rows, empty_note="No PAN-OS security profile rules were extracted.")

    def _build_custom_url_categories(self, workbook: Any) -> None:
        rows = [(
            item.name, item.source_context, item.category_type, ", ".join(item.entries), item.description,
            item.support_level, item.migration_status, item.requires_manual_review,
            ", ".join(item.review_reasons), self._format_settings(item.source_attributes),
        ) for item in self.ir.custom_url_categories]
        self._table_sheet(workbook, "Custom URL Categories", (
            "Name", "Source Context", "Type", "Entries", "Description", "Support Level",
            "Extraction Status", "Manual Review", "Review Reasons", "Additional Settings",
        ), rows, empty_note="No PAN-OS custom URL categories were extracted.")

    def _build_vip_groups(self, workbook: Any) -> None:
        self._table_sheet(
            workbook,
            "VIP Groups",
            (
                "Name", "Address Family", "Source UUID", "Interface", "Members", "Source Color",
                "Extraction Status", "Manual Review", "Review Reason",
                "Additional Settings", "Description", "Source VDOM", "Unresolved Members",
            ),
            (
                (
                    item.name, item.address_family, item.source_uuid, item.interface, item.members,
                    item.source_color, item.migration_status,
                    self._optional_bool_literal(item.requires_manual_review),
                    item.audit_note, self._format_settings(item.source_attributes), item.description,
                    item.source_context, item.unresolved_members,
                )
                for item in self.ir.virtual_ip_groups
            ),
        )

    def _build_sdwan(self, workbook: Any) -> None:
        sdwans = self._vendor_extension_value("sdwans") or []
        self._table_sheet(
            workbook,
            "SD-WAN",
            (
                "Status", "Load Balance Mode", "Extraction Status", "Manual Review",
                "Additional Settings", "VDOM",
            ),
            (
                (
                    sdwan.status, sdwan.load_balance_mode, sdwan.migration_status,
                    self._optional_bool_literal(sdwan.requires_manual_review),
                    self._format_settings(sdwan.source_attributes),
                    sdwan.source_context,
                )
                for sdwan in sdwans
            ),
        )

        self._table_sheet(
            workbook,
            "SD-WAN Zones",
            ("Zone Name", "Additional Settings", "VDOM"),
            (
                (zone.name, self._format_settings(zone.source_attributes), zone.source_context)
                for sdwan in sdwans
                for zone in sdwan.zones
            ),
        )
        self._table_sheet(
            workbook,
            "SD-WAN Members",
            (
                "ID", "Interface", "Zone", "Gateway", "Source", "IPv6 Gateway",
                "IPv6 Source", "Cost", "Weight", "Priority", "IPv6 Priority",
                "Spillover Threshold", "Ingress Spillover Threshold", "Volume Ratio",
                "Status", "Description", "Additional Settings", "VDOM",
            ),
            (
                (
                    item.source_id, item.interface, item.zone, item.gateway,
                    item.source, item.gateway6, item.source6, item.cost, item.weight,
                    item.priority, item.priority6, item.spillover_threshold,
                    item.ingress_spillover_threshold, item.volume_ratio, item.status,
                    item.description, self._format_settings(item.source_attributes),
                    item.source_context,
                )
                for sdwan in sdwans
                for item in sdwan.members
            ),
        )
        self._table_sheet(
            workbook,
            "SD-WAN Health Checks",
            (
                "Name", "Server", "Members", "Protocol", "Port", "Interval",
                "Probe Timeout", "Fail Time", "Recovery Time", "Update Static Route",
                "VRF", "Source Address", "SLA Count", "Additional Settings", "VDOM",
            ),
            (
                (
                    item.name, item.server, item.member_ids, item.protocol, item.port,
                    item.interval, item.probe_timeout, item.failtime, item.recoverytime,
                    item.update_static_route, item.vrf, item.source, len(item.sla),
                    self._format_settings(item.source_attributes), item.source_context,
                )
                for sdwan in sdwans
                for item in sdwan.health_checks
            ),
        )
        self._table_sheet(
            workbook,
            "SD-WAN SLAs",
            ("Health Check", "SLA ID", "Additional Settings", "VDOM"),
            (
                (check.name, sla.source_id, self._format_settings(sla.source_attributes), sla.source_context)
                for sdwan in sdwans
                for check in sdwan.health_checks
                for sla in check.sla
            ),
        )
        self._table_sheet(
            workbook,
            "SD-WAN Rules",
            (
                "ID", "Name", "Mode", "Status", "Source", "Destination",
                "Health Checks", "Priority Members", "Priority Zones",
                "Internet Service", "Internet Service Names",
                "Internet Service App Control", "SLA Compare Method", "Tie Break",
                "Use Shortcut SLA", "Additional Settings", "VDOM",
            ),
            (
                (
                    item.source_id, item.name, item.mode, item.status, item.source_addresses,
                    item.destination_addresses, item.health_checks, item.priority_member_ids,
                    item.priority_zones,
                    item.internet_service, item.internet_service_names,
                    item.internet_service_app_ctrl, item.sla_compare_method, item.tie_break,
                    item.use_shortcut_sla,
                    self._format_settings(item.source_attributes), item.source_context,
                )
                for sdwan in sdwans
                for item in sdwan.rules
            ),
        )
        self._build_sdwan_source_details(workbook, sdwans)

    def _build_sdwan_source_details(self, workbook: Any, sdwans: list[Any]) -> None:
        self._table_sheet(
            workbook,
            "SD-WAN Duplication",
            (
                "ID", "Service ID", "Source Addresses", "Destination Addresses",
                "IPv6 Source Addresses", "IPv6 Destination Addresses",
                "Source Interfaces", "Destination Interfaces", "Services",
                "Packet Duplication", "SLA Match Service", "Packet De-duplication",
                "Extraction Status", "Manual Review", "Additional Settings", "VDOM",
            ),
            (
                (
                    item.source_id, item.service_id, item.source_addresses,
                    item.destination_addresses, item.source_addresses6,
                    item.destination_addresses6, item.source_interfaces,
                    item.destination_interfaces, item.services, item.packet_duplication,
                    item.sla_match_service, item.packet_de_duplication,
                    item.migration_status,
                    self._optional_bool_literal(item.requires_manual_review),
                    self._format_settings(item.source_attributes), item.source_context,
                )
                for sdwan in sdwans
                for item in sdwan.duplication_rules
            ),
        )
        self._table_sheet(
            workbook,
            "SD-WAN Neighbors",
            ("Name", "Extraction Status", "Manual Review", "Additional Settings", "VDOM"),
            (
                (
                    item.name, item.migration_status,
                    self._optional_bool_literal(item.requires_manual_review),
                    self._format_settings(item.source_attributes), item.source_context,
                )
                for sdwan in sdwans
                for item in sdwan.neighbors
            ),
        )
        self._table_sheet(
            workbook,
            "SD-WAN Rule SLAs",
            ("Rule ID", "Rule Name", "SLA", "SLA ID", "Additional Settings", "VDOM"),
            (
                (
                    rule.source_id, rule.name, sla.name, sla.source_id,
                    self._format_settings(sla.source_attributes), sla.source_context,
                )
                for sdwan in sdwans
                for rule in sdwan.rules
                for sla in rule.sla
            ),
        )

    def _structured_security_items(self) -> list[Any]:
        if self.extraction is None:
            return []
        return [
            item for item in self.extraction.inventory_items
            if "structured-security-profile" in item.notes
        ]

    def _profile_policy_references(self, name: Any) -> list[str]:
        if not name:
            return []
        references = []
        for policy in self.ir.policies:
            known = {
                policy.antivirus,
                policy.ips_sensor,
                policy.webfilter,
                policy.application_list,
                policy.ssl_ssh_profile,
            }
            if name in known:
                references.append(policy.source_rule_id or policy.name)
        return references

    def _flatten_source_profile_settings(self, item: Any) -> list[tuple[Any, ...]]:
        rows = []

        def walk(node: Any, subsections: list[str], entry: Any) -> None:
            node_type = next(
                (note.split(":", 1)[1] for note in node.notes if note.startswith("source-node:")),
                "edit",
            )
            current_subsections = subsections
            current_entry = entry
            if node is not item:
                if node_type == "config":
                    current_subsections = [*subsections, node.name]
                elif node_type == "edit":
                    current_entry = node.name
            for command in node.commands:
                rows.append((
                    item.source_path,
                    item.name,
                    " / ".join(current_subsections),
                    current_entry,
                    command.operation,
                    command.key,
                    command.values,
                ))
            for child in node.children:
                walk(child, current_subsections, current_entry)

        walk(item, [], item.name)
        return rows

    def _build_source_security_profiles(self, workbook: Any) -> None:
        items = self._structured_security_items()
        self._table_sheet(
            workbook,
            "Source Security Profiles",
            ("Profile Type", "Name", "Extraction Status", "Manual Review", "Referenced By Policy"),
            (
                (
                    item.source_path, item.name, item.status,
                    self._optional_bool_literal(item.requires_manual_review),
                    self._profile_policy_references(item.name),
                )
                for item in items
            ),
        )
        self._table_sheet(
            workbook,
            "Source Security Profile Setting",
            ("Profile Type", "Profile Name", "Subsection", "Entry", "Operation", "Setting", "Value"),
            (row for item in items for row in self._flatten_source_profile_settings(item)),
        )

    def _build_identity_inventory(self, workbook: Any) -> None:
        self._table_sheet(
            workbook, "LDAP Servers",
            ("Name", "Server", "Secondary Server", "Tertiary Server", "CNID", "DN", "Type", "Username", "Password Configured", "Port", "Secure", "CA Certificate", "CA Certificate Resolved", "Client Certificate", "Client Certificate Resolved", "Unresolved Certificate References", "Server Identity Check", "Source IP", "Source Port", "Interface Selection", "Interface", "Group Filter", "Group Search Base", "Group Member Check", "Group Object Filter", "Member Attribute", "Password Attribute", "Obtain User Info", "Password Expiry Warning", "Password Renewal", "Account Key Certificate Field", "Account Key Filter", "Account Key Processing", "Antiphish", "Search Type", "SSL Minimum Protocol", "Extraction Status", "Manual Review", "Additional Settings"),
            (
                (
                    item.name, item.server, item.secondary_server, item.tertiary_server,
                    item.cnid, item.dn, item.source_type, item.username,
                    item.has_password, item.port, item.secure, item.ca_cert,
                    self._optional_bool_literal(item.ca_certificate_resolved),
                    item.client_cert,
                    self._optional_bool_literal(item.client_certificate_resolved),
                    item.unresolved_certificate_references, item.server_identity_check, item.source_ip,
                    item.source_port, item.interface_select_method, item.interface, item.group_filter,
                    item.group_search_base, item.group_member_check, item.group_object_filter,
                    item.member_attr, item.password_attr, item.obtain_user_info,
                    item.password_expiry_warning, item.password_renewal,
                    item.account_key_cert_field, item.account_key_filter,
                    item.account_key_processing, item.antiphish, item.search_type,
                    item.ssl_min_proto_version,
                    item.migration_status,
                    item.requires_manual_review, self._format_settings(item.source_attributes),
                ) for item in self._vendor_extension_value("user_ldap_servers") or []
            ),
        )
        self._table_sheet(
            workbook, "RADIUS Servers",
            (
                "Name", "Source Context", "Primary Server", "Secondary Server",
                "Tertiary Server", "Authentication Type", "Port", "NAS IP",
                "Source IP", "Accounting Interim Interval", "Secret Configured", "Extraction Status",
                "Manual Review", "Additional Settings",
            ),
            (
                (
                    item.name, item.source_context, item.server,
                    item.secondary_server, item.tertiary_server, item.auth_type,
                    item.port, item.nas_ip, item.source_ip, item.acct_interim_interval,
                    item.has_secret,
                    item.migration_status, item.requires_manual_review,
                    self._format_settings(item.source_attributes),
                )
                for item in self._vendor_extension_value("user_radius_servers") or []
            ),
        )
        self._table_sheet(
            workbook, "RADIUS Accounting Servers",
            (
                "RADIUS Server", "Accounting ID", "Status", "Server", "Port",
                "Source IP", "Interface Selection", "Interface", "Secret Configured",
                "Extraction Status", "Manual Review", "Additional Settings",
            ),
            (
                (
                    item.name, accounting.id, accounting.status, accounting.server,
                    accounting.port, accounting.source_ip,
                    accounting.interface_select_method, accounting.interface,
                    accounting.has_secret, accounting.migration_status,
                    accounting.requires_manual_review,
                    self._format_settings(accounting.source_attributes),
                )
                for item in self._vendor_extension_value("user_radius_servers") or []
                for accounting in item.accounting_servers
            ),
        )
        self._table_sheet(
            workbook, "TACACS+ Servers",
            (
                "Name", "Source Context", "Primary Server", "Secondary Server",
                "Tertiary Server", "Port", "Authentication Type", "Authorization",
                "Source IP", "Interface Selection", "Interface", "Status TTL", "Secret Configured",
                "Extraction Status", "Manual Review", "Additional Settings",
            ),
            (
                (
                    item.name, item.source_context, item.server,
                    item.secondary_server, item.tertiary_server, item.port,
                    item.authentication_type, item.authorization, item.source_ip,
                    item.interface_select_method, item.interface, item.status_ttl, item.has_secret,
                    item.migration_status, item.requires_manual_review,
                    self._format_settings(item.source_attributes),
                )
                for item in self._vendor_extension_value("user_tacacs_servers") or []
            ),
        )
        self._table_sheet(
            workbook, "SAML Servers",
            ("Name", "Entity ID", "SSO URL", "SLO URL", "IdP Entity ID", "IdP SSO URL", "IdP SLO URL", "IdP Certificate", "SP Certificate", "IdP Certificate Resolved", "SP Certificate Resolved", "Unresolved Certificate References", "User Name", "Group Name", "Digest Method", "Clock Tolerance", "ADFS Claim", "Limit Relaystate", "Reauth", "User Claim Type", "Group Claim Type", "Extraction Status", "Manual Review", "Additional Settings"),
            (
                (
                    item.name, item.entity_id, item.single_sign_on_url, item.single_logout_url,
                    item.idp_entity_id, item.idp_single_sign_on_url,
                    item.idp_single_logout_url, item.idp_cert, item.cert,
                    self._optional_bool_literal(item.idp_certificate_resolved),
                    self._optional_bool_literal(item.cert_certificate_resolved),
                    item.unresolved_certificate_references, item.user_name,
                    item.group_name, item.digest_method, item.clock_tolerance,
                    item.adfs_claim, item.limit_relaystate, item.reauth,
                    item.user_claim_type, item.group_claim_type,
                    item.migration_status,
                    item.requires_manual_review, self._format_settings(item.source_attributes),
                ) for item in self._vendor_extension_value("user_saml_servers") or []
            ),
        )
        self._table_sheet(
            workbook, "FSSO Servers",
            ("Name", "Server", "Server 2", "Server 3", "Server 4", "Server 5", "Port", "Port 2", "Port 3", "Port 4", "Port 5", "Interface Selection", "Interface", "LDAP Poll", "LDAP Poll Filter", "LDAP Poll Interval", "LDAP Server", "Logon Timeout", "Source IP", "Source IPv6", "SSL", "SSL Host/IP Check", "SSL Trusted Certificate", "Type", "User Info Server", "Password Configured", "Extraction Status", "Manual Review", "Additional Settings"),
            (
                (
                    item.name, item.server, item.server2, item.server3, item.server4,
                    item.server5, item.port, item.port2, item.port3, item.port4,
                    item.port5, item.interface_select_method, item.interface,
                    item.ldap_poll, item.ldap_poll_filter, item.ldap_poll_interval,
                    item.ldap_server, item.logon_timeout, item.source_ip,
                    item.source_ip6, item.ssl, item.ssl_server_host_ip_check,
                    item.ssl_trusted_cert, item.source_type, item.user_info_server,
                    item.has_password,
                    item.migration_status, item.requires_manual_review,
                    self._format_settings(item.source_attributes),
                ) for item in self._vendor_extension_value("fsso_providers") or []
            ),
        )
        self._table_sheet(
            workbook, "FSSO AD Groups",
            ("Name", "FSSO Server", "Server Resolved", "Extraction Status", "Manual Review", "Additional Settings"),
            (
                (
                    item.name, item.provider_name, item.provider_resolved,
                    item.migration_status, item.requires_manual_review,
                    self._format_settings(item.source_attributes),
                ) for item in self._vendor_extension_value("fsso_ad_groups") or []
            ),
        )
        self._table_sheet(
            workbook, "Local Users",
            ("Name", "ID", "Status", "Type", "Password Configured", "Password Time", "Two Factor", "Two Factor Authentication", "Two Factor Notification", "FortiToken", "Email", "SMS Server", "SMS Custom Server", "SMS Phone", "LDAP Server", "RADIUS Server", "TACACS+ Server", "Auth Concurrent Override", "Auth Concurrent Value", "Auth Timeout", "Password Policy", "Workstation", "Username Sensitivity", "PPK Identity", "PPK Secret Configured", "Extraction Status", "Manual Review", "Additional Settings"),
            (
                (
                    item.name, item.id, item.status, item.source_type, item.has_password,
                    item.source_passwd_time,
                    item.two_factor, item.two_factor_authentication,
                    item.two_factor_notification, item.fortitoken, item.email_to,
                    item.sms_server, item.sms_custom_server, item.sms_phone,
                    item.ldap_server, item.radius_server, item.tacacs_server,
                    item.auth_concurrent_override, item.auth_concurrent_value,
                    item.authtimeout, item.passwd_policy, item.workstation,
                    item.username_sensitivity, item.ppk_identity, item.has_ppk_secret,
                    item.migration_status, item.requires_manual_review,
                    self._format_settings(item.source_attributes),
                ) for item in self._vendor_extension_value("local_users") or []
            ),
        )
        self._table_sheet(
            workbook, "User Groups",
            ("Name", "ID", "Type", "Members", "Resolved Members", "Unresolved Members", "Match Count", "Auth Concurrent Override", "Auth Concurrent Value", "Auth Timeout", "Company", "Email", "Expire", "Expire Type", "HTTP Digest Realm", "Max Accounts", "Mobile Phone", "Multiple Guest Add", "Password Mode", "SMS Server", "SMS Custom Server", "Sponsor", "SSO Attribute", "User ID", "User Name", "Unresolved Match Servers", "Extraction Status", "Manual Review", "Additional Settings"),
            (
                (
                    item.name, item.id, item.group_type, item.members, item.resolved_members,
                    item.unresolved_members, len(item.matches), item.auth_concurrent_override,
                    item.auth_concurrent_value, item.authtimeout, item.company, item.email,
                    item.expire, item.expire_type, item.http_digest_realm, item.max_accounts,
                    item.mobile_phone, item.multiple_guest_add, item.password, item.sms_server,
                    item.sms_custom_server, item.sponsor, item.sso_attribute_value,
                    item.user_id, item.user_name,
                    item.unresolved_match_servers,
                    item.migration_status, item.requires_manual_review,
                    self._format_settings(item.source_attributes),
                ) for item in self._vendor_extension_value("user_groups") or []
            ),
        )
        self._table_sheet(
            workbook, "User Group Matches",
            ("User Group", "ID", "Server Name", "Server Resolved", "Group Name"),
            (
                (
                    group.name, match.source_id, match.server_name,
                    self._optional_bool_literal(
                        None if match.server_name is None
                        else match.server_name not in group.unresolved_match_servers
                    ),
                    match.group_name,
                )
                for group in self._vendor_extension_value("user_groups") or [] for match in group.matches
            ),
        )

    def _build_ssl_vpn(self, workbook: Any) -> None:
        settings = self._vendor_extension_value("ssl_vpn_settings")
        self._table_sheet(
            workbook, "SSL VPN Settings",
            ("Status", "Minimum Protocol", "Maximum Protocol", "Algorithm", "Banned Ciphers", "Client Signature Algorithms", "Require Client Certificate", "DTLS Tunnel", "Login Attempt Limit", "Login Block Time", "Authentication Timeout", "Idle Timeout", "Port", "DNS Server 1", "DNS Server 2", "WINS Server 1", "WINS Server 2", "Server Certificate", "Server Certificate Configured", "Source Interfaces", "Source Addresses", "Tunnel IP Pools", "Default Portal", "Extraction Status", "Manual Review", "Typed Source Fields", "Additional Settings"),
            [] if settings is None else [(
                settings.status, settings.ssl_min_proto_ver, settings.ssl_max_proto_ver,
                settings.algorithm, settings.banned_cipher,
                settings.client_signature_algorithms,
                settings.require_client_certificate, settings.dtls_tunnel,
                settings.login_attempt_limit, settings.login_block_time,
                settings.auth_timeout, settings.idle_timeout, settings.port,
                settings.dns_server1, settings.dns_server2,
                settings.wins_server1, settings.wins_server2,
                settings.server_certificate,
                "TRUE" if settings.server_certificate_configured else "FALSE",
                settings.source_interfaces,
                settings.source_addresses, settings.tunnel_ip_pools,
                settings.default_portal, settings.migration_status,
                settings.requires_manual_review, self._format_settings(settings.source_fields), self._format_settings(settings.source_attributes),
            )],
        )
        self._table_sheet(
            workbook, "SSL VPN Portals",
            ("Name", "Tunnel Mode", "IPv6 Tunnel Mode", "IP Pools", "IPv6 Pools", "Split Tunneling", "Limit User Logins", "FortiClient Download", "Host Check", "Host Check Policies", "Host Check Interval", "Unresolved Host Check Policies", "Allow User Access", "Auto Connect", "Exclusive Routing", "IP Mode", "Service Restriction", "Split Tunneling Routing Addresses", "Split Tunneling Routing Negate", "Extraction Status", "Manual Review", "Typed Source Fields", "Additional Settings"),
            (
                (
                    item.name, item.tunnel_mode, item.ipv6_tunnel_mode, item.ip_pools,
                    item.ipv6_pools, item.split_tunneling, item.limit_user_logins,
                    item.forticlient_download, item.host_check,
                    item.host_check_policies, item.host_check_interval,
                    item.unresolved_host_check_policies, item.allow_user_access,
                    item.auto_connect, item.exclusive_routing, item.ip_mode,
                    item.service_restriction,
                    item.split_tunneling_routing_addresses,
                    item.split_tunneling_routing_negate, item.migration_status,
                    item.requires_manual_review, self._format_settings(item.source_fields), self._format_settings(item.source_attributes),
                ) for item in self._vendor_extension_value("ssl_vpn_portals") or []
            ),
        )
        self._table_sheet(
            workbook, "SSL VPN Authentication Rules",
            ("ID", "Auth", "Cipher", "Client Certificate", "Realm", "Source Interfaces", "Source Addresses", "Source Address Negate", "IPv6 Source Addresses", "IPv6 Source Address Negate", "Users", "User Peer", "Groups", "Unresolved Groups", "Portal", "Extraction Status", "Manual Review", "Additional Settings"),
            [] if settings is None else (
                (
                    item.source_id, item.auth, item.cipher, item.client_cert,
                    item.realm, item.source_interfaces, item.source_addresses,
                    item.source_address_negate, item.source_addresses6,
                    item.source_address6_negate, item.users, item.user_peer,
                    item.groups, item.unresolved_groups, item.portal, item.migration_status,
                    item.requires_manual_review,
                    self._format_settings(item.source_attributes),
                )
                for item in settings.authentication_rules
            ),
        )
        self._table_sheet(
            workbook, "User Group Guests",
            ("User Group", "ID", "Name", "User ID", "Company", "Email", "Expiration", "Mobile Phone", "Sponsor", "Password Configured", "Additional Settings"),
            (
                (
                    group.name, guest.id, guest.name, guest.user_id, guest.company,
                    guest.email, guest.expiration, guest.mobile_phone, guest.sponsor,
                    guest.has_password, self._format_settings(guest.source_attributes),
                )
                for group in self._vendor_extension_value("user_groups") or [] for guest in group.guests
            ),
        )
        self._table_sheet(
            workbook, "FSSO Polling",
            ("Name", "Source Context", "Status", "Server", "Default Domain", "Port", "User", "Password Configured", "LDAP Server", "Logon History", "Polling Frequency", "SMBv1", "SMB NTLMv1 Auth", "AD Groups", "Extraction Status", "Manual Review", "Additional Settings"),
            (
                (
                    item.name, item.source_context, item.status, item.server,
                    item.default_domain, item.port, item.user, item.has_password,
                    item.ldap_server, item.logon_history, item.polling_frequency,
                    item.smbv1, item.smb_ntlmv1_auth,
                    ", ".join(group.name for group in item.ad_groups),
                    item.migration_status, item.requires_manual_review,
                    self._format_settings(item.source_attributes),
                ) for item in self._vendor_extension_value("fsso_polling") or []
            ),
        )
        portals = self._vendor_extension_value("ssl_vpn_portals") or []
        self._table_sheet(workbook, "SSL VPN Portal Split DNS",
            ("Portal", "ID", "Domains", "DNS Server 1", "DNS Server 2", "IPv6 DNS Server 1", "IPv6 DNS Server 2", "Extraction Status", "Manual Review", "Additional Settings"),
            ((portal.name, item.id, item.domains, item.dns_server1, item.dns_server2, item.ipv6_dns_server1, item.ipv6_dns_server2, item.migration_status, item.requires_manual_review, self._format_settings(item.source_attributes)) for portal in portals for item in portal.split_dns))
        self._table_sheet(workbook, "SSL VPN Portal MAC Rules",
            ("Portal", "ID", "Extraction Status", "Manual Review", "Additional Settings"),
            ((portal.name, item.id, item.migration_status, item.requires_manual_review, self._format_settings(item.source_attributes)) for portal in portals for item in portal.mac_address_check_rules))
        self._table_sheet(workbook, "SSL VPN Portal OS Checks",
            ("Portal", "ID", "Extraction Status", "Manual Review", "Additional Settings"),
            ((portal.name, item.id, item.migration_status, item.requires_manual_review, self._format_settings(item.source_attributes)) for portal in portals for item in portal.os_check_list))
        self._table_sheet(workbook, "SSL VPN Bookmark Groups",
            ("Portal", "Bookmark Group", "Extraction Status", "Manual Review", "Additional Settings"),
            ((portal.name, group.name, group.migration_status, group.requires_manual_review, self._format_settings(group.source_attributes)) for portal in portals for group in portal.bookmark_groups))
        self._table_sheet(workbook, "SSL VPN Bookmarks",
            ("Portal", "Bookmark Group", "Bookmark", "Has Logon Password", "Has SSO Password", "Extraction Status", "Manual Review", "Additional Settings"),
            ((portal.name, group.name, bookmark.name, bookmark.has_logon_password, bookmark.has_sso_password, bookmark.migration_status, bookmark.requires_manual_review, self._format_settings(bookmark.source_attributes)) for portal in portals for group in portal.bookmark_groups for bookmark in group.bookmarks))
        self._table_sheet(workbook, "SSL VPN Bookmark Form Data",
            ("Portal", "Bookmark Group", "Bookmark", "Form Data", "Value Configured", "Extraction Status", "Manual Review", "Additional Settings"),
            ((portal.name, group.name, bookmark.name, item.name, item.value_configured, item.migration_status, item.requires_manual_review, self._format_settings(item.source_attributes)) for portal in portals for group in portal.bookmark_groups for bookmark in group.bookmarks for item in bookmark.form_data))
        self._table_sheet(workbook, "SSL VPN Landing Pages",
            ("Portal", "Landing Page", "Extraction Status", "Manual Review", "Additional Settings"),
            ((portal.name, page.name, page.migration_status, page.requires_manual_review, self._format_settings(page.source_attributes)) for portal in portals for page in portal.landing_pages))
        self._table_sheet(workbook, "SSL VPN Landing Form Data",
            ("Portal", "Landing Page", "Form Data", "Value Configured", "Extraction Status", "Manual Review", "Additional Settings"),
            ((portal.name, page.name, item.name, item.value_configured, item.migration_status, item.requires_manual_review, self._format_settings(item.source_attributes)) for portal in portals for page in portal.landing_pages for item in page.form_data))

    def _build_user_identity_settings(self, workbook: Any) -> None:
        settings = self._vendor_extension_value("user_authentication_settings")
        self._table_sheet(
            workbook,
            "User Authentication Settings",
            (
                "Auth Certificate", "Auth Certificate Resolved",
                "Auth CA Certificate", "Auth CA Certificate Resolved",
                "Auth Timeout", "Auth Lockout Threshold", "Auth Lockout Duration",
                "Minimum TLS Version", "Management Authentication Profile",
                "Management Authentication Profile Resolved", "Unresolved Management Authentication Profile",
                "Extraction Status", "Manual Review",
                "Additional Settings",
            ),
            [] if settings is None else [(
                settings.auth_certificate,
                self._optional_bool_literal(settings.auth_certificate_resolved),
                settings.auth_ca_certificate,
                self._optional_bool_literal(settings.auth_ca_certificate_resolved),
                settings.auth_timeout, settings.auth_lockout_threshold,
                settings.auth_lockout_duration, settings.ssl_min_proto_version,
                settings.management_authentication_profile,
                self._optional_bool_literal(settings.management_authentication_profile_resolved),
                settings.unresolved_management_authentication_profile,
                settings.migration_status,
                self._optional_bool_literal(settings.requires_manual_review),
                self._format_settings(settings.source_attributes),
            )],
        )
        quarantine = self._vendor_extension_value("user_quarantine_settings")
        self._table_sheet(
            workbook,
            "User Quarantine",
            (
                "Firewall Groups", "Resolved Firewall Groups",
                "Unresolved Firewall Groups", "Extraction Status",
                "Manual Review", "Additional Settings",
            ),
            [] if quarantine is None else [(
                quarantine.firewall_groups,
                quarantine.resolved_firewall_groups,
                quarantine.unresolved_firewall_groups,
                quarantine.migration_status,
                self._optional_bool_literal(quarantine.requires_manual_review),
                self._format_settings(quarantine.source_attributes),
            )],
        )

    def _build_security_identity_dependencies(self, workbook: Any) -> None:
        rows: list[tuple[Any, ...]] = []

        def add(
            consumer_type: str, consumer_name: str, dependency_type: str,
            reference: str, resolved: bool, impact: str, notes: str,
        ) -> None:
            rows.append((
                consumer_type, consumer_name, dependency_type, reference,
                self._optional_bool_literal(resolved),
                "RESOLVED" if resolved else "UNRESOLVED",
                impact, notes,
            ))

        for group in self._vendor_extension_value("user_groups") or []:
            for dependency in group.member_dependencies:
                add(
                    "User Group", group.name, dependency.dependency_type,
                    dependency.reference, dependency.resolved, "EXTRACT_ONLY",
                    "Source identity dependency preserved; target identity mapping requires review.",
                )
            for match in group.matches:
                if match.server_name:
                    add(
                        "User Group Match", group.name, "authentication-provider",
                        match.server_name,
                        match.server_name not in group.unresolved_match_servers,
                        "EXTRACT_ONLY",
                        "External directory group name is preserved but is not a local FortiGate reference.",
                    )
        for ad_group in self._vendor_extension_value("fsso_ad_groups") or []:
            if ad_group.provider_name:
                add(
                    "FSSO AD Group", ad_group.name, "fsso-provider",
                    ad_group.provider_name, ad_group.provider_resolved,
                    "EXTRACT_ONLY", "FSSO provider reference preserved.",
                )
        for saml in self._vendor_extension_value("user_saml_servers") or []:
            if saml.idp_cert and saml.idp_certificate_resolved is not None:
                add(
                    "SAML Server", saml.name, "certificate", saml.idp_cert,
                    saml.idp_certificate_resolved, "EXTRACT_ONLY",
                    "Certificate existence only; trust semantics are not inferred.",
                )
        for scheme in self.ir.authentication_schemes:
            for dependency in scheme.user_database_dependencies:
                add(
                    "Authentication Scheme", scheme.name,
                    dependency.dependency_type, dependency.reference,
                    dependency.resolved, "EXTRACT_ONLY",
                    "Authentication database dependency preserved.",
                )
        for rule in self.ir.authentication_rules:
            if rule.active_auth_method and rule.active_auth_method_resolved is not None:
                add(
                    "Authentication Rule", rule.name, "authentication-scheme",
                    rule.active_auth_method, rule.active_auth_method_resolved,
                    "EXTRACT_ONLY", "Authentication scheme reference preserved.",
                )
        for admin in self._vendor_extension_value("administrators") or []:
            if admin.token_reference and admin.fortitoken_resolved is not None:
                add(
                    "Administrator", admin.name, "fortitoken",
                    admin.token_reference, admin.fortitoken_resolved,
                    "EXTRACT_ONLY", "FortiToken assignment metadata only; no token secret retained.",
                )
        settings = self._vendor_extension_value("user_authentication_settings")
        if settings is not None:
            for reference, resolved, label in (
                (settings.auth_certificate, settings.auth_certificate_resolved, "authentication certificate"),
                (settings.auth_ca_certificate, settings.auth_ca_certificate_resolved, "authentication CA certificate"),
            ):
                if reference is not None and resolved is not None:
                    add(
                        "User Authentication Settings", "global", "certificate",
                        reference, resolved, "EXTRACT_ONLY", f"{label.title()} reference preserved.",
                    )
        quarantine = self._vendor_extension_value("user_quarantine_settings")
        if quarantine is not None:
            for reference in quarantine.firewall_groups:
                add(
                    "User Quarantine", "global", "address-group", reference,
                    reference in quarantine.resolved_firewall_groups,
                    "EXTRACT_ONLY", "Quarantine firewall-group dependency preserved.",
                )
        for policy in self.ir.policies:
            for reference in policy.source_user_groups:
                add(
                    "Policy", policy.source_rule_id or policy.name, "user-group",
                    reference, reference not in policy.unresolved_user_groups,
                    "REVIEW_REQUIRED",
                    "Target identity enforcement is not normalized; policy is withheld.",
                )
            for reference in policy.source_users:
                add(
                    "Policy", policy.source_rule_id or policy.name, "local-user",
                    reference, reference not in policy.unresolved_users,
                    "REVIEW_REQUIRED",
                    "Target identity enforcement is not normalized; policy is withheld.",
                )
        for tunnel in self.ir.vpn_tunnels:
            if tunnel.source_auth_user_group:
                add(
                    "VPN Tunnel", tunnel.name, "user-group",
                    tunnel.source_auth_user_group,
                    tunnel.source_auth_user_group not in tunnel.unresolved_auth_user_groups,
                    "REVIEW_REQUIRED", "VPN authentication group reference preserved.",
                )
        ssl_vpn_settings = self._vendor_extension_value("ssl_vpn_settings")
        if ssl_vpn_settings is not None:
            for rule in ssl_vpn_settings.authentication_rules:
                for reference in rule.groups:
                    add(
                        "SSL VPN Authentication Rule", str(rule.source_id),
                        "user-group", reference,
                        reference not in rule.unresolved_groups,
                        "EXTRACT_ONLY", "SSL VPN group reference preserved.",
                    )

        self._table_sheet(
            workbook,
            "Security Identity Dependencies",
            (
                "Consumer Type", "Consumer Name", "Dependency Type", "Reference",
                "Resolved", "Dependency Status", "Migration Impact", "Notes",
            ),
            rows,
            empty_note="No Security/Identity dependencies were extracted.",
        )
        self._table_sheet(
            workbook, "SSL VPN Host Checks",
            ("Name", "Type", "OS Type", "Version", "GUID", "Check Item Count", "Extraction Status", "Manual Review", "Additional Settings"),
            (
                (
                    item.name, item.check_type, item.os_type, item.version, item.guid,
                    len(item.check_items),
                    item.migration_status, item.requires_manual_review,
                    self._format_settings(item.source_attributes),
                ) for item in self._vendor_extension_value("ssl_vpn_host_checks") or []
            ),
        )
        self._table_sheet(
            workbook, "SSL VPN Host Check Items",
            ("Host Check", "ID", "Action", "Type", "Target", "MD5s", "Version", "Extraction Status", "Manual Review", "Additional Settings"),
            (
                (
                    host_check.name, item.source_id, item.action, item.check_type,
                    item.target, item.md5s, item.version, item.migration_status,
                    item.requires_manual_review,
                    self._format_settings(item.source_attributes),
                )
                for host_check in self._vendor_extension_value("ssl_vpn_host_checks") or []
                for item in host_check.check_items
            ),
        )

    def _build_administrator_inventory(self, workbook: Any) -> None:
        self._table_sheet(
            workbook,
            "Administrators",
            (
                "Name", "Access Profile", "VDOMs", "IPv4 Trusted Hosts", "IPv6 Trusted Hosts",
                "Two Factor", "FortiToken", "Guest User Groups", "Remote Auth", "Remote Group",
                "FortiToken Resolved", "Access Profile Resolved", "Unresolved References",
                "Schedule", "Peer Auth", "Peer Group", "SSH Certificate", "SSH Public Keys",
                "Credential Configured", "Migration Status", "Manual Review",
                "Authentication Profile", "Authentication Profile Resolved",
                "Authentication Sequence", "Authentication Sequence Resolved",
                "Additional Settings",
            ),
            (
                (
                    item.name, item.access_profile, item.vdoms, item.trusted_hosts_ipv4,
                    item.trusted_hosts_ipv6, item.two_factor, item.token_reference,
                    item.guest_user_groups, item.remote_auth, item.remote_group,
                    self._optional_bool_literal(item.fortitoken_resolved),
                    self._optional_bool_literal(item.access_profile_resolved),
                    item.unresolved_references,
                    item.schedule, item.peer_auth, item.peer_group, item.ssh_certificate,
                    item.ssh_public_keys,
                    item.credential_configured, item.migration_status,
                    item.requires_manual_review,
                    item.authentication_profile, self._optional_bool_literal(item.authentication_profile_resolved),
                    item.authentication_sequence, self._optional_bool_literal(item.authentication_sequence_resolved),
                    self._format_settings(item.source_attributes),
                )
                for item in self._vendor_extension_value("administrators") or []
            ),
        )
        self._table_sheet(
            workbook,
            "Admin Profiles",
            ("Name", "Migration Status", "Manual Review", "Additional Settings"),
            (
                (
                    item.name, item.migration_status, item.requires_manual_review,
                    self._format_settings(item.source_attributes),
                )
                for item in self._vendor_extension_value("admin_profiles") or []
            ),
        )
    def _structured_routing_dependency_items(self) -> list[Any]:
        if self.extraction is None:
            return []
        return [
            item
            for item in self.extraction.inventory_items
            if "structured-routing-dependency" in item.notes
        ]

    def _build_routing_dependencies(self, workbook: Any) -> None:
        items = self._structured_routing_dependency_items()
        self._table_sheet(
            workbook,
            "Routing Dependencies",
            (
                "Type",
                "Name / Source ID",
                "Source Path",
                "Source Block Present",
                "Configured",
                "Extraction Status",
                "Manual Review",
            ),
            (
                (
                    item.source_path.removeprefix("router "),
                    item.name or item.source_id,
                    item.source_path,
                    "Yes",
                    (
                        "Yes"
                        if self._flatten_source_profile_settings(item)
                        else "No"
                    ),
                    item.status,
                    self._optional_bool_literal(item.requires_manual_review),
                )
                for item in items
            ),
        )
        self._table_sheet(
            workbook,
            "Routing Dependency Settings",
            (
                "Type",
                "Object",
                "Parent / Subsection",
                "Entry",
                "Operation",
                "Setting",
                "Value",
            ),
            (
                (row[0].removeprefix("router "), *row[1:])
                for item in items
                for row in self._flatten_source_profile_settings(item)
            ),
        )
        self._table_sheet(
            workbook,
            "Admin Profile Permissions",
            ("Profile", "Permission Group", "Setting", "Value", "Extraction Status", "Additional Settings"),
            (
                (
                    profile.name, block.name, setting, value, "EXTRACT_ONLY",
                    self._format_settings(block.source_attributes),
                )
                for profile in self._vendor_extension_value("admin_profiles") or []
                for block in profile.permission_blocks
                for setting, value in {**block.settings, **block.source_attributes}.items()
            ),
            empty_note="No admin profile permissions were extracted from the source configuration.",
        )
    def _build_dos_inventory(self, workbook: Any) -> None:
        self._table_sheet(
            workbook, "DoS Policies",
            ("Policy ID", "Policy Name", "Status", "Interface", "Source Addresses", "Destination Addresses", "Services", "Anomaly Count", "Description", "Extraction Status", "Manual Review", "Additional Settings"),
            (
                (
                    item.source_id, item.name, item.status, item.interface, item.source_addresses,
                    item.destination_addresses, item.services, len(item.anomalies),
                    item.description, item.migration_status, item.requires_manual_review,
                    self._format_settings(item.source_attributes),
                ) for item in self._vendor_extension_value("dos_policies") or []
            ),
        )
        self._table_sheet(
            workbook, "DoS Anomalies",
            ("Policy ID", "Name", "Status", "Log", "Action", "Quarantine", "Quarantine Expiry", "Quarantine Log", "Threshold", "Default Threshold", "Additional Settings"),
            (
                (
                    policy.source_id, item.name, item.status, item.log, item.action,
                    item.quarantine, item.quarantine_expiry, item.quarantine_log,
                    item.threshold, item.threshold_default,
                    self._format_settings(item.source_attributes),
                ) for policy in self._vendor_extension_value("dos_policies") or [] for item in policy.anomalies
            ),
        )

    def _build_firewall_sniffers(self, workbook: Any) -> None:
        self._table_sheet(
            workbook, "Firewall Sniffer",
            ("ID", "Source UUID", "Log Traffic", "IPv6", "Non-IP", "Application List Status", "Application List", "IPS Sensor Status", "IPS Sensor", "AV Profile Status", "AV Profile", "Web Filter Status", "Web Filter Profile", "Extraction Status", "Manual Review", "Additional Settings"),
            (
                (
                    item.source_id, item.source_uuid, item.logtraffic, item.ipv6,
                    item.non_ip, item.application_list_status, item.application_list,
                    item.ips_sensor_status, item.ips_sensor, item.av_profile_status,
                    item.av_profile, item.webfilter_profile_status,
                    item.webfilter_profile, item.migration_status,
                    item.requires_manual_review, self._format_settings(item.source_attributes),
                ) for item in self._vendor_extension_value("firewall_sniffers") or []
            ),
        )

    def _build_authentication_inventory(self, workbook: Any) -> None:
        self._table_sheet(
            workbook, "Authentication Schemes",
            ("Name", "Method", "User Database", "Resolved User Databases", "Unresolved User Databases", "Extraction Status", "Manual Review", "Additional Settings"),
            (
                (
                    item.name, item.method, item.user_database,
                    item.resolved_user_databases, item.unresolved_user_databases,
                    item.migration_status,
                    item.requires_manual_review, self._format_settings(item.source_attributes),
                ) for item in self.ir.authentication_schemes
            ),
        )
        self._table_sheet(
            workbook, "Authentication Rules",
            ("Name", "Source Interfaces", "Source Addresses", "Active Auth Method", "Auth Method Resolved", "Unresolved Auth Methods", "Extraction Status", "Manual Review", "Additional Settings"),
            (
                (
                    item.name, item.source_interfaces, item.source_addresses,
                    item.active_auth_method,
                    self._optional_bool_literal(item.active_auth_method_resolved),
                    item.unresolved_auth_methods, item.migration_status,
                    item.requires_manual_review, self._format_settings(item.source_attributes),
                ) for item in self.ir.authentication_rules
            ),
        )

    def _build_pan_phase9_sheets(self, workbook: Any) -> None:
        pan_ext = self.ir.vendor_extensions.panos
        def add(name, headers, rows):
            self._table_sheet(workbook, name, headers, rows,
                              empty_note="No PAN-OS Phase 9 source objects were extracted.",
                              subtitle="PAN-OS source-only inventory; not target configuration.")
        add("PAN Log Servers", ("Profile", "Type", "Endpoint", "Address", "Transport", "Port", "Format", "Facility", "SNMP Version", "Community Configured", "Auth Password Configured", "Privacy Password Configured"), [
            (p.name, p.profile_type, e.name, e.address, e.transport, e.port, e.format, e.facility, e.snmp_version, e.community_configured, e.authentication_password_configured, e.privacy_password_configured)
            for p in pan_ext.pan_log_server_profiles for e in p.servers
        ])
        add("PAN Log Forwarding", ("Profile", "Match Count", "Status"), [(p.name, len(p.matches), p.migration_status) for p in pan_ext.pan_log_forwarding_profiles])
        add("PAN Log Forward Matches", ("Profile", "Match", "Log Type", "Filter", "Send to Panorama", "Syslog", "Email", "SNMP", "HTTP", "Resolution State"), [
            (p.name, m.name, m.log_type, m.filter, m.send_to_panorama, ", ".join(m.syslog_profiles), ", ".join(m.email_profiles), ", ".join(m.snmptrap_profiles), ", ".join(m.http_profiles), "review" if m.review_reasons else "retained")
            for p in pan_ext.pan_log_forwarding_profiles for m in p.matches
        ])
        add("PAN DNS Proxies", ("Name", "Enabled", "Cache", "Max TTL", "Default Primary", "Default Secondary", "TCP Queries", "Interfaces", "Resolved Interfaces"), [(p.name,p.enabled,p.cache_enabled,p.max_ttl_enabled,p.default_primary,p.default_secondary,p.tcp_queries_enabled,", ".join(p.interfaces),", ".join(p.resolved_interfaces)) for p in pan_ext.pan_dns_proxies])
        add("PAN DNS Proxy Domains", ("Proxy", "Domain Rule", "Domains", "Primary", "Secondary", "Cacheable"), [(p.name,d.name,", ".join(d.domain_names),d.primary,d.secondary,d.cacheable) for p in pan_ext.pan_dns_proxies for d in p.domain_servers])
        add("PAN Monitor Profiles", ("Name", "Interval Seconds", "Threshold", "Action"), [(p.name,p.interval_seconds,p.threshold,p.action) for p in pan_ext.pan_monitor_profiles])
        add("PAN QoS Profiles", ("Name", "Bandwidth Type", "Class Count"), [(p.name,p.bandwidth_type,len(p.classes)) for p in pan_ext.pan_qos_profiles])
        add("PAN QoS Classes", ("Profile", "Class", "Priority", "Egress Max", "Egress Guaranteed"), [(p.name,c.name,c.priority,c.egress_max,c.egress_guaranteed) for p in pan_ext.pan_qos_profiles for c in p.classes])
        ha=pan_ext.pan_high_availability
        add("PAN High Availability", ("Context", "Enabled", "Group ID", "Description", "Peer IP", "Preemptive", "HA2 Keepalive"), [] if ha is None else [(ha.source_context,ha.enabled,ha.group_id,ha.description,ha.peer_ip,ha.preemptive,ha.ha2_keep_alive_enabled)])
        add("PAN HA Monitoring", ("Kind", "Name", "References", "Resolved", "Unresolved", "Destinations"), [] if ha is None else [("link",g.name,", ".join(g.interfaces),", ".join(g.resolved_interfaces),", ".join(g.unresolved_interfaces),"") for g in ha.link_groups]+[("path",g.name,g.routing_instance or "",g.resolved_routing_instance or "", "", ", ".join(g.destination_ips)) for g in ha.path_groups])
        ds=pan_ext.pan_device_operational_settings
        add("PAN Device Settings", ("Context", "Rematch", "Hostname Type", "Commit Lock", "WildFire Benign", "WildFire Grayware", "Urgent Data", "Asymmetric Path", "Default Timeout", "TCP Timeout"), [] if ds is None else [(ds.source_context,ds.rematch_sessions,ds.hostname_type_in_syslog,ds.auto_acquire_commit_lock,ds.wildfire_report_benign_file,ds.wildfire_report_grayware_file,ds.tcp_urgent_data,ds.tcp_asymmetric_path,ds.session_timeout_default_seconds,ds.session_timeout_tcp_seconds)])
        add("PAN VSYS Settings", ("Context", "Allow Forward Decrypted Content"), [(p.source_context,p.allow_forward_decrypted_content) for p in pan_ext.pan_vsys_settings])
        br=pan_ext.pan_botnet_report_settings
        add("PAN Botnet Report", ("Dynamic DNS", "Malware", "Recent Domains", "IP Domains", "Unknown Sites", "IRC", "Top N", "Scheduled"), [] if br is None else [(br.dynamic_dns_enabled,br.malware_sites_enabled,br.recent_domains_enabled,br.ip_domains_enabled,br.executables_unknown_sites_enabled,br.irc_enabled,br.topn,br.scheduled)])
        add("PAN Custom Reports", ("Name", "Type", "Sort By", "Group By", "Aggregate Count", "Top N", "Top M", "Caption", "Start", "End"), [(p.name,p.report_type,p.sort_by,p.group_by,len(p.aggregate_by),p.topn,p.topm,p.caption,p.start_time,p.end_time) for p in pan_ext.pan_custom_reports])

    def _build_pan_sdwan_sheets(self, workbook: Any) -> None:
        pan_ext = self.ir.vendor_extensions.panos
        self._table_sheet(
            workbook, "PAN SD-WAN Interface Profiles",
            ("Name", "Source Context", "VPN Failover Metric", "Probe Settings", "Migration Status", "Manual Review", "Review Reasons"),
            ((p.name, p.source_context, p.vpn_failover_metric, p.probe_settings,
              p.migration_status, p.requires_manual_review, p.review_reasons)
             for p in pan_ext.pan_sdwan_interface_profiles),
            empty_note="No PAN-OS SD-WAN interface profiles were extracted.",
        )
        self._table_sheet(
            workbook, "PAN SD-WAN Link Settings",
            ("Interface", "Interface Profile", "Path Quality Profile", "Traffic Distribution Profile", "SaaS Quality Profile", "Migration Status", "Manual Review", "Review Reasons"),
            ((p.interface, p.interface_profile, p.path_quality_profile,
              p.traffic_distribution_profile, p.saas_quality_profile,
              p.migration_status, p.requires_manual_review, p.review_reasons)
             for p in pan_ext.pan_sdwan_link_settings),
            empty_note="No PAN-OS SD-WAN link settings were extracted.",
        )
        self._table_sheet(
            workbook, "PAN SD-WAN Path Quality",
            ("Name", "Source Context", "Latency", "Jitter", "Packet Loss", "Sensitivity", "Migration Status", "Manual Review", "Review Reasons"),
            ((p.name, p.source_context, p.latency, p.jitter, p.packet_loss,
              p.sensitivity, p.migration_status, p.requires_manual_review, p.review_reasons)
             for p in pan_ext.pan_sdwan_path_quality_profiles),
            empty_note="No PAN-OS SD-WAN path-quality profiles were extracted.",
        )
        self._table_sheet(
            workbook, "PAN SD-WAN Traffic Distribution",
            ("Name", "Source Context", "Method", "Link Tags", "Weights", "Migration Status", "Manual Review", "Review Reasons"),
            ((p.name, p.source_context, p.method, p.link_tags, p.weights,
              p.migration_status, p.requires_manual_review, p.review_reasons)
             for p in pan_ext.pan_sdwan_traffic_distribution_profiles),
            empty_note="No PAN-OS SD-WAN traffic-distribution profiles were extracted.",
        )
        self._table_sheet(
            workbook, "PAN SD-WAN Rules",
            ("Rule #", "Name", "Source Context", "Rulebase Position", "Source Order", "Source", "Destination", "Source User", "Applications", "Services", "Input Interfaces", "Input Zones", "Path Quality Profile", "Traffic Distribution Profile", "SaaS Quality Profile", "Action", "Disabled", "Migration Status", "Manual Review", "Review Reasons"),
            ((i, p.name, p.source_context, p.rulebase_position, p.source_order,
              p.source, p.destination, p.source_user, p.application, p.service,
              p.input_interface, p.input_zone, p.path_quality_profile,
              p.traffic_distribution_profile, p.saas_quality_profile, p.action,
              p.disabled, p.migration_status, p.requires_manual_review,
              p.review_reasons)
             for i, p in enumerate(pan_ext.pan_sdwan_rules, 1)),
            empty_note="No PAN-OS SD-WAN rules were extracted.",
        )

    def _build_extraction_coverage(self, workbook: Any) -> None:
        if self.extraction is not None:
            rows = (
                (
                    section.path,
                    section.present,
                    section.object_count_source,
                    section.object_count_parsed,
                    section.object_count_normalized,
                    section.status.value,
                    fortigate_semantic_support_level(section.path),
                    section.parser_handler,
                    section.line_start,
                    section.line_end,
                    "; ".join(section.semantic_unknowns),
                    section.unresolved_dependencies,
                    "; ".join(section.notes),
                )
                for section in self.extraction.source_sections
            )
            sheet = self._table_sheet(
                workbook,
                "Extraction Coverage",
                (
                    "Source Section",
                    "Found",
                    "Source Objects",
                    "Parsed Objects",
                    "Normalized Objects",
                    "Status",
                    "Semantic Level",
                    "Parser Handler",
                    "Line Start",
                    "Line End",
                    "Semantic Unknowns",
                    "Unresolved Dependencies",
                    "Notes",
                ),
                rows,
                empty_note="No FortiGate config sections were discovered in the source.",
                subtitle="Coverage correlates independent source discovery with typed parsing and canonical IR.",
            )
            for row in range(4, sheet.max_row + 1):
                status = str(sheet.cell(row, 6).value or "").lower()
                if "unsupported" in status or "parse_error" in status:
                    fill = self._LIGHT_RED
                elif "partial" in status or "extract_only" in status:
                    fill = self._LIGHT_AMBER
                else:
                    continue
                for column in range(1, 14):
                    sheet.cell(row, column).fill = PatternFill("solid", fgColor=fill)
            return

        collections = (
            ("Interfaces", self.ir.interfaces),
            (
                "Interface Secondary IPs",
                sum(
                    len(getattr(intf, "secondary_ips", []))
                    for intf in self.ir.interfaces
                ),
            ),
            (
                "DHCP Servers",
                self.ir.dhcp_servers,
            ),
            (
                "DHCP IP Ranges",
                sum(len(server.ip_ranges) for server in self.ir.dhcp_servers),
            ),
            (
                "DHCP Reservations",
                sum(len(server.reservations) for server in self.ir.dhcp_servers),
            ),
            ("Zones", self.ir.zones),
            ("Addresses", self.ir.addresses),
            (
                "Address Groups",
                self.ir.address_groups,
            ),
            ("Proxy Addresses", self.ir.proxy_addresses),
            (
                "Service Categories",
                self.ir.service_categories,
            ),
            ("Services", self.ir.services),
            (
                "Service Groups",
                self.ir.service_groups,
            ),
            (
                "Session TTL Overrides",
                self._vendor_extension_value("session_ttl_overrides") or [],
            ),
            ("Schedules", self.ir.schedules),
            ("Policies", self.ir.policies),
            (
                "ZTNA Providers",
                self.ir.ztna_providers,
            ),
            ("IP Pools", self.ir.ip_pools),
            ("Virtual IPs", self.ir.virtual_ips),
            ("VIP Groups", self.ir.virtual_ip_groups),
            (
                "VIP Real Servers",
                sum(len(vip.real_servers) for vip in self.ir.virtual_ips),
            ),
            ("NAT Rules", self.ir.nat_rules),
            ("VPN Tunnels", self.ir.vpn_tunnels),
            ("VPN Phase 2", self.ir.vpn_phase2),
            ("SSL VPN Portals", self._vendor_extension_value("ssl_vpn_portals") or []),
            ("SSL VPN Host Checks", self._vendor_extension_value("ssl_vpn_host_checks") or []),
            (
                "SSL VPN Host Check Items",
                sum(
                    len(host_check.check_items)
                    for host_check in self._vendor_extension_value("ssl_vpn_host_checks") or []
                ),
            ),
            (
                "SSL VPN Settings",
                [] if self._vendor_extension_value("ssl_vpn_settings") is None else [self._vendor_extension_value("ssl_vpn_settings")],
            ),
            ("SD-WAN", self._vendor_extension_value("sdwans") or []),
            ("LDAP Servers", self._vendor_extension_value("user_ldap_servers") or []),
            ("SAML Servers", self._vendor_extension_value("user_saml_servers") or []),
            ("FSSO Servers", self._vendor_extension_value("fsso_providers") or []),
            ("FSSO AD Groups", self._vendor_extension_value("fsso_ad_groups") or []),
            ("Local Users", self._vendor_extension_value("local_users") or []),
            ("User Groups", self._vendor_extension_value("user_groups") or []),
            ("DoS Policies", self._vendor_extension_value("dos_policies") or []),
            ("Firewall Sniffer", self._vendor_extension_value("firewall_sniffers") or []),
            ("Authentication Schemes", self.ir.authentication_schemes),
            ("Authentication Sequences", self.ir.authentication_sequences),
            ("Authentication Rules", self.ir.authentication_rules),
            ("Certificates", self.ir.certificates),
            ("SSL TLS Service Profiles", self._vendor_extension_value("ssl_tls_service_profiles") or []),
            ("Routes", self.ir.routes),
            (
                "Security Profiles",
                self.ir.security_profile_groups,
            ),
        )
        def collection_count(items: Any) -> int:
            return items if isinstance(items, int) else len(items)

        rows = [
            (
                name,
                None,
                collection_count(items),
                "Not reported",
                "Populated" if collection_count(items) else "Empty / unknown",
                "Phase 1 exports IR only; source-section coverage awaits ExtractionResult in Phase 2.",
            )
            for name, items in collections
        ]
        note = "Source-section evidence is unavailable in IR; inferred counts are not proof of parser completeness."
        sheet = self._table_sheet(
            workbook,
            "Extraction Coverage",
            ("Source Section", "Found", "Objects", "Parser Status", "IR Status", "Notes"),
            rows,
            subtitle=note,
        )
        for row in range(4, sheet.max_row + 1):
            status = str(sheet.cell(row, 5).value or "").lower()
            if "unsupported" in status or "not transformed" in status:
                fill = self._LIGHT_RED
            elif "partial" in status or "unknown" in status:
                fill = self._LIGHT_AMBER
            else:
                continue
            for column in range(1, 7):
                sheet.cell(row, column).fill = PatternFill("solid", fgColor=fill)

    def _build_unresolved_references(self, workbook: Any) -> None:
        """Export the complete dependency registry without an unresolved-only view."""
        dependencies = list(self.extraction.dependencies) if self.extraction is not None else []
        headers = (
            "Source VDOM", "Source Type", "Source Object", "Field",
            "Reference", "Expected Type", "Result", "Target Path", "Target UID",
            "Target Name", "Semantic Kind", "Normalization Status", "Reason", "Notes",
        )
        rows = [
            (
                dependency.source_context or "root",
                dependency.source_path,
                dependency.source_object or "",
                dependency.source_field,
                dependency.reference,
                dependency.expected_type,
                dependency.result,
                dependency.target_path or "",
                dependency.target_uid or "",
                dependency.target_name or "",
                dependency.semantic_kind or "",
                dependency.normalization_status or "",
                dependency.reason or "",
                dependency.notes or "",
            )
            for dependency in dependencies
        ]
        self._table_sheet(
            workbook,
            "Dependency Registry",
            headers,
            rows,
            empty_note="No FortiGate dependency references were discovered.",
            subtitle="References are resolved within the source VDOM/context; unresolved entries require manual review.",
        )

    def _build_phase7_identity_sheets(self, workbook: Any) -> None:
        self._table_sheet(workbook, "Identity Server Endpoints",
            ("Profile Type", "Profile Name", "Endpoint Name", "Address", "Port", "Secret Configured", "Extraction Status", "Manual Review", "Additional Settings"),
            ((kind, item.name, endpoint.name, endpoint.address, endpoint.port, endpoint.has_secret,
              item.migration_status, item.requires_manual_review, self._format_settings(endpoint.source_attributes))
             for kind, items in (("LDAP", self._vendor_extension_value("user_ldap_servers") or []), ("RADIUS", self._vendor_extension_value("user_radius_servers") or []), ("TACACS+", self._vendor_extension_value("user_tacacs_servers") or []))
             for item in items for endpoint in item.server_entries))
        self._table_sheet(workbook, "Authentication Sequences",
            ("Name", "Source Context", "Authentication Profiles", "Resolved Authentication Profiles", "Unresolved Authentication Profiles", "Extraction Status", "Manual Review", "Review Reasons", "Additional Settings"),
            ((item.name, item.source_context, item.authentication_profiles, item.resolved_authentication_profiles,
              item.unresolved_authentication_profiles, item.migration_status, item.requires_manual_review,
              item.review_reasons, self._format_settings(item.source_attributes)) for item in self.ir.authentication_sequences))
        self._table_sheet(workbook, "SSL TLS Service Profiles",
            ("Name", "Source Context", "Certificate", "Certificate Resolved", "Source Certificate Profile Reference", "Certificate Profile Name Resolved", "Minimum TLS Version", "Maximum TLS Version", "Extraction Status", "Manual Review", "Review Reasons", "Additional Settings"),
            ((item.name, item.source_context, item.certificate, self._optional_bool_literal(item.certificate_resolved),
              item.certificate_profile, self._optional_bool_literal(item.certificate_profile_resolved),
              item.minimum_tls_version, item.maximum_tls_version, item.migration_status, item.requires_manual_review,
              item.review_reasons, self._format_settings(item.source_attributes)) for item in self._vendor_extension_value("ssl_tls_service_profiles") or []))

    def _build_globalprotect_sheets(self, workbook: Any) -> None:
        pan_ext = self.ir.vendor_extensions.panos
        portals = pan_ext.global_protect_portals
        gateways = pan_ext.global_protect_gateways
        network_gateways = pan_ext.global_protect_network_gateways
        self._table_sheet(workbook, "GlobalProtect Portals", (
            "Name", "Source Context", "Local Interface", "Local Interface Resolved", "Local IPv4",
            "Local IPv6", "Local Address Resolved", "SSL/TLS Service Profile", "SSL/TLS Profile Resolved",
            "Custom Login Page", "Custom Home Page", "Client Auth Count", "Client Config Count",
            "Root CA Count", "Agent Override Key Configured", "Extraction Status", "Manual Review",
            "Review Reasons", "Additional Settings",
        ), ((item.name, item.source_context, item.local_interface,
              self._optional_bool_literal(item.local_interface_resolved), item.local_ipv4, item.local_ipv6,
              self._optional_bool_literal(item.local_address_resolved), item.ssl_tls_service_profile,
              self._optional_bool_literal(item.ssl_tls_service_profile_resolved), item.custom_login_page,
              item.custom_home_page, len(item.client_authentication), len(item.client_configs),
              len(item.root_ca_certificates), self._optional_bool_literal(item.has_agent_user_override_key),
              item.migration_status, self._optional_bool_literal(item.requires_manual_review), item.review_reasons,
              self._format_settings(item.source_attributes)) for item in portals))
        self._table_sheet(workbook, "GlobalProtect Gateways", (
            "Name", "Source Context", "SSL/TLS Service Profile", "SSL/TLS Profile Resolved", "Tunnel Mode",
            "Remote User Tunnel", "Remote User Tunnel Resolved", "Role Count", "Client Auth Count",
            "Tunnel Config Count", "Extraction Status", "Manual Review", "Review Reasons", "Additional Settings",
        ), ((item.name, item.source_context, item.ssl_tls_service_profile,
              self._optional_bool_literal(item.ssl_tls_service_profile_resolved), self._optional_bool_literal(item.tunnel_mode),
              item.remote_user_tunnel, self._optional_bool_literal(item.remote_user_tunnel_resolved), len(item.roles),
              len(item.client_authentication), len(item.remote_user_tunnel_configs), item.migration_status,
              self._optional_bool_literal(item.requires_manual_review), item.review_reasons,
              self._format_settings(item.source_attributes)) for item in gateways))
        self._table_sheet(workbook, "GlobalProtect Client Auth", (
            "Consumer Type", "Consumer Name", "Source Context", "Name", "OS", "Authentication Profile",
            "Authentication Profile Resolved", "Resolved Authentication Profile", "Authentication Message",
            "Username Label", "Password Label", "Manual Review", "Review Reasons", "Additional Settings",
        ), ((kind, owner.name, owner.source_context, auth.name, auth.os, auth.authentication_profile,
              self._optional_bool_literal(auth.authentication_profile_resolved), auth.resolved_authentication_profile,
              auth.authentication_message, auth.username_label, auth.password_label,
              self._optional_bool_literal(owner.requires_manual_review or bool(auth.review_reasons)), auth.review_reasons,
              self._format_settings(auth.source_attributes))
             for kind, owners in (("Portal", portals), ("Gateway", gateways))
             for owner in owners for auth in owner.client_authentication))
        self._table_sheet(workbook, "GlobalProtect Portal Configs", (
            "Portal", "Source Context", "Name", "Source Users", "Operating Systems", "External Gateway Count",
            "External Gateway Cutoff Time", "Generate Cookie", "Max Agent User Overrides", "Agent Override Timeout",
            "HIP Collect Data", "HIP Max Wait Time", "Save User Credentials", "Portal 2FA",
            "Manual-Only Gateway 2FA", "Internal Gateway 2FA", "Auto-Discovery External Gateway 2FA",
            "MDM Enrollment Port", "Extraction Status", "Manual Review", "Review Reasons", "Additional Settings",
        ), ((portal.name, portal.source_context, config.name, config.source_users, config.operating_systems,
              len(config.external_gateways), config.external_gateway_cutoff_time,
              self._optional_bool_literal(config.authentication_override_generate_cookie), config.max_agent_user_overrides,
              config.agent_user_override_timeout, self._optional_bool_literal(config.hip_collect_data), config.hip_max_wait_time,
              config.save_user_credentials, self._optional_bool_literal(config.portal_2fa),
              self._optional_bool_literal(config.manual_only_gateway_2fa), self._optional_bool_literal(config.internal_gateway_2fa),
              self._optional_bool_literal(config.auto_discovery_external_gateway_2fa), config.mdm_enrollment_port,
              "EXTRACT_ONLY", self._optional_bool_literal(bool(config.review_reasons)), config.review_reasons,
              self._format_settings(config.source_attributes)) for portal in portals for config in portal.client_configs))
        self._table_sheet(workbook, "GlobalProtect External Gateways", (
            "Portal", "Client Config", "Name", "IPv4", "IPv6", "Manual", "Priority Rules", "Additional Settings",
        ), ((portal.name, config.name, gateway.name, gateway.ipv4, gateway.ipv6,
              self._optional_bool_literal(gateway.manual),
              [(rule.name, rule.priority) for rule in gateway.priority_rules],
              self._format_settings(gateway.source_attributes))
             for portal in portals for config in portal.client_configs for gateway in config.external_gateways))
        self._table_sheet(workbook, "GlobalProtect App Settings", (
            "Portal", "Client Config", "Order", "Setting Name", "Values", "Additional Settings",
        ), ((portal.name, config.name, setting.source_order, setting.name, setting.values,
              self._format_settings(setting.source_attributes))
             for portal in portals for config in portal.client_configs for setting in config.app_settings))
        self._table_sheet(workbook, "GlobalProtect Root CAs", (
            "Portal", "Source Context", "Certificate", "Certificate Resolved", "Resolved Certificate",
            "Install in Cert Store", "Manual Review", "Review Reasons", "Additional Settings",
        ), ((portal.name, portal.source_context, ca.certificate, self._optional_bool_literal(ca.certificate_resolved),
              ca.resolved_certificate, self._optional_bool_literal(ca.install_in_cert_store),
              self._optional_bool_literal(bool(ca.review_reasons)), ca.review_reasons,
              self._format_settings(ca.source_attributes)) for portal in portals for ca in portal.root_ca_certificates))
        self._table_sheet(workbook, "GlobalProtect Gateway Roles", (
            "Gateway", "Source Context", "Name", "Login Lifetime Days", "Inactivity Logout Hours",
            "Disconnect on Idle Minutes", "Additional Settings",
        ), ((gateway.name, gateway.source_context, role.name, role.login_lifetime_days,
              role.inactivity_logout_hours, role.disconnect_on_idle_minutes,
              self._format_settings(role.source_attributes)) for gateway in gateways for role in gateway.roles))
        self._table_sheet(workbook, "GlobalProtect Tunnel Configs", (
            "Gateway", "Source Context", "Name", "Source Users", "Operating Systems", "IP Pools",
            "Split Include Routes", "Resolved Split Include Routes", "Unresolved Split Include Routes",
            "Split Exclude Routes", "Resolved Split Exclude Routes", "Unresolved Split Exclude Routes",
            "Retrieve Framed IP", "No Direct Access to Local Network", "Extraction Status", "Manual Review",
            "Review Reasons", "Additional Settings",
        ), ((gateway.name, gateway.source_context, config.name, config.source_users, config.operating_systems,
              config.ip_pools, config.split_include_routes, config.resolved_split_include_routes,
              config.unresolved_split_include_routes, config.split_exclude_routes, config.resolved_split_exclude_routes,
              config.unresolved_split_exclude_routes, self._optional_bool_literal(config.retrieve_framed_ip_address),
              self._optional_bool_literal(config.no_direct_access_to_local_network), config.migration_status,
              self._optional_bool_literal(config.requires_manual_review), config.review_reasons,
              self._format_settings(config.source_attributes)) for gateway in gateways for config in gateway.remote_user_tunnel_configs))
        self._table_sheet(workbook, "GlobalProtect Network Gateways", (
            "Name", "Source Context", "Local Interface", "Local Interface Resolved", "Tunnel Interface",
            "Tunnel Interface Resolved", "IP Pools", "DNS Primary", "DNS Secondary", "DNS Suffixes",
            "DNS Suffix Inherited", "Third-Party Client Enabled", "Third-Party Group Name",
            "Third-Party Group Password Configured", "Extraction Status", "Manual Review", "Review Reasons",
            "Additional Settings",
        ), ((item.name, item.source_context, item.local_interface,
              self._optional_bool_literal(item.local_interface_resolved), item.tunnel_interface,
              self._optional_bool_literal(item.tunnel_interface_resolved), item.ip_pools, item.client_dns_primary,
              item.client_dns_secondary, item.dns_suffixes, self._optional_bool_literal(item.dns_suffix_inherited),
              self._optional_bool_literal(item.third_party_client_enabled), item.third_party_group_name,
              self._optional_bool_literal(item.third_party_group_password_configured), item.migration_status,
              self._optional_bool_literal(item.requires_manual_review), item.review_reasons,
              self._format_settings(item.source_attributes)) for item in network_gateways))

    def _table_styles(self) -> dict[str, Any]:
        styles = getattr(self, "_table_style_registry", None)
        if styles is None:
            thin = Side(style="thin", color=self._BORDER)
            styles = self._table_style_registry = {
                "title_font": Font(
                    name="Aptos Display", size=16, bold=True, color=self._WHITE
                ),
                "title_fill": PatternFill("solid", fgColor=self._NAVY),
                "title_alignment": Alignment(vertical="center"),
                "subtitle_font": Font(
                    name="Aptos", size=9, italic=True, underline="single", color="0563C1"
                ),
                "subtitle_alignment": Alignment(wrap_text=True, vertical="center"),
                "header_font": Font(name="Aptos", bold=True, color=self._WHITE),
                "header_fill": PatternFill("solid", fgColor=self._TEAL),
                "header_alignment": Alignment(wrap_text=True, vertical="center"),
                "body_font": Font(name="Aptos", size=10, color=self._TEXT),
                "body_alignment": Alignment(wrap_text=True, vertical="top"),
                "stripe_fill": PatternFill("solid", fgColor="F8FAFC"),
                "header_border": Border(bottom=thin),
            }
        return styles

    def _register_partitioned_sheets(self, title: str, names: list[str]) -> None:
        previous = self._partitioned_sheet_names.get(title, ())
        partitioned = tuple(names)
        self._partitioned_sheet_names[title] = partitioned
        order = list(self.SHEET_ORDER)
        anchor = previous[0] if previous else title
        if anchor not in order:
            return
        index = order.index(anchor)
        order[index:index + max(len(previous), 1)] = names
        self.SHEET_ORDER = tuple(order)

    def _is_full_profile(self) -> bool:
        return self._export_options.profile is ExcelExportProfile.FULL

    def _is_fast_profile(self) -> bool:
        return self._export_options.profile is ExcelExportProfile.FAST

    def _is_data_only_profile(self) -> bool:
        return self._export_options.profile is ExcelExportProfile.DATA_ONLY

    def _expanded_sheet_order(self, order: Sequence[str]) -> tuple[str, ...]:
        return tuple(
            partition
            for name in order
            for partition in self._partitioned_sheet_names.get(name, (name,))
        )

    def _table_sheet(
        self,
        workbook: Any,
        title: str,
        headers: Sequence[str],
        rows: Iterable[Sequence[Any]],
        empty_note: str = (
            "No objects were represented in this IR collection."
        ),
        subtitle: str = (
            "Vendor-neutral IR inventory exported before migration optimization."
        ),
    ) -> Any:
        row_iterator = iter(rows)
        sentinel = object()
        first_row = next(row_iterator, sentinel)
        has_rows = first_row is not sentinel
        if has_rows:
            from itertools import chain

            row_iterator = chain((first_row,), row_iterator)

        styles = self._table_styles()
        minimal = self._is_data_only_profile()
        fast = self._is_fast_profile()
        full = self._is_full_profile()
        metrics_enabled = getattr(self, "_last_export_metrics", None) is not None
        limit = self._export_options.max_rows_per_sheet or self.MAX_ROWS_PER_DATA_SHEET
        sheets: list[tuple[Any, int]] = []
        names = [title]
        sheet_starts: dict[int, float] = {}
        sheet_nonempty: dict[int, int] = {}
        classifiers: dict[str, AuditSheetClassifier] = {}
        category = getattr(self, "_sheet_category", lambda _: "Inventory")
        if self._audit_accumulator is not None and has_rows:
            classifiers[title] = build_audit_classifier(title, headers, category)

        def create_sheet(sheet_title: str, note: str) -> Any:
            sheet = workbook.create_sheet(sheet_title)
            sheet.sheet_view.showGridLines = False
            last_column = get_column_letter(len(headers))
            sheet.merge_cells(f"A1:{last_column}1")
            sheet["A1"] = title
            if not minimal:
                sheet["A1"].font = styles["title_font"]
                sheet["A1"].fill = styles["title_fill"]
                sheet["A1"].alignment = styles["title_alignment"]
            sheet.row_dimensions[1].height = 30
            sheet.merge_cells(f"A2:{last_column}2")
            sheet["A2"] = self._safe_value(f"Back to Summary  |  {note}")
            self._set_internal_link(sheet["A2"], "Summary")
            if not minimal:
                sheet["A2"].font = styles["subtitle_font"]
                sheet["A2"].alignment = styles["subtitle_alignment"]
            sheet.row_dimensions[2].height = 26
            for column, header in enumerate(headers, 1):
                cell = sheet.cell(3, column, str(header))
                if not minimal:
                    cell.font = styles["header_font"]
                    cell.fill = styles["header_fill"]
                    cell.alignment = styles["header_alignment"]
            sheet.row_dimensions[3].height = 28
            if metrics_enabled:
                sheet_starts[id(sheet)] = time.perf_counter()
                sheet_nonempty[id(sheet)] = 0
            return sheet

        note = subtitle if has_rows else empty_note
        sheet = create_sheet(title, note)
        sheets.append((sheet, 0))
        current_count = 0
        for values in row_iterator:
            if current_count >= limit:
                if len(names) == 1:
                    first_name = f"{title} 1"
                    old_title = sheet.title
                    names[0] = first_name
                    sheets[0][0].title = first_name
                    if self._audit_accumulator is not None:
                        self._audit_accumulator.rename_sheet(old_title, first_name)
                    if old_title in classifiers:
                        classifiers[first_name] = build_audit_classifier(
                            first_name,
                            headers,
                            category,
                        )
                        classifiers.pop(old_title)
                next_name = f"{title} {len(names) + 1}"
                names.append(next_name)
                self._register_partitioned_sheets(title, names)
                sheet = create_sheet(next_name, subtitle)
                sheets.append((sheet, 0))
                if self._audit_accumulator is not None:
                    classifiers[next_name] = build_audit_classifier(
                        next_name,
                        headers,
                        category,
                    )
                current_count = 0

            current_count += 1
            sheets[-1] = (sheet, current_count)
            worksheet_row = current_count + 3
            safe_values = tuple(self._safe_value(value) for value in values)
            for column, value in enumerate(safe_values, 1):
                cell = sheet.cell(worksheet_row, column, value)
                if metrics_enabled:
                    sheet_nonempty[id(sheet)] += value not in (None, "")
                if not minimal and not fast:
                    cell.font = styles["body_font"]
                    cell.alignment = styles["body_alignment"]
                if full and current_count % 2 == 0:
                    cell.fill = styles["stripe_fill"]
            accumulator = self._audit_accumulator
            if accumulator is not None:
                classifier = classifiers[sheet.title]
                if classifier.can_produce_review or classifier.can_produce_evidence:
                    accumulator.add_row(
                        sheet.title,
                        headers,
                        safe_values,
                        worksheet_row,
                        category,
                        classifier,
                    )

        for sheet, row_count in sheets:
            last_column = get_column_letter(len(headers))
            sheet.auto_filter.ref = f"A3:{last_column}{max(3, row_count + 3)}"
            sheet.freeze_panes = "A4"
            sizing_started = time.perf_counter() if metrics_enabled else 0.0
            if not minimal:
                self._size_table(sheet, len(headers), row_count)
            if metrics_enabled:
                recorder = getattr(self, "_record_table_metrics", None)
                if recorder is not None:
                    recorder(
                        sheet.title,
                        row_count,
                        len(headers),
                        row_count * len(headers),
                        sheet_nonempty[id(sheet)],
                        sizing_started and sizing_started - sheet_starts[id(sheet)],
                        time.perf_counter() - sizing_started if sizing_started else 0.0,
                    )
        return sheets[0][0]

    def _size_table(
        self,
        sheet: Any,
        column_count: int,
        row_count: int,
    ) -> None:
        header_border = self._table_styles()["header_border"]

        verbose_tokens = (
            "description",
            "message",
            "note",
            "reason",
            "additional settings",
            "audit",
            "migration instruction",
            "raw capture",
            "parse error",
        )

        for column in range(
            1,
            column_count + 1,
        ):
            header = str(
                sheet.cell(
                    3,
                    column,
                ).value
                or ""
            )

            max_length = len(header)

            # Inspect only a bounded number of records. Large source inventories
            # must not make workbook generation progressively more expensive.
            for row in range(4, min(row_count + 4, 4 + self.COLUMN_WIDTH_SAMPLE_ROWS)):
                value = str(
                    sheet.cell(
                        row,
                        column,
                    ).value
                    or ""
                )

                max_length = max(
                    max_length,
                    max(
                        (
                            len(line)
                            for line
                            in value.splitlines()
                        ),
                        default=0,
                    ),
                )

            header_lower = header.lower()

            width_cap = (
                48
                if any(
                    token in header_lower
                    for token
                    in verbose_tokens
                )
                else 32
            )

            sheet.column_dimensions[
                get_column_letter(column)
            ].width = min(
                max(
                    max_length + 2,
                    11,
                ),
                width_cap,
            )

            sheet.cell(3, column).border = header_border

        if not self._is_full_profile() or row_count > self.LARGE_TABLE_ROW_THRESHOLD:
            return

        for row in range(4, row_count + 4):
            max_lines = max(
                (
                    str(
                        sheet.cell(
                            row,
                            column,
                        ).value
                        or ""
                    ).count("\n")
                    + 1
                    for column
                    in range(
                        1,
                        column_count + 1,
                    )
                ),
                default=1,
            )

            sheet.row_dimensions[row].height = min(
                15 * max_lines + 5,
                90,
            )

    def _format_port(self, port: Any) -> str:
        protocol = self._enum_value(port.protocol)
        result = f"{protocol}/{port.port}"
        details = []
        if port.icmptype is not None:
            details.append(f"type={port.icmptype}")
        if port.icmpcode is not None:
            details.append(f"code={port.icmpcode}")
        if details:
            result += f" ({', '.join(details)})"
        return result

    @staticmethod
    def _format_settings(settings: dict[str, Any]) -> str:
        settings = sanitize_source_attributes(settings)

        def format_value(value: Any) -> str:
            if isinstance(value, (list, tuple, set)):
                return " ".join(str(item) for item in value)
            return str(value)

        return "; ".join(
            f"{key.replace('_', '-')}={format_value(value)}"
            for key, value in sorted(settings.items())
        )

    @staticmethod
    def _format_nat_ports(ports: list[Any]) -> str:
        return " ".join(
            f"{port.start}-{port.end}" if port.end is not None else str(port.start)
            for port in ports
        )

    @staticmethod
    def _format_ip_pool_ranges(ranges: list[Any]) -> str:
        return " ".join(f"{item.start_ip} to {item.end_ip}" for item in ranges)

    @staticmethod
    def _format_nat_source_section(item: Any) -> Any:
        attributes = item.source_attributes or {}
        return attributes.get("fmc_nat_section") or attributes.get("section")

    @staticmethod
    def _format_nat_attachment_values(attachments: Any, field: str) -> Any:
        values = []
        for attachment in attachments or []:
            value = getattr(attachment, field, None)
            if value is None and isinstance(attachment, dict):
                value = attachment.get(field)
            if value is not None:
                values.append(str(value))
        return "\n".join(values) or None

    def _nat_attachment_resolution_status(self, item: Any) -> Any:
        source = self._object_extension_value(item, "source_attachments")
        destination = self._object_extension_value(item, "destination_attachments")
        if source is None and destination is None:
            return None
        attachments = list(source or []) + list(destination or [])
        if not attachments:
            return "UNRESOLVED"
        return (
            "RESOLVED"
            if all(
                bool(
                    getattr(attachment, "resolved", None)
                    if not isinstance(attachment, dict)
                    else attachment.get("resolved")
                )
                for attachment in attachments
            )
            else "UNRESOLVED"
        )

    @staticmethod
    def _format_nat_address_ranges(mappings: list[Any], field: str) -> Any:
        values = [getattr(mapping, field) for mapping in mappings]
        if not values:
            return None
        return values[0] if len(values) == 1 else json.dumps(values, ensure_ascii=False)

    @staticmethod
    def _format_pbr_table_routes(routes: list[Any]) -> str:
        return "\n".join(
            " | ".join(
                f"{label}: {route.get(key) or ''}"
                for label, key in (
                    ("Destination", "destination"),
                    ("Next Hop", "next_hop"),
                    ("Output Interface", "outgoing_interface"),
                    ("Priority", "priority"),
                )
            )
            for route in routes
            if isinstance(route, dict)
        )

    @staticmethod
    def _policy_source_settings(policy: Any) -> dict[str, Any]:
        """Include typed source-only ZTNA fields in the policy audit view."""
        settings = dict(policy.source_extra_settings)
        typed_ztna_settings = {
            "ztna-device-ownership": policy.source_ztna_device_ownership,
            "ztna-ems-tag-secondary": policy.source_ztna_ems_tags_secondary,
            "ztna-geo-tag": policy.source_ztna_geo_tags,
            "ztna-policy-redirect": policy.source_ztna_policy_redirect,
            "ztna-tags-match-logic": policy.source_ztna_tags_match_logic,
        }
        settings.update({
            key: value
            for key, value in typed_ztna_settings.items()
            if value not in (None, "", []) and key not in settings
        })
        return settings

    def _safe_value(self, value: Any) -> Any:
        if value is None:
            return ""
        if isinstance(value, bool):
            return "Yes" if value else "No"
        if isinstance(value, (int, float)):
            return value
        if isinstance(value, datetime):
            value = value.isoformat()
        elif isinstance(value, Enum):
            value = value.value
        elif isinstance(value, (list, tuple, set)):
            value = "\n".join(str(self._enum_value(item)) for item in value)
        else:
            value = str(value)

        value = _ILLEGAL_XML_CHARS.sub(" ", value)
        if len(value) > _MAX_CELL_TEXT:
            suffix = "\n[truncated for Excel cell limit]"
            value = value[: _MAX_CELL_TEXT - len(suffix)] + suffix
        if value.lstrip().startswith(_FORMULA_PREFIXES):
            value = "'" + value
        return value

    @staticmethod
    def _enum_value(value: Any) -> Any:
        return value.value if isinstance(value, Enum) else value

    @staticmethod
    def _optional_bool_literal(value: bool | None) -> str | None:
        if value is None:
            return None
        return "TRUE" if value else "FALSE"
